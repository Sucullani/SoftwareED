"""construir_pptx.py — arma la presentación en PowerPoint a partir de la escena que extrae
extraer_pptx.mjs de la versión 2 (estilo PowerPoint) de la defensa.

Cada paso de la presentación HTML se vuelve un clic. Todo lo que es texto, caja, tabla o
imagen es una forma nativa y editable (Calibri); los gráficos SVG, los lienzos del MEF y las
fórmulas son imágenes al doble de resolución. La capa de animación se arma con
plantillas_pptx (el XML que escribe el propio PowerPoint, verificado por COM):

- entradas por grupo (el elemento que se revela en la versión HTML), con el efecto que
  corresponde a su clase: flotar, ascender, barrido, zoom, rebote o palabra por palabra,
  escalonadas en el orden del documento; al llegar a la lámina, en el orden de data-entrada;
- salidas rápidas antes de las entradas, y fundido cruzado sin parpadeo cuando una forma
  cambia de un paso al siguiente (la de abajo aparece, la de arriba se funde);
- bucles hasta el final de la diapositiva: latido y respiración (puntos calientes, insignias
  «se cumple», el punto de SAP2000 en el margen de CR.2, el «Gracias») y los videos en bucle de
  las regiones animadas (carátula, malla de los separadores, funciones de forma, modos de
  Gauss, arquitectura, campo manufacturado, ciclo Q4 → Q9);
- trayectorias de movimiento (los puntos del margen de CR.2 salen del umbral);
- transiciones: Transformación (Morph) entre láminas de contenido: la cabecera y el pie se
  emparejan por nombre («!!…») y se deslizan, el resto se funde en su lugar; Galería en los
  separadores; fundido a negro en la carátula; Cortinas en el cierre; Forma circular en el
  respaldo.

Lleva además los hipervínculos (botones de bloque, filas de la matriz de consistencia,
criterios del veredicto, enlaces de CR.1, «Volver»), las notas del orador con la narración de
cada paso, los dos videos del software y un índice oculto.

    python herramientas/construir_pptx.py [--trabajo=_trabajo/pptx] [--salida=../Defensa_EduFEM_final4.pptx]
                                          [--pasos-separados]   # una diapositiva por paso, sin animación (revisión)
"""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import uuid
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import plantillas_pptx as pp

AQUI = Path(__file__).resolve().parent
FINAL = AQUI.parent
PX = 6350                      # EMU por px de diseño: 1920 px = 13,333 in
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def emu(v: float) -> Emu:
    return Emu(int(round(v * PX)))


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def poner_alfa(elemento, alfa: float):
    """Agrega <a:alpha> a un <a:srgbClr> (transparencia de relleno, línea o texto)."""
    if alfa >= 0.995:
        return
    for c in elemento.iter(qn("a:srgbClr")):
        for viejo in c.findall(qn("a:alpha")):
            c.remove(viejo)
        etree.SubElement(c, qn("a:alpha")).set("val", str(int(round(max(0.0, alfa) * 100000))))


# ----------------------------------------------------------------------------- escena
def clave(it: dict) -> str:
    d = {k: v for k, v in it.items() if k not in ("orden", "entrada", "z", "rid")}
    return json.dumps(d, sort_keys=True, ensure_ascii=False)


def unir_pasos(pasos: list[dict]) -> tuple[list[dict], int]:
    """Une los pasos en una lista de formas con el paso en que aparecen y el último en que siguen.

    Una forma idéntica en pasos consecutivos es la misma forma; si desaparece y vuelve, es otra.
    El orden de apilado respeta el orden del DOM de cada paso (fusión estable)."""
    ocurrencias: list[dict] = []
    activas: dict[str, dict] = {}
    global_orden: list[dict] = []
    n = len(pasos)
    for p, paso in enumerate(pasos):
        vistas = set()
        previa = None
        for it in paso["items"]:
            k = clave(it)
            if k in vistas:          # duplicado exacto dentro del mismo paso
                continue
            vistas.add(k)
            oc = activas.get(k)
            if oc is None or oc["hasta"] != p - 1:
                oc = {"it": it, "desde": p, "hasta": p, "k": k}
                ocurrencias.append(oc)
                activas[k] = oc
                # insertar justo después de la forma que la precede en este paso
                if previa is not None:
                    global_orden.insert(global_orden.index(previa) + 1, oc)
                elif p > 0:
                    global_orden.insert(0, oc)
                else:
                    global_orden.append(oc)
            else:
                oc["hasta"] = p
            previa = oc
    # enlaces arriba de todo, y lo que tenga z-index positivo por encima del resto. Un bucle en video
    # que cae dentro de una región grabada (z 50) va encima de ella: el video de la región lo
    # muestra congelado, el bucle sigue vivo arriba.
    regiones = [oc["it"] for oc in global_orden if oc["it"]["t"] == "region"]

    def sobre_region(it):
        return it.get("clipVideo") and any(
            it["x"] < r["x"] + r["w"] and it["x"] + it["w"] > r["x"] and it["y"] < r["y"] + r["h"] and it["y"] + it["h"] > r["y"]
            for r in regiones)

    def llave(oc):
        it = oc["it"]
        z = 51 if sobre_region(it) else it.get("z", 0)
        return (2 if it["t"] == "enlace" else (1 if z > 0 else 0), z, global_orden.index(oc))
    return sorted(global_orden, key=llave), n


# ----------------------------------------------------------------------------- formas
def _formas(destino):
    """Colección donde se agregan las formas: la de una diapositiva o la de un grupo."""
    return destino.shapes if hasattr(destino, "shapes") and not hasattr(destino, "add_shape") else destino


def forma_caja(slide, it: dict):
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    lados = it["lados"]
    radios = it["radios"]
    alfa = it.get("alfa", 1)
    uniforme = all(lados) and len({(l["w"], l["c"]["hex"], l["estilo"]) for l in lados}) == 1
    bw = lados[0]["w"] if uniforme else 0
    # el borde de PowerPoint va centrado en el contorno; el de CSS, por dentro
    if uniforme:
        x, y, w, h = x + bw / 2, y + bw / 2, max(0.5, w - bw), max(0.5, h - bw)
    rmax = max(radios) if radios else 0
    rmin = min(radios) if radios else 0
    tipo = MSO_SHAPE.RECTANGLE
    girar = 0
    voltearV = False
    if rmax > 0.6:
        if abs(w - h) < 1 and rmin >= min(w, h) / 2 - 1:
            tipo = MSO_SHAPE.OVAL
        elif rmax - rmin <= 10 or rmin > 0.6:
            tipo = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            tl, tr, br, bl = radios
            if tl > 0.6 and tr > 0.6 and br < 0.6 and bl < 0.6:
                tipo = MSO_SHAPE.ROUND_2_SAME_RECTANGLE
            elif bl > 0.6 and br > 0.6 and tl < 0.6 and tr < 0.6:
                tipo, voltearV = MSO_SHAPE.ROUND_2_SAME_RECTANGLE, True
            else:
                tipo = MSO_SHAPE.ROUNDED_RECTANGLE
    shp = _formas(slide).add_shape(tipo, emu(x), emu(y), emu(w), emu(h))
    shp.shadow.inherit = False
    if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        r = max(0.0, rmax - (bw / 2 if uniforme else 0))
        shp.adjustments[0] = min(0.5, r / max(1e-6, min(w, h)))
    elif tipo == MSO_SHAPE.ROUND_2_SAME_RECTANGLE:
        shp.adjustments[0] = min(0.5, rmax / max(1e-6, min(w, h)))
        shp.adjustments[1] = 0.0
        if voltearV:
            shp._element.spPr.get_or_add_xfrm().set("flipV", "1")
    if it.get("giro"):
        shp.rotation = it["giro"]
    # relleno
    fondo, grad = it.get("fondo"), it.get("grad")
    if grad and grad.get("rep"):
        p = [q for q in grad["paradas"] if q["c"]["a"] > 0]
        if len(p) >= 2:
            shp.fill.patterned()
            shp.fill.pattern = MSO_PATTERN_TYPE.WIDE_UPWARD_DIAGONAL
            shp.fill.fore_color.rgb = rgb(p[-1]["c"]["hex"])
            shp.fill.back_color.rgb = rgb(p[0]["c"]["hex"])
        elif fondo:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fondo["hex"])
        else:
            shp.fill.background()
    elif grad and len(grad["paradas"]) >= 2:
        paradas = grad["paradas"]
        # franja dura con el primer color transparente (resaltado tipo «mark»): solo la parte coloreada
        if paradas[0]["c"]["a"] == 0 and paradas[1]["pos"] is not None and len(paradas) == 2:
            f = paradas[1]["pos"] / 100.0
            shp.top = emu(y + h * f); shp.height = emu(h * (1 - f))
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(paradas[1]["c"]["hex"])
            poner_alfa(shp.fill._xPr, paradas[1]["c"]["a"] * alfa)
        else:
            shp.fill.gradient()
            shp.fill.gradient_angle = (grad["ang"] - 90) % 360
            # todas las paradas (la barra jet tiene muchas), con las posiciones de CSS
            largo = w if grad["ang"] % 180 == 90 else h
            pos = [None if q["pos"] is None else (q["pos"] / largo * 100 if q.get("px") else q["pos"]) for q in paradas]
            if pos[0] is None:
                pos[0] = 0.0
            if pos[-1] is None:
                pos[-1] = 100.0
            k = 0
            while k < len(pos):
                if pos[k] is None:
                    j = k
                    while pos[j] is None:
                        j += 1
                    for m in range(k, j):
                        pos[m] = pos[k - 1] + (pos[j] - pos[k - 1]) * (m - k + 1) / (j - k + 1)
                    k = j
                k += 1
            gs = shp.fill._xPr.find(qn("a:gradFill")).find(qn("a:gsLst"))
            for viejo in gs.findall(qn("a:gs")):
                gs.remove(viejo)
            for q, p in zip(paradas, pos):
                g_el = etree.SubElement(gs, qn("a:gs"))
                g_el.set("pos", str(int(round(max(0.0, min(100.0, p)) * 1000))))
                etree.SubElement(g_el, qn("a:srgbClr")).set("val", q["c"]["hex"].lstrip("#").upper())
                poner_alfa(g_el, q["c"]["a"] * alfa)
    elif fondo:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fondo["hex"])
        poner_alfa(shp.fill._xPr, fondo["a"] * alfa)
    else:
        shp.fill.background()
    # borde
    if uniforme:
        l = lados[0]
        shp.line.color.rgb = rgb(l["c"]["hex"])
        shp.line.width = Pt(max(0.25, l["w"] * 0.5))
        if l["estilo"] in ("dashed", "dotted"):
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH if l["estilo"] == "dashed" else MSO_LINE_DASH_STYLE.ROUND_DOT
        poner_alfa(shp.line._ln, l["c"]["a"] * alfa)
    else:
        shp.line.fill.background()
    if it.get("sombra"):
        sombra_exterior(shp, it["sombra"], alfa)
    formas = [shp]
    if not uniforme:
        # bordes de un solo lado: barras o trazos por dentro de la caja
        X, Y, W, H = it["x"], it["y"], it["w"], it["h"]
        for k, l in enumerate(lados):
            if not l:
                continue
            bw = l["w"]
            if l["estilo"] in ("dashed", "dotted"):
                if k == 0: x1, y1, x2, y2 = X, Y + bw / 2, X + W, Y + bw / 2
                elif k == 1: x1, y1, x2, y2 = X + W - bw / 2, Y, X + W - bw / 2, Y + H
                elif k == 2: x1, y1, x2, y2 = X, Y + H - bw / 2, X + W, Y + H - bw / 2
                else: x1, y1, x2, y2 = X + bw / 2, Y, X + bw / 2, Y + H
                ln = _formas(slide).add_connector(1, emu(x1), emu(y1), emu(x2), emu(y2))
                ln.line.color.rgb = rgb(l["c"]["hex"])
                ln.line.width = Pt(max(0.25, bw * 0.5))
                ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH if l["estilo"] == "dashed" else MSO_LINE_DASH_STYLE.ROUND_DOT
                poner_alfa(ln.line._ln, l["c"]["a"] * alfa)
                formas.append(ln)
                continue
            if k == 0: rx, ry, rw, rh = X, Y, W, bw
            elif k == 1: rx, ry, rw, rh = X + W - bw, Y, bw, H
            elif k == 2: rx, ry, rw, rh = X, Y + H - bw, W, bw
            else: rx, ry, rw, rh = X, Y, bw, H
            b = _formas(slide).add_shape(MSO_SHAPE.RECTANGLE, emu(rx), emu(ry), emu(rw), emu(rh))
            b.shadow.inherit = False
            b.fill.solid(); b.fill.fore_color.rgb = rgb(l["c"]["hex"])
            poner_alfa(b.fill._xPr, l["c"]["a"] * alfa)
            b.line.fill.background()
            formas.append(b)
    return formas


def sombra_exterior(shp, s: dict, alfa: float = 1.0):
    """box-shadow difuminada de CSS -> <a:outerShdw> (distancia, dirección y difuminado en EMU)."""
    import math
    spPr = shp._element.spPr
    ef = spPr.find(qn("a:effectLst"))
    if ef is None:
        ef = etree.SubElement(spPr, qn("a:effectLst"))
        ln = spPr.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(ef)
    for viejo in list(ef):
        ef.remove(viejo)
    dist = math.hypot(s["dx"], s["dy"])
    ang = (math.degrees(math.atan2(s["dy"], s["dx"])) % 360) if dist else 90.0
    sh = etree.SubElement(ef, qn("a:outerShdw"))
    sh.set("blurRad", str(int(s["blur"] * PX)))
    sh.set("dist", str(int(dist * PX)))
    sh.set("dir", str(int(ang * 60000)))
    sh.set("algn", "ctr")
    sh.set("rotWithShape", "0")
    c = etree.SubElement(sh, qn("a:srgbClr"))
    c.set("val", s["c"]["hex"].lstrip("#").upper())
    etree.SubElement(c, qn("a:alpha")).set("val", str(int(round(s["c"]["a"] * alfa * 100000))))


ALINEAR = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}


def forma_texto(slide, it: dict):
    L, s = it["interlineado"], it["tamBase"]
    # calibración Chrome ↔ PowerPoint (interlineado exacto, Calibri): la primera línea de
    # PowerPoint cae 0,24·L − 0,325·s px más abajo que la del navegador
    dy = 0.24 * L - 0.325 * s
    tb = _formas(slide).add_textbox(emu(it["x"]), emu(it["y"] - dy), emu(it["w"]), emu(it["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = ALINEAR.get(it["al"], PP_ALIGN.LEFT)
    pPr = p._p.get_or_add_pPr()
    ln = etree.SubElement(pPr, qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPts")).set("val", str(int(round(L / 2 * 100))))
    for s_ in ("a:spcBef", "a:spcAft"):
        e = etree.SubElement(pPr, qn(s_))
        etree.SubElement(e, qn("a:spcPts")).set("val", "0")
    alfa = it.get("alfa", 1)
    for k, linea in enumerate(it["lineas"]):
        for j, run in enumerate(linea):
            txt = run["txt"]
            if j == len(linea) - 1:
                txt = txt.rstrip()
            if not txt:
                continue
            r = p.add_run()
            r.text = txt
            e = run["e"]
            f = r.font
            f.name = e["f"]
            f.size = Pt(e["tam"] / 2)
            f.bold = bool(e["n"])
            f.italic = bool(e["i"])
            f.color.rgb = rgb(e["c"])
            rPr = r._r.get_or_add_rPr()
            if e.get("may"):
                rPr.set("cap", "all")
            if e.get("esp"):
                rPr.set("spc", str(int(round(e["esp"] * 50))))
            if e.get("desp") == "sub":
                rPr.set("baseline", "-25000")
            elif e.get("desp") == "sup":
                rPr.set("baseline", "30000")
            if e.get("sub"):
                rPr.set("u", "sng")
            poner_alfa(rPr, e.get("a", 1) * alfa)
    # los saltos van entre corridas: mover cada <a:br> a su lugar (python-pptx agrega las corridas al final)
    reordenar_saltos(p._p, it["lineas"])
    return [tb]


def reordenar_saltos(p_el, lineas):
    """Reconstruye el párrafo en orden: corridas de la línea 1, <a:br>, corridas de la línea 2…"""
    runs = p_el.findall(qn("a:r"))
    for br in p_el.findall(qn("a:br")):
        p_el.remove(br)
    for r in runs:
        p_el.remove(r)
    k = 0
    for i, linea in enumerate(lineas):
        if i > 0:
            p_el.append(etree.Element(qn("a:br")))
        n = sum(1 for j, run in enumerate(linea) if (run["txt"].rstrip() if j == len(linea) - 1 else run["txt"]))
        for _ in range(n):
            p_el.append(runs[k]); k += 1
    end = p_el.find(qn("a:endParaRPr"))
    if end is not None:
        p_el.remove(end); p_el.append(end)


def forma_imagen(slide, it: dict, raiz: Path):
    ruta = Path(it["ruta"])
    if not ruta.exists():
        print("  falta la imagen", ruta, file=sys.stderr)
        return []
    pic = _formas(slide).add_picture(str(ruta), emu(it["x"]), emu(it["y"]), emu(it["w"]), emu(it["h"]))
    rc = it.get("recorte") or {}
    pic.crop_left, pic.crop_right = rc.get("l", 0), rc.get("r", 0)
    pic.crop_top, pic.crop_bottom = rc.get("t", 0), rc.get("b", 0)
    if it.get("radio", 0) > 0.6:
        # esquinas redondeadas: geometría roundRect con su ajuste (python-pptx no lo expone en imágenes)
        geo = pic._element.spPr.find(qn("a:prstGeom"))
        geo.set("prst", "roundRect")
        av = geo.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geo, qn("a:avLst"))
        for g in list(av):
            av.remove(g)
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(min(0.5, it['radio'] / min(it['w'], it['h'])) * 100000)}")
    if it.get("alfa", 1) < 0.995:
        blip = pic._element.find(".//" + qn("a:blip"))
        etree.SubElement(blip, qn("a:alphaModFix")).set("amt", str(int(it["alfa"] * 100000)))
    return [pic]


MEDIOS = AQUI / "_trabajo" / "medios_pptx"


def clip_local(ruta: str, img_dir: Path) -> Path:
    """Video grabado por el extractor: su ruta absoluta o, si la carpeta de trabajo se movio, el
    mismo archivo en clips/ al lado de img/."""
    p = Path(ruta)
    return p if p.exists() else img_dir.parent / "clips" / p.name


def para_powerpoint(mp4: Path) -> tuple[Path, Path]:
    """Copia del video sin cuadros B y su póster (el cuadro 0 de esa copia), en una caché.

    Con cuadros B, PowerPoint arranca el video 160 ms tarde y nunca muestra los últimos cuadros;
    con un póster distinto del cuadro 0, el video exportado los mezcla al empezar (probado por COM)."""
    MEDIOS.mkdir(parents=True, exist_ok=True)
    st = mp4.stat()
    llave = hashlib.sha1(f"{mp4.resolve()}|{st.st_size}|{st.st_mtime_ns}".encode("utf-8")).hexdigest()[:16]
    dst, poster = MEDIOS / f"{llave}.mp4", MEDIOS / f"{llave}.png"
    if not dst.exists():
        fps = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=r_frame_rate",
                              "-of", "csv=p=0", str(mp4)], capture_output=True, text=True).stdout.strip() or "25/1"
        num, _, den = fps.partition("/")
        g = max(1, round(float(num) / float(den or 1)))
        r = subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(mp4), "-an", "-c:v", "libx264", "-preset", "slow",
                            "-tune", "animation", "-crf", "18", "-bf", "0", "-g", str(g), "-keyint_min", str(g),
                            "-sc_threshold", "0", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(dst)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            raise RuntimeError(f"ffmpeg no pudo recodificar {mp4}: {r.stderr[-300:]}")
    if not poster.exists():
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(dst), "-vf", r"select=eq(n\,0)", "-frames:v", "1", str(poster)],
                       capture_output=True)
    return dst, poster


def forma_raster(slide, it: dict, img_dir: Path):
    """Región fotografiada. Entra como video, con su cuadro de póster, si en la versión HTML se
    anima al llegar el paso (clipPaso: se reproduce una vez con el clic), si se anima sin fin
    (clipVideo: su bucle) o si es un video del software. Dentro de un grupo, siempre imagen."""
    ruta = img_dir / it["archivo"]
    formas = _formas(slide)
    if hasattr(formas, "add_movie"):
        cp = it.get("clipPaso")
        mp4 = clip_local(cp["mp4"], img_dir) if cp else (clip_local(it["clipVideo"], img_dir) if it.get("clipVideo") else
                                          (FINAL / it["video"] if it.get("video") else None))
        if mp4 is not None and mp4.exists():
            video, poster = para_powerpoint(mp4)
            return [formas.add_movie(str(video), emu(it["x"]), emu(it["y"]), emu(it["w"]), emu(it["h"]),
                                     poster_frame_image=str(poster), mime_type="video/mp4")]
    return [formas.add_picture(str(ruta), emu(it["x"]), emu(it["y"]), emu(it["w"]), emu(it["h"]))]


def forma_region(slide, it: dict, img_dir: Path):
    """Región nativa que en la HTML se anima (una tabla que se marca, un contador): su video, que
    se reproduce una vez con el clic y queda en su último cuadro."""
    mp4 = clip_local(it["clipPaso"]["mp4"], img_dir)
    formas = _formas(slide)
    if not hasattr(formas, "add_movie") or not mp4.exists():
        return []
    video, poster = para_powerpoint(mp4)
    return [formas.add_movie(str(video), emu(it["x"]), emu(it["y"]), emu(it["w"]), emu(it["h"]),
                             poster_frame_image=str(poster), mime_type="video/mp4")]


RT_AUDIO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"


def insertar_audio(s, ruta: Path, nombre: str):
    """Audio incrustado (la voz de la lámina en la versión narrada). python-pptx solo sabe
    insertar videos: se inserta como video y se convierte (a:audioFile y relación de audio, que
    hay que crear de nuevo: la de video queda fijada al insertarla). El icono queda fuera de la
    diapositiva, así no se ve al presentar."""
    pic = s.shapes.add_movie(str(ruta), emu(-300), emu(40), emu(64), emu(64),
                             poster_frame_image=str(icono_audio()), mime_type="audio/mpeg")
    vf = pic._element.find(".//" + qn("a:videoFile"))
    viejo = vf.get(qn("r:link"))
    vf.tag = qn("a:audioFile")
    vf.set(qn("r:link"), s.part.relate_to(s.part.rels[viejo].target_part, RT_AUDIO))
    s.part.drop_rel(viejo)
    pic.name = nombre
    return pic


def hipervinculo(shape, accion: str | None = None, destino=None):
    if destino is not None:
        shape.click_action.target_slide = destino
        return
    cNvPr = shape._element.find(".//" + qn("p:cNvPr"))
    h = etree.SubElement(cNvPr, qn("a:hlinkClick"))
    h.set(qn("r:id"), "")
    h.set("action", accion)


def zona(slide, x, y, w, h, nombre: str, elipse: bool = False):
    """Zona sensible invisible: relleno 100 % transparente (sin relleno no recibiría el clic)."""
    shp = _formas(slide).add_shape(MSO_SHAPE.OVAL if elipse else MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    shp.shadow.inherit = False
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    poner_alfa(shp.fill._xPr, 0.0)
    shp.line.fill.background()
    shp.name = nombre
    return shp


def forma_enlace(slide, it: dict, destinos: dict):
    ir = it["ir"]
    if ir != "__siguiente" and ir not in destinos:
        return []
    shp = zona(slide, it["x"], it["y"], it["w"], it["h"], f"Enlace a {ir}")
    if ir == "__siguiente":
        hipervinculo(shp, "ppaction://hlinkshowjump?jump=nextslide")
    else:
        hipervinculo(shp, destino=destinos[ir])
    return [shp]


def crear_formas(slide, it: dict, img_dir: Path, destinos: dict) -> list:
    """Las formas nativas de un elemento de la escena (en una diapositiva o en un grupo)."""
    t = it["t"]
    if t == "caja":
        return forma_caja(slide, it)
    if t == "texto":
        return forma_texto(slide, it)
    if t == "img":
        return forma_imagen(slide, it, FINAL)
    if t == "raster":
        return forma_raster(slide, it, img_dir)
    if t == "region":
        return forma_region(slide, it, img_dir)
    if t == "enlace":
        return forma_enlace(slide, it, destinos)
    return []


# ----------------------------------------------------------------------------- animaciones
# Efecto de entrada de cada grupo según la pista «anim» del extractor (data-pptx-anim o la
# clase del elemento): preset de plantillas_pptx y duración en ms.
ENTRADAS = {
    "flotar": ("float_in_up", 650),
    "subir": ("rise_up", 800),
    "barrido": ("wipe_left", 500),
    "zoom": ("zoom", 500),
    "zoom-suave": ("zoom", 750),
    "desvanecer": ("fade", 600),
    "rebote": ("bounce", 1000),
    "palabras": ("fade", 450),      # el texto entra palabra por palabra; su caja, con un fundido
}
PALABRA_PCT = 12                   # retraso entre palabras, en % de la duración del efecto
# Bucles de énfasis hasta el final de la diapositiva: preset, escala y ms de ida (vuelve solo).
BUCLES = {
    "respirar": ("grow_shrink", 1.05, 1100),
    "latido": ("grow_shrink", 1.18, 520),
}
SALIDA_MS = 300                    # salida de un grupo (antes de que entre lo nuevo)
CAMBIO_MS = 420                    # fundido cruzado de una forma que cambia de un paso al siguiente
TRAYECTO_MS = 1100                 # entrada por trayectoria (data-pptx-desde)
MEDIA_MS = 60000                   # duración nominal de «reproducir» en los videos del software


def tipo_transicion(L: dict) -> tuple[str, int]:
    """Transición de entrada a la lámina y su duración en ms."""
    if L["id"] == "portada":
        return "fade_through_black", 1200
    if L["id"].startswith("sep-"):
        return "gallery_left", 1600
    if L["id"] == "cierre":
        return "curtains", 2000
    if L["respaldo"]:
        return "shape_circle", 1000
    return "morph_objeto", 1000


def transicion_de(L: dict, avance_ms: int | None = None) -> str:
    tipo, dur = tipo_transicion(L)
    return pp.transicion(tipo, dur, avance_ms=avance_ms)


def clase_forma(f) -> str:
    """imagen (imágenes, videos, conectores y grupos: sin grpId ni bldP), texto (cuadro de texto:
    bldP sin animBg) o caja (autoforma con relleno o línea)."""
    el = f._element
    if el.tag in (qn("p:pic"), qn("p:cxnSp"), qn("p:grpSp")):
        return "imagen"
    c = el.find(f"{qn('p:nvSpPr')}/{qn('p:cNvSpPr')}")
    return "texto" if c is not None and c.get("txBox") == "1" else "caja"


def _kw(f, nodo: str, retraso: int) -> dict:
    c = clase_forma(f)
    return {"nodo": nodo, "retraso": int(round(retraso)), "grpId": None if c == "imagen" else 0, "fondo": c == "caja"}


def ef_entrada(f, preset: str, dur: int):
    return lambda nodo, t: pp.entrada(preset, f.shape_id, dur=dur, **_kw(f, nodo, t))


def ef_por_palabra(f, dur: int, n: int):
    return lambda nodo, t: pp.texto_por_palabra("fade", f.shape_id, dur=dur, retraso_pct=PALABRA_PCT,
                                                n_unidades=n, **_kw(f, nodo, t))


def ef_salida(f, preset: str = "fade", dur: int = SALIDA_MS):
    return lambda nodo, t: pp.salida(preset, f.shape_id, dur=None if preset == "disappear" else dur, **_kw(f, nodo, t))


def ef_trayecto(f, dx: float, dy: float, dur: int = TRAYECTO_MS):
    ruta = f"M {dx / 1920:.5f} {dy / 1080:.5f} L 0 0 E"
    return lambda nodo, t: pp.trayectoria(ruta=ruta, spid=f.shape_id, dur=dur, decel=0.8, ptsTypes="AA",
                                          **_kw(f, nodo, t))


def ef_pulso(f, dur: int = 600):
    return lambda nodo, t: pp.enfasis("pulse", f.shape_id, dur=dur, **_kw(f, nodo, t))


def ef_bucle(f, clase: str):
    preset, escala, ms = BUCLES[clase]
    return lambda nodo, t: pp.enfasis(preset, f.shape_id, dur=ms, escala=escala, autorev=True,
                                      repetir="indefinite", **_kw(f, nodo, t))


def ef_reproducir(spid: int, dur_ms: int, bucle: bool = True):
    """«Reproducir» un video o un audio (lo que escribe PowerPoint con Inicio: automáticamente).
    bucle y ms quedan anotados en la fábrica: un bucle no alarga el paso de la versión narrada."""
    def fab(nodo, t):
        return (f'<p:par><p:cTn id="1" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="{nodo}">'
                f'<p:stCondLst><p:cond delay="{int(round(t))}"/></p:stCondLst><p:childTnLst>'
                f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="2" dur="{dur_ms}" fill="hold"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>')
    fab.bucle, fab.ms = bucle, dur_ms
    return fab


class Linea:
    """Efectos de un grupo de clic (o de la llegada) con su instante de inicio en ms."""

    def __init__(self):
        self.items: list[tuple[float, int, object]] = []

    def poner(self, t: float, fabrica):
        self.items.append((t, len(self.items), fabrica))

    def efectos(self, al_clic: bool) -> list:
        orden = sorted(self.items, key=lambda x: (x[0], x[1]))
        return [fab("clickEffect" if (k == 0 and al_clic) else "withEffect", t) for k, (t, _, fab) in enumerate(orden)]


def palabras(it: dict) -> int:
    return max(1, sum(len(run["txt"].split()) for linea in it.get("lineas", []) for run in linea))


def solapa(a: dict, b: dict) -> bool:
    """Las cajas de dos formas se cubren en más del 40 % de la menor."""
    ia, ib = a["it"], b["it"]
    w = min(ia["x"] + ia["w"], ib["x"] + ib["w"]) - max(ia["x"], ib["x"])
    h = min(ia["y"] + ia["h"], ib["y"] + ib["h"]) - max(ia["y"], ib["y"])
    if w <= 0 or h <= 0:
        return False
    menor = min(ia["w"] * ia["h"], ib["w"] * ib["h"])
    return menor > 0 and w * h / menor > 0.4


def bucle_de(r: dict) -> tuple[str | None, int]:
    """(clase, clic desde el que late): «latido@1» -> ('latido', 1); sin «@», desde que entra."""
    b = r["it"].get("bucle")
    if not b:
        return None, 0
    clase, _, k = b.partition("@")
    return (clase if clase in BUCLES else None), (int(k) if k else 0)


def agrupar(regs: list[dict], llave) -> list[list[dict]]:
    """Agrupa por el elemento que se revela (grupo) y ordena los grupos con llave(regs)."""
    grupos: dict[str, list[dict]] = {}
    for r in regs:
        grupos.setdefault(r["it"].get("grupo") or f"_{id(r)}", []).append(r)
    return sorted(grupos.values(), key=llave)


def entrar_grupo(linea: Linea, rs: list[dict], t: float, ini: dict, fin: dict) -> None:
    """Programa la entrada de todas las formas de un grupo en el instante t."""
    anim = rs[0]["it"].get("anim") or "flotar"
    preset, dur = ENTRADAS.get(anim, ENTRADAS["flotar"])
    for r in rs:
        it = r["it"]
        ini[id(r)] = t
        if it["t"] == "enlace":                      # zona sensible invisible: aparece con su grupo
            for f in r["formas"]:
                linea.poner(t, ef_entrada(f, "appear", 1))
            fin[id(r)] = t
            continue
        if it.get("desde"):                          # sale de otro lugar y llega al suyo
            dx, dy = it["desde"]
            for f in r["formas"]:
                linea.poner(t, ef_entrada(f, "fade", 350))
                linea.poner(t, ef_trayecto(f, dx, dy))
            fin[id(r)] = t + TRAYECTO_MS
            continue
        d = dur
        for f in r["formas"]:
            if anim == "palabras" and clase_forma(f) == "texto":
                n = palabras(it)
                linea.poner(t, ef_por_palabra(f, dur, n))
                d = max(d, dur * (1 + PALABRA_PCT / 100 * (n - 1)))
            else:
                linea.poner(t, ef_entrada(f, preset, dur))
        fin[id(r)] = t + d


def animar(regs: list[dict], n: int, al_final: list | None = None) -> tuple[Linea, list[Linea], list[tuple[int, bool]]]:
    """Línea de tiempo de una lámina: lo que pasa al llegar, lo que pasa con cada clic y los
    videos [(spid, en bucle)]. al_final: formas que aparecen con el último clic (las zonas
    sensibles de los botones de un laboratorio: antes de ese paso, el clic sigue la presentación)."""
    vis: list[set] = [set() for _ in range(n)]
    for z, r in enumerate(regs):
        r["z"] = z
        g = r["it"].get("grupo")
        if g:
            for p in range(r["desde"], r["hasta"] + 1):
                vis[p].add(g)
    ini: dict[int, float] = {}
    fin: dict[int, float] = {}
    videos = [(f.shape_id, r.get("bucle", True)) for r in regs if r.get("video_ms") for f in r["formas"]]
    es_clip = lambda r: bool(r["it"].get("clipPaso"))   # noqa: E731

    # 1) al llegar: grupos de data-entrada por su orden; en el mismo orden, en cascada
    llegada = Linea()
    ent0 = [r for r in regs if r["desde"] == 0 and r["it"].get("entrada") is not None and r["it"]["t"] != "enlace"]
    rangos = sorted({r["it"]["entrada"] for r in ent0})
    en_rango: Counter = Counter()
    for rs in agrupar(ent0, lambda rs: (rs[0]["it"]["entrada"], min(r["it"]["orden"] for r in rs))):
        k = rangos.index(rs[0]["it"]["entrada"])
        j = en_rango[k]
        en_rango[k] += 1
        entrar_grupo(llegada, rs, 150 + 220 * k + 70 * min(j, 5), ini, fin)
    # trayectorias sueltas (los puntos del margen de CR.2), de arriba abajo
    tray = [r for r in regs if r["desde"] == 0 and r["it"].get("desde") and id(r) not in ini]
    filas = sorted({round(r["it"]["y"] + r["it"]["h"] / 2) for r in tray})
    for r in tray:
        t = 350 + 140 * filas.index(round(r["it"]["y"] + r["it"]["h"] / 2))
        entrar_grupo(llegada, [r], t, ini, fin)
    for r in regs:
        if r["desde"] == 0 and r.get("video_ms"):
            for f in r["formas"]:
                llegada.poner(ini.get(id(r), 0), ef_reproducir(f.shape_id, r["video_ms"], r.get("bucle", True)))
        clase, kb = bucle_de(r)
        if clase and r["desde"] == 0 and kb == 0:
            for f in r["formas"]:
                llegada.poner(fin.get(id(r), 300) + 150, ef_bucle(f, clase))

    # 2) cada clic: se va lo que sale, se funde lo que cambia, entra lo nuevo
    clics: list[Linea] = []
    for k in range(1, n):
        linea = Linea()
        sal = [r for r in regs if r["hasta"] == k - 1]
        ent = [r for r in regs if r["desde"] == k]
        g_sal = [r for r in sal if r["it"].get("grupo") and r["it"]["grupo"] not in vis[k]]
        g_ent = [r for r in ent if r["it"].get("grupo") and r["it"]["grupo"] not in vis[k - 1]]
        c_sal = [r for r in sal if r not in g_sal]
        c_ent = [r for r in ent if r not in g_ent]
        clips = [r for r in ent if es_clip(r)]
        for r in g_sal:
            for f in r["formas"]:
                instantaneo = r["it"]["t"] in ("enlace", "region") or any(solapa(x, r) for x in clips)
                linea.poner(0, ef_salida(f, "disappear" if instantaneo else "fade"))
        # fundido cruzado: si la forma vieja queda encima, la nueva aparece de golpe debajo y la
        # vieja se funde; si la nueva queda encima, se funde sobre la vieja, que se va al final. Un
        # video de paso empieza en el cuadro de lo que había: aparece de golpe y lo de abajo se va.
        for r in c_ent:
            debajo = r["it"]["t"] == "enlace" or es_clip(r) or any(x["z"] > r["z"] and solapa(x, r) for x in c_sal)
            ini[id(r)] = 0
            fin[id(r)] = 0 if debajo else CAMBIO_MS
            for f in r["formas"]:
                linea.poner(0, ef_entrada(f, "appear", 1) if debajo else ef_entrada(f, "fade", CAMBIO_MS))
        for r in c_sal:
            if r["it"]["t"] in ("enlace", "region") or any(solapa(x, r) for x in clips):
                for f in r["formas"]:
                    linea.poner(0, ef_salida(f, "disappear"))
                continue
            tapada = any(x["z"] > r["z"] and solapa(x, r) for x in c_ent)
            for f in r["formas"]:
                linea.poner(CAMBIO_MS if tapada else 0,
                            ef_salida(f, "disappear") if tapada else ef_salida(f, "fade", CAMBIO_MS))
        grupos = agrupar(g_ent, lambda rs: min(r["it"]["orden"] for r in rs))
        m = len(grupos)
        paso = 120 if m <= 6 else max(40.0, 700 / (m - 1))
        t0 = 250 if g_sal else 0
        for j, rs in enumerate(grupos):
            entrar_grupo(linea, rs, t0 + j * paso, ini, fin)
        for r in ent:
            if r.get("video_ms"):
                for f in r["formas"]:
                    linea.poner(ini[id(r)], ef_reproducir(f.shape_id, r["video_ms"], r.get("bucle", True)))
        for r in regs:
            clase, kb = bucle_de(r)
            if not clase:
                continue
            if r["desde"] == k and kb <= k:
                t = fin[id(r)] + 150
            elif kb == k and r["desde"] < k <= r["hasta"]:
                t = 250
            else:
                continue
            for f in r["formas"]:
                linea.poner(t, ef_bucle(f, clase))
        if not linea.items:
            # paso cuya animación en la HTML vuelve a su estado: un pulso en cascada sobre los
            # gráficos de la lámina conserva el clic de la narración
            vivos = sorted((r for r in regs if r["it"]["t"] == "raster" and r["desde"] <= k <= r["hasta"]),
                           key=lambda r: r["it"]["orden"])
            for j, r in enumerate(vivos):
                for f in r["formas"]:
                    linea.poner(350 * j, ef_pulso(f))
        clics.append(linea)
    if al_final and n > 1:
        for f in al_final:
            clics[-1].poner(0, ef_entrada(f, "appear", 1))
    return llegada, clics, videos


def principal_de(llegada: Linea, clics: list[Linea]) -> list[list]:
    """Grupos de la secuencia principal: el de la llegada arranca solo; cada clic, con un clic."""
    principal = []
    if llegada.items:
        principal.append(llegada.efectos(al_clic=False))
    principal += [c.efectos(al_clic=True) for c in clics if c.items]
    return principal


def armar_timing(principal: list[list], videos: list[tuple[int, bool]], interactivas=(), audios=()) -> str | None:
    """<p:timing> con plantillas_pptx y, al final del tmRoot, un nodo por video: en bucle hasta
    detenerlo, o una sola vez (queda en su último cuadro)."""
    if not principal and not videos and not interactivas and not audios:
        return None
    raiz = etree.fromstring(pp.timing(principal, interactivas=list(interactivas)).encode("utf-8"))
    hijos = raiz.find(f".//{qn('p:cTn')}[@nodeType='tmRoot']/{qn('p:childTnLst')}")
    n = max(int(c.get("id")) for c in raiz.iter(qn("p:cTn")))
    for spid, bucle in videos:
        n += 1
        rep = ' repeatCount="indefinite"' if bucle else ""
        hijos.append(etree.fromstring(
            f'<p:video xmlns:p="{NS_P}"><p:cMediaNode vol="80000"><p:cTn id="{n}"{rep} fill="hold" display="0">'
            f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>'))
    for spid in audios:
        n += 1
        hijos.append(etree.fromstring(
            f'<p:audio xmlns:p="{NS_P}"><p:cMediaNode vol="80000"><p:cTn id="{n}" fill="hold" display="0">'
            f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:audio>'))
    return etree.tostring(raiz, encoding="unicode")


def nombrar(f, base: str, usados: Counter) -> None:
    """Nombre único en la lámina. Las formas de la cabecera y del pie llevan el mismo nombre
    «!!…» en todas las láminas: Morph las empareja y las desliza. Las demás llevan un «!!» con
    el id de su lámina, así Morph no las confunde con otras y las funde en su lugar (probado:
    sin el prefijo, Morph empareja rectángulos de nombre distinto)."""
    usados[base] += 1
    f.name = base if usados[base] == 1 else f"{base}·{usados[base]}"


# ----------------------------------------------------------------------------- armado
def notas(lamina: dict, con_cambio: list[bool] | None = None) -> str:
    """Guion de la lámina, un párrafo por paso, rotulado con el clic de PowerPoint que lo muestra.

    con_cambio[k] dice si el paso k+1 cambia algo en pantalla (si no, no consume clic)."""
    pasos = lamina.get("narracion") or []
    if not pasos:
        return ""
    if len(pasos) == 1:
        return pasos[0]
    partes, clic = [], 0
    for k, t in enumerate(pasos):
        if k == 0:
            rot = "Al entrar"
        elif con_cambio is None or (k - 1 < len(con_cambio) and con_cambio[k - 1]):
            clic += 1
            rot = f"Clic {clic}"
        else:
            rot = f"Sigue con el clic {clic}" if clic else "Sigue al entrar"
        partes.append(f"[{rot}] {t}")
    return "\n\n".join(partes)


def boton_nativo(slide, x, y, w, h, texto, relleno, color_txt, accion=None, destino=None, tam=21):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    shp.shadow.inherit = False
    shp.adjustments[0] = 0.2
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(relleno)
    shp.line.fill.background()
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = emu(10); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = texto
    r.font.name = "Calibri"; r.font.size = Pt(tam / 2); r.font.bold = True; r.font.color.rgb = rgb(color_txt)
    hipervinculo(shp, accion, destino)
    return shp


MORPH_ESTADO_MS = 500              # transición entre los estados de un laboratorio
LUPA_MAX = 10                      # imágenes ampliables por lámina
BLOQUES = {"portada": "Carátula", "intro": "Introducción", "teoria": "Capítulo 1 · Fundamentación teórica",
           "diseno": "Capítulo 2 · Desarrollo de EduFEM", "resultados": "Capítulo 3 · Verificación y validación",
           "conclusiones": "Conclusiones y recomendaciones"}


def nombre_bloque(L: dict) -> str:
    """Nombre del bloque del Taller 6 (secciones de PowerPoint y capítulos del video)."""
    if L["respaldo"]:
        return "Respaldo para preguntas"
    if L["bloque"] in BLOQUES:
        return BLOQUES[L["bloque"]]
    return {"portada": "Carátula", "cierre": "Cierre"}.get(L["id"], (L["bloque"] or "Otros").capitalize())


def huella(it: dict) -> str:
    return hashlib.sha1(clave(it).encode("utf-8")).hexdigest()[:10]


def orden_items(items: list[dict]) -> list[dict]:
    """Orden de apilado de un estado suelto (el mismo criterio que unir_pasos)."""
    ocs, _ = unir_pasos([{"items": items}])
    return [oc["it"] for oc in ocs]


def marcar_video(r: dict) -> None:
    """Duración y modo de reproducción de un video: el de un paso se reproduce una vez; los bucles
    y los videos del software, en bucle."""
    it = r["it"]
    if r["formas"][0]._element.find(".//" + qn("a:videoFile")) is None:
        return
    if it.get("clipPaso"):
        r["video_ms"], r["bucle"] = it["clipPaso"]["ms"], False
    elif it.get("clipVideo") and it.get("clip"):
        c = it["clip"]
        r["video_ms"], r["bucle"] = c["ms"] * (2 if c.get("modo") == "idavuelta" else 1), True
    else:
        r["video_ms"], r["bucle"] = MEDIA_MS, True


def union(bb: list, it: dict) -> list:
    return [min(bb[0], it["x"]), min(bb[1], it["y"]), max(bb[2], it["x"] + it["w"]), max(bb[3], it["y"] + it["h"])]


def en_region(it: dict, rg: dict, margen: float = 4) -> bool:
    """El elemento cae en la región: se cruzan y su centro horizontal está dentro de ella."""
    cx = it["x"] + it["w"] / 2
    return (it["x"] < rg["x"] + rg["w"] + margen and it["x"] + it["w"] > rg["x"] - margen and
            it["y"] < rg["y"] + rg["h"] + margen and it["y"] + it["h"] > rg["y"] - margen and
            rg["x"] - margen <= cx <= rg["x"] + rg["w"] + margen)


def armar_lupas(s, regs: list[dict], prefijo: str) -> list:
    """Imagen ampliable de la HTML (lupa): con un clic se abre en grande sobre un velo oscuro y
    con otro clic se cierra. Devuelve los disparadores: primero los de abrir (así el velo y la
    imagen grande empiezan ocultos), después los de cerrar."""
    abrir, cerrar = [], []
    k = 0
    for r in regs:
        it = r["it"]
        if it["t"] != "img" or not it.get("zoomable") or it["w"] > 1400 or k >= LUPA_MAX:
            continue
        f_img = r["formas"][0]
        velo = _formas(s).add_shape(MSO_SHAPE.RECTANGLE, 0, 0, emu(1920), emu(1080))
        velo.shadow.inherit = False
        velo.fill.solid(); velo.fill.fore_color.rgb = rgb("#0B1222"); poner_alfa(velo.fill._xPr, 0.86)
        velo.line.fill.background()
        esc = min(1720 / it["w"], 860 / it["h"])
        W, H = it["w"] * esc, it["h"] * esc
        x, y = (1920 - W) / 2, (1080 - H) / 2 - 30
        grande = _formas(s).add_picture(str(Path(it["ruta"])), emu(x), emu(y), emu(W), emu(H))
        rc = it.get("recorte") or {}
        grande.crop_left, grande.crop_right = rc.get("l", 0), rc.get("r", 0)
        grande.crop_top, grande.crop_bottom = rc.get("t", 0), rc.get("b", 0)
        grande.line.color.rgb = rgb("#FFFFFF"); grande.line.width = Pt(1.5)
        cierre = [velo, grande]
        if it.get("alt"):
            tb = _formas(s).add_textbox(emu(100), emu(y + H + 18), emu(1720), emu(60))
            tf = tb.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            pa = tf.paragraphs[0]; pa.alignment = PP_ALIGN.CENTER
            run = pa.add_run(); run.text = it["alt"] + "  ·  clic para cerrar"
            run.font.name = "Calibri"; run.font.size = Pt(12); run.font.color.rgb = rgb("#DDE6F3")
            cierre.append(tb)
        for j, f in enumerate(cierre):
            f.name = f"!!{prefijo}:lupa {k}·{j}"
        la = Linea()
        la.poner(0, ef_entrada(velo, "fade", 300))
        la.poner(60, ef_entrada(grande, "zoom", 450))
        for f in cierre[2:]:
            la.poner(250, ef_entrada(f, "fade", 300))
        abrir.append((f_img.shape_id, [la.efectos(al_clic=True)]))
        for disparador in cierre:
            lc = Linea()
            lc.poner(0, ef_salida(grande, "zoom", 300))
            for f in cierre[2:]:
                lc.poner(0, ef_salida(f, "fade", 200))
            lc.poner(80, ef_salida(velo, "fade", 300))
            cerrar.append((disparador.shape_id, [lc.efectos(al_clic=True)]))
        k += 1
    return abrir + cerrar


def armar_disparadores(s, L: dict, est: dict, img_dir: Path, destinos: dict) -> tuple[list, list]:
    """Laboratorio de una dimensión (los ítems de la cobertura, las etapas del procedimiento, fórmula
    o valores...): cada estado es un grupo con sus regiones sobre un fondo opaco, oculto al empezar;
    el clic en un control muestra el grupo de su estado y oculta los demás."""
    nodos, regiones = est["nodos"], est["regiones"]
    ext = {}
    for cl, nd in nodos.items():
        for i, rg in enumerate(regiones):
            bb = [rg["x"], rg["y"], rg["x"] + rg["w"], rg["y"] + rg["h"]]
            for it in nd["items"]:
                if it["t"] not in ("enlace", "region") and en_region(it, rg):
                    bb = union(bb, it)
            ext[(cl, i)] = bb
    grupos = {}
    for cl, nd in nodos.items():
        g = _formas(s).add_group_shape()
        for i, rg in enumerate(regiones):
            a, b = ext[(cl, i)], ext[(est["base"], i)]
            x0, y0, x1, y1 = min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3])
            fondo = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x0), emu(y0), emu(x1 - x0), emu(y1 - y0))
            fondo.shadow.inherit = False
            fondo.fill.solid(); fondo.fill.fore_color.rgb = rgb(rg["fondo"]); fondo.line.fill.background()
        for it in orden_items(nd["items"]):
            if it["t"] in ("enlace", "region") or not any(en_region(it, rg) for rg in regiones):
                continue
            crear_formas(g, it, img_dir, destinos)
        g.name = f"!!{L['id']}:estado {cl}"
        grupos[cl] = g
    zonas = []
    for c in nodos[est["base"]]["controles"]:
        if c["destino"] not in grupos:
            continue
        z = zona(s, c["x"], c["y"], c["w"], c["h"], f"!!{L['id']}:botón {c['dim']}={c['valor']}", elipse=c.get("forma") == "elipse")
        zonas.append((z, c["destino"]))
    # disparador ficticio, fuera de la diapositiva: su secuencia nunca corre, pero al ser la primera
    # que nombra cada grupo (con una entrada), PowerPoint los deja ocultos al empezar
    ficticio = zona(s, -300, -300, 12, 12, f"!!{L['id']}:disparador inicial")
    l0 = Linea()
    for g in grupos.values():
        l0.poner(0, ef_entrada(g, "appear", 1))
    inter = [(ficticio.shape_id, [l0.efectos(al_clic=True)])]
    for z, dest in zonas:
        l = Linea()
        l.poner(0, ef_entrada(grupos[dest], "appear", 1))
        for cl, g in grupos.items():
            if cl != dest:
                l.poner(0, ef_salida(g, "disappear"))
        inter.append((z.shape_id, [l.efectos(al_clic=True)]))
    return [z for z, _ in zonas], inter


def etiqueta_estado(e: dict) -> str:
    return " · ".join(f"{k} = {v}" for k, v in e.items())


def llenar_estado(s, L: dict, est: dict, cl: str, img_dir: Path, destinos: dict, estados_diapo: dict,
                  siguiente=None) -> None:
    """Diapositiva de un estado de un laboratorio (oculta): lo que muestra la lámina en su último
    paso con ese estado, y sus controles enlazados a las diapositivas de los otros estados. Las
    formas iguales llevan el mismo nombre en todas, así Morph las deja quietas y funde lo que cambia."""
    nd = est["nodos"][cl]
    if nd.get("fondo"):
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(nd["fondo"]["hex"])
    usados: Counter = Counter()
    regs = []
    for it in orden_items(nd["items"]):
        formas = crear_formas(s, it, img_dir, destinos)
        if not formas:
            continue
        for f in formas:
            nombrar(f, it.get("nombre") or f"!!{L['id']}:{huella(it)}", usados)
        r = {"it": it, "formas": formas}
        marcar_video(r)
        regs.append(r)
    for c in nd["controles"]:
        dest = estados_diapo.get((L["id"], c["destino"]))
        if dest is None:
            continue
        z = zona(s, c["x"], c["y"], c["w"], c["h"], f"!!{L['id']}:botón {c['dim']}={c['valor']}", elipse=c.get("forma") == "elipse")
        hipervinculo(z, destino=dest)
    inter = armar_lupas(s, regs, L["id"])
    videos = [(f.shape_id, r.get("bucle", True)) for r in regs if r.get("video_ms") for f in r["formas"]]
    llegada = Linea()
    for r in regs:
        if r.get("video_ms"):
            for f in r["formas"]:
                llegada.poner(0, ef_reproducir(f.shape_id, r["video_ms"], r.get("bucle", True)))
    pp.inyectar(s._element, transicion=pp.transicion("morph_objeto", MORPH_ESTADO_MS),
                timing=armar_timing(principal_de(llegada, []), videos, inter))
    # salida explícita del laboratorio: a la lámina que sigue en la presentación
    if siguiente is not None:
        b = boton_nativo(s, 100, 958, 190, 40, "Seguir ›", "#1B2A4A", "#FFFFFF", destino=siguiente, tam=20)
        b.name = f"!!{L['id']}:seguir"
    s._element.set("show", "0")
    s.notes_slide.notes_text_frame.text = (
        f"Laboratorio «{L['titulo']}» · {etiqueta_estado(nd['estado'])}. Los botones cambian el estado "
        "(transición Transformación); «Seguir ›» continúa con la presentación.")


def poner_secciones(prs, secciones: list[tuple[str, list]]) -> None:
    """Secciones de PowerPoint (bloques del Taller 6) en presentation.xml."""
    NS14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    pres = prs.part._element
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn("p:extLst"))
    for e in ext_lst.findall(qn("p:ext")):
        if e.get("uri") == "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}":
            ext_lst.remove(e)
    ext = etree.SubElement(ext_lst, qn("p:ext"))
    ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    lst = etree.SubElement(ext, f"{{{NS14}}}sectionLst", nsmap={"p14": NS14})
    for nombre, diapos in secciones:
        sec = etree.SubElement(lst, f"{{{NS14}}}section")
        sec.set("name", nombre)
        sec.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        ids = etree.SubElement(sec, f"{{{NS14}}}sldIdLst")
        for d in diapos:
            etree.SubElement(ids, f"{{{NS14}}}sldId").set("id", str(d.slide_id))


def construir(escena: dict, img_dir: Path, salida: Path, separados: bool, narrada: dict | None = None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(1920), emu(1080)
    blanco = prs.slide_layouts[6]
    laminas = escena["laminas"]
    # 1) crear todas las diapositivas primero (los hipervínculos necesitan el destino); las de los
    #    estados de un laboratorio van justo después de su lámina, ocultas
    diapos: dict[str, object] = {}
    estados_diapo: dict[tuple, object] = {}
    plan = []
    for L in laminas:
        ocs, n = unir_pasos(L["pasos"])
        if separados:
            for p in range(n):
                s = prs.slides.add_slide(blanco)
                diapos.setdefault(L["id"], s)
                plan.append((L, s, [oc for oc in ocs if oc["desde"] <= p <= oc["hasta"]], p, None))
            continue
        s = prs.slides.add_slide(blanco)
        diapos[L["id"]] = s
        plan.append((L, s, ocs, None, None))
        if narrada is None:
            for est in L.get("estados", []):
                if est["via"] != "diapositivas":
                    continue
                for cl in est["nodos"]:
                    se = prs.slides.add_slide(blanco)
                    estados_diapo[(L["id"], cl)] = se
                    plan.append((L, se, None, None, (est, cl)))
    indice = prs.slides.add_slide(blanco)
    destinos = dict(diapos)
    destinos["__indice"] = indice
    objetivos_de_enlace = {oc["it"]["ir"] for _, _, ocs, _, e in plan if ocs for oc in ocs if oc["it"]["t"] == "enlace"}
    resumen: Counter = Counter()
    # 2) llenar cada diapositiva
    for L, s, ocs, p_sep, est in plan:
        if est is not None:
            i = next(j for j, x in enumerate(laminas) if x["id"] == L["id"])
            sig = next((diapos[x["id"]] for x in laminas[i + 1:] if not x["respaldo"]), None)
            llenar_estado(s, L, est[0], est[1], img_dir, destinos, estados_diapo, siguiente=sig)
            resumen["estados (diapositivas)"] += 1
            continue
        fondo = L["pasos"][0].get("fondo")
        if fondo:
            s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(fondo["hex"])
        n = len(L["pasos"])
        regs = []
        usados: Counter = Counter()
        for oc in ocs:
            it = oc["it"]
            formas = crear_formas(s, it, img_dir, destinos)
            if not formas:
                continue
            # «!!» + lámina + huella del contenido: Morph no empareja formas de nombre distinto si
            # las dos llevan «!!» (sin él, empareja rectángulos al azar y los hace volar); entre la
            # lámina y las diapositivas de sus estados, lo que no cambia se llama igual y queda quieto
            base = it.get("nombre") or f"!!{L['id']}:{huella(it)}"
            for f in formas:
                nombrar(f, base, usados)
            r = {"it": it, "formas": formas, "desde": oc["desde"], "hasta": oc["hasta"]}
            marcar_video(r)
            regs.append(r)
        con_cambio = None
        if not separados:
            al_final, inter = [], []
            if narrada is None:
                for e in L.get("estados", []):
                    base_n = e["nodos"][e["base"]]
                    if e["via"] == "diapositivas":
                        for c in base_n["controles"]:
                            dest = estados_diapo.get((L["id"], c["destino"]))
                            if dest is None:
                                continue
                            z = zona(s, c["x"], c["y"], c["w"], c["h"], f"!!{L['id']}:botón {c['dim']}={c['valor']}",
                                     elipse=c.get("forma") == "elipse")
                            hipervinculo(z, destino=dest)
                            al_final.append(z)
                    else:
                        zs, i2 = armar_disparadores(s, L, e, img_dir, destinos)
                        al_final += zs
                        inter += i2
                        resumen["estados (disparadores)"] += len(e["nodos"])
                inter += armar_lupas(s, regs, L["id"])
            llegada, clics, videos = animar(regs, n, al_final)
            con_cambio = [bool(c.items) for c in clics]
            if narrada is None:
                principal = principal_de(llegada, clics)
                tr = transicion_de(L)
                audios = ()
            else:
                principal, tr, audios = narrada["programar"](L, s, llegada, clics)
            pp.inyectar(s._element, transicion=tr, timing=armar_timing(principal, videos, inter, audios))
            for grupo in principal:
                for e in grupo:
                    if isinstance(e, pp.Efecto):
                        resumen["bucles" if e.repetir == "indefinite" else e.clase] += 1
            resumen["videos en bucle"] += sum(1 for _, b in videos if b)
            resumen["videos de paso"] += sum(1 for _, b in videos if not b)
            resumen["disparadores"] += len(inter)
        # notas del orador: la narración de cada paso
        if p_sep is None:
            txt = notas(L, con_cambio)
            if narrada is None and L.get("estados"):
                dims = sorted({c["dim"] for e in L["estados"] for c in e["nodos"][e["base"]]["controles"]})
                txt += ("\n\n[Laboratorio] Después del último clic, los botones de la lámina funcionan ("
                        + ", ".join(dims) + "): cada uno muestra su estado, como en la versión HTML. "
                        "Avanzar sigue con la presentación.")
            if narrada is None and any(oc["it"].get("zoomable") and oc["it"]["w"] <= 1400 for oc in ocs):
                txt += "\n\n[Lupa] Un clic sobre una captura la amplía; otro clic la cierra."
        else:
            txt = (L.get("narracion") or [])[p_sep] if p_sep < len(L.get("narracion") or []) else ""
        if txt:
            s.notes_slide.notes_text_frame.text = txt
        # «Volver» en las láminas a las que se llega por un enlace y en el respaldo
        if not separados and narrada is None and (L["id"] in objetivos_de_enlace or L["respaldo"]) and not L.get("sinPie"):
            b = boton_nativo(s, 1480, 22, 170, 40, "↩ Volver", "#E8741A", "#FFFFFF",
                             accion="ppaction://hlinkshowjump?jump=lastslideviewed")
            b.name = "!!volver"
        if L["respaldo"] and not separados:
            s._element.set("show", "0")
    # 3) índice (oculto): un botón por lámina
    armar_indice(indice, laminas, diapos)
    indice._element.set("show", "0")
    pp.inyectar(indice._element, transicion=pp.transicion("fade_smoothly", 750))
    # 4) secciones por bloque (las diapositivas de estado, con su lámina)
    if not separados:
        secciones: list[tuple[str, list]] = []
        for L, s, _, _, _ in plan:
            nombre = nombre_bloque(L)
            if not secciones or secciones[-1][0] != nombre:
                secciones.append((nombre, []))
            secciones[-1][1].append(s)
        secciones.append(("Índice", [indice]))
        poner_secciones(prs, secciones)
    prs.save(salida)
    if resumen:
        print("  " + " · ".join(f"{k} {v}" for k, v in sorted(resumen.items())))
    return len(prs.slides._sldIdLst)


# ----------------------------------------------------------------------------- versión narrada
# Una copia que se presenta sola: cada clic pasa a ser automático, en el instante en que la voz
# llega a ese paso (lo mismo que el video HTML: cada paso dura su pista más la pausa, o lo que
# tarde su animación); cada diapositiva lleva la voz de sus pasos como un audio que arranca solo
# y avanza al terminar. PowerPoint la exporta como video (grabar_video_pptx.py).
EXTRA_MAX = 8.0                    # s: lo más que se alarga un paso esperando a su animación


def icono_audio() -> Path:
    """Icono del audio de la voz (queda fuera de la diapositiva: no se ve al presentar)."""
    ruta = AQUI / "_trabajo" / "icono_audio.png"
    if not ruta.exists():
        from PIL import Image, ImageDraw
        ruta.parent.mkdir(parents=True, exist_ok=True)
        im = Image.new("RGBA", (96, 96), (0, 0, 0, 0))
        d = ImageDraw.Draw(im)
        d.rounded_rectangle((4, 4, 92, 92), 18, fill=(27, 42, 74, 255))
        d.polygon([(24, 38), (38, 38), (56, 22), (56, 74), (38, 58), (24, 58)], fill=(255, 255, 255, 255))
        d.arc((50, 30, 78, 66), -60, 60, fill=(232, 116, 26, 255), width=5)
        im.save(ruta)
    return ruta


def cargar_narracion() -> dict:
    t = (FINAL / "assets" / "data" / "narracion.js").read_text(encoding="utf-8")
    return json.loads(t[t.index("{"): t.rindex("}") + 1])


def fin_linea(linea: Linea) -> float:
    """ms en que termina lo último de una línea (los bucles no cuentan)."""
    fin = 0.0
    for t, _, fab in linea.items:
        if getattr(fab, "bucle", None) is not None:          # reproducir un medio
            if not fab.bucle:
                fin = max(fin, t + fab.ms)
            continue
        e = fab("withEffect", t)
        if isinstance(e, pp.Efecto) and e.repetir != "indefinite":
            fin = max(fin, t + (e.dur_total or 0))
    return fin


def armar_audio(pistas: list[tuple[Path, float]], destino: Path) -> float:
    """Une las pistas de voz, cada una rellenada con silencio hasta la duración de su paso, en un
    solo audio (mp3). Devuelve su duración en s."""
    args = ["ffmpeg", "-v", "error", "-y"]
    filtros, etiquetas = [], []
    for k, (ruta, dur) in enumerate(pistas):
        if ruta is not None and ruta.exists():
            args += ["-i", str(ruta)]
        else:
            args += ["-f", "lavfi", "-t", f"{dur:.3f}", "-i", "anullsrc=r=24000:cl=mono"]
        filtros.append(f"[{k}:a]aresample=24000,aformat=channel_layouts=mono,apad=whole_dur={dur:.3f},atrim=0:{dur:.3f}[a{k}]")
        etiquetas.append(f"[a{k}]")
    filtros.append("".join(etiquetas) + f"concat=n={len(pistas)}:v=0:a=1[s]")
    args += ["-filter_complex", ";".join(filtros), "-map", "[s]", "-c:a", "libmp3lame", "-b:a", "96k", str(destino)]
    r = subprocess.run(args, capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError("ffmpeg no pudo armar el audio: " + r.stderr[-400:])
    return sum(d for _, d in pistas)


class Narrada:
    """Programación de la versión narrada y su línea de tiempo (para subtítulos y capítulos)."""

    def __init__(self, trabajo: Path):
        self.nar = cargar_narracion()
        self.pausa = float(self.nar.get("pausa_s", 0.6))
        self.dir = trabajo / "narrada"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.laminas: list[dict] = []

    def programar(self, L: dict, s, llegada: Linea, clics: list[Linea]):
        if L["respaldo"]:   # oculto: no entra en el video; queda como en la versión con clics
            return principal_de(llegada, clics), transicion_de(L), []
        pasos_n = self.nar["laminas"].get(L["id"], {}).get("pasos", [])
        lineas = [llegada] + clics
        total = Linea()
        T, pistas, reg = 0.0, [], []
        for k, lin in enumerate(lineas):
            p = pasos_n[k] if k < len(pasos_n) else None
            voz = float(p["dur"]) if p else 2.0
            anim = fin_linea(lin) / 1000
            D = max(voz + self.pausa, min(anim + 0.35, voz + self.pausa + EXTRA_MAX))
            for t, _, fab in lin.items:
                total.poner(T * 1000 + t, fab)
            pistas.append((FINAL / "assets" / "audio" / p["audio"] if p and p.get("audio") else None, D))
            reg.append({"t": round(T, 3), "dur": round(D, 3), "texto": p["texto"] if p else "",
                        "subtitulos": p.get("subtitulos", []) if p else []})
            T += D
        audio = self.dir / f"{L['id']}.mp3"
        armar_audio(pistas, audio)
        forma = insertar_audio(s, audio, f"!!{L['id']}:voz")
        total.poner(0, ef_reproducir(forma.shape_id, int(T * 1000), bucle=False))
        tipo, dur_tr = tipo_transicion(L)
        self.laminas.append({"id": L["id"], "titulo": L["titulo"], "bloque": L["bloque"],
                             "bloque_nombre": nombre_bloque(L),
                             "transicion_ms": dur_tr, "dur": round(T, 3), "pasos": reg})
        return [total.efectos(al_clic=False)], transicion_de(L, avance_ms=int(round(T * 1000))), [forma.shape_id]

    def escribir_linea(self, ruta: Path) -> None:
        """Línea de tiempo del video exportado: dónde empieza cada lámina (t0, s) y cuánto tardan en
        empezar su animación y su voz (retraso: la transición). La usa grabar_video_pptx.py."""
        t = 0.0
        for L in self.laminas:
            tr = L["transicion_ms"] / 1000
            L["t0"], L["retraso"] = round(t, 3), round(tr, 3)
            t += tr + L["dur"]
        ruta.write_text(json.dumps({"laminas": self.laminas}, ensure_ascii=False, indent=1), encoding="utf-8")


def armar_indice(s, laminas, diapos):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb("#FFFFFF")
    tb = s.shapes.add_textbox(emu(89), emu(60), emu(1500), emu(80))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = "Índice de la defensa"
    r.font.name = "Calibri"; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = rgb("#1B2A4A")
    regla = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(89), emu(150), emu(116), emu(6))
    regla.fill.solid(); regla.fill.fore_color.rgb = rgb("#E8741A"); regla.line.fill.background()
    principales = [L for L in laminas if not L["respaldo"]]
    respaldo = [L for L in laminas if L["respaldo"]]
    filas = principales + respaldo
    cols, alto = 3, 27
    por_col = (len(filas) + cols - 1) // cols
    for k, L in enumerate(filas):
        c, f = divmod(k, por_col)
        x, y = 89 + c * 590, 190 + f * (alto + 3)
        n = principales.index(L) + 1 if L in principales else None
        etiqueta = f"{n:>2}. {L['titulo']}" if n else f"R. {L['titulo']}"
        shp = s.shapes.add_textbox(emu(x), emu(y), emu(570), emu(alto))
        tf = shp.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        run = tf.paragraphs[0].add_run(); run.text = etiqueta
        run.font.name = "Calibri"; run.font.size = Pt(10.5)
        run.font.color.rgb = rgb("#1F2937" if n else "#1F77B4")
        shp.click_action.target_slide = diapos[L["id"]]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--trabajo", default=str(AQUI / "_trabajo" / "pptx"))
    ap.add_argument("--salida", default=str(FINAL / "Defensa_EduFEM_final4.pptx"))
    ap.add_argument("--pasos-separados", action="store_true")
    ap.add_argument("--narrada", action="store_true",
                    help="versión que se presenta sola, con la voz (-> ..._narrada.pptx y su línea de tiempo)")
    a = ap.parse_args()
    trabajo = Path(a.trabajo)
    escena = json.loads((trabajo / "escena.json").read_text(encoding="utf-8"))
    salida = Path(a.salida)
    narrada = None
    if a.narrada:
        if a.salida == ap.get_default("salida"):
            salida = FINAL / "Defensa_EduFEM_final4_narrada.pptx"
        narr = Narrada(trabajo)
        narrada = {"programar": narr.programar}
    n = construir(escena, trabajo / "img", salida, a.pasos_separados, narrada)
    if narrada:
        narr.escribir_linea(salida.with_suffix(".linea.json"))
    print(f"[OK] {salida} · {n} diapositivas · {salida.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
