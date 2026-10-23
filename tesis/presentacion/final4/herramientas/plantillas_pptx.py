# -*- coding: utf-8 -*-
"""
plantillas_pptx — XML de animaciones y transiciones de PowerPoint, tal como lo escribe
PowerPoint 16.0 (Microsoft 365, build 18730).

Cada preset de este modulo se extrajo de un .pptx que PowerPoint creo por COM
(referencia.pptx, ver catalogo.md) y se verifico abriendo con PowerPoint un .pptx generado
con estas plantillas (prueba.pptx -> validacion.txt). Solo depende de la biblioteca
estandar y de lxml.

Uso minimo con python-pptx
--------------------------
    from pptx import Presentation
    import plantillas_pptx as pp

    s = prs.slides.add_slide(prs.slide_layouts[6])
    caja = s.shapes.add_shape(...)             # caja.shape_id -> spid
    e1 = pp.entrada("fly_bottom", caja.shape_id, nodo="withEffect")      # arranca solo
    e2 = pp.enfasis("spin", caja.shape_id, nodo="afterEffect",
                    repetir="indefinite", dur=2000)                      # bucle infinito
    e3 = pp.salida("fade", caja.shape_id)                                # en el clic
    t = pp.timing([[e1, e2], [e3]])            # grupos de clic
    tr = pp.transicion("vortex_left", dur=1200)
    pp.inyectar(s._element, transicion=tr, timing=t)

Convenciones
------------
* Tiempos en milisegundos (enteros). ``dur=None`` = duracion por defecto del preset.
* ``nodo``: "clickEffect" (Al hacer clic), "withEffect" (Con la anterior),
  "afterEffect" (Despues de la anterior).
* ``repetir``: None | n (int/float, veces) | "indefinite" (hasta el final de la diapositiva)
  | "untilNextClick" (hasta el siguiente clic).
* Las funciones de efecto devuelven un ``Efecto`` (subclase de ``str``: el ``<p:par>`` del
  efecto con prefijos p:/a: sin declarar) con metadatos que usa ``timing()``.
* ``timing()`` devuelve ``<p:timing>`` completo (con xmlns) y renumera ids de cTn, grpId y
  arma ``<p:bldLst>``. ``transicion()`` devuelve ``<p:transition>`` o el
  ``<mc:AlternateContent>`` que usa PowerPoint.
* ``inyectar()`` coloca ambos en el orden correcto dentro de ``<p:sld>``:
  cSld, clrMapOvr, (transition | mc:AlternateContent), timing, extLst.

Verificacion: 438 de 443 comparaciones C14N identicas al XML de PowerPoint (las 5 restantes:
Grow/Shrink escrito en su forma limpia ``<p:by>`` y 1 ms de redondeo en Bounce); prueba.pptx
(136 diapositivas, 231 efectos de los 149 presets + trayectoria personalizada, 102 tipos de
transicion + avance automatico + transicion con animaciones) abre sin reparar, 367/367
comprobaciones COM correctas y PowerPoint re-guarda el mismo XML en las 136 diapositivas.

English summary
---------------
entrada/salida/enfasis(type, spid, **kw) -> entrance/exit/emphasis effect <p:par>;
trayectoria(path, spid=..., preset=...) -> motion path; texto_por_palabra(...) -> by word/letter;
efecto(key, spid, **kw) -> any preset by key. Keyword args: ids (cTn id generator), nodo
(nodeType: clickEffect/withEffect/afterEffect), retraso (delay ms), dur (duration ms), grpId
(None for pictures), repetir (None | count | "indefinite" | "untilNextClick"), autorev,
rep_dur (repeatDur ms), accel/decel (0..1), suave_inicio/suave_fin, rebobinar (fill=remove),
unidad ("wd"/"lt"/None), retraso_unidad (% between words), n_unidades, parrafo (paragraph
index), fondo (shape has fill/line -> animBg), angulo, escala, color, opacidad, fuente,
radios, ruta. timing(principal=[[...click group...], ...], interactivas=[(trigger_spid,
[[...]])], sin_fondo={spids}) -> full <p:timing>. transicion(key, dur_ms, avance_ms,
avance_clic) -> <p:transition> or <mc:AlternateContent>. inyectar(slide._element, ...).
"""
import math
import re
from fractions import Fraction

from lxml import etree

__all__ = ["NS", "Ids", "Efecto", "entrada", "salida", "enfasis", "trayectoria",
           "texto_por_palabra", "efecto", "timing", "transicion", "inyectar", "elemento",
           "nombre_morph", "PRESETS", "TRANSICIONES", "claves", "esperado_com",
           "INDEFINIDO_COM"]

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p15": "http://schemas.microsoft.com/office/powerpoint/2012/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
}
_P = "{%s}" % NS["p"]
_A = "{%s}" % NS["a"]
_MC = "{%s}" % NS["mc"]
_DECL = " ".join(f'xmlns:{k}="{v}"' for k, v in NS.items())
#: Valor que devuelve COM (Timing.RepeatCount) para repeatCount="indefinite".
INDEFINIDO_COM = -2147483648
_COLORES_TEMA = {"accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "tx1", "tx2",
                 "bg1", "bg2", "dk1", "dk2", "lt1", "lt2", "hlink", "folHlink"}
NODOS = ("clickEffect", "withEffect", "afterEffect")


# ======================================================================================
# utilidades
# ======================================================================================
class Ids:
    """Generador de ids de ``p:cTn``. ``timing()`` renumera todo al final, asi que basta con
    que sean unicos mientras se construye; se admite para mantener ids legibles."""

    def __init__(self, inicio=1):
        self._n = inicio - 1

    def __call__(self):
        self._n += 1
        return self._n

    @property
    def ultimo(self):
        return self._n


_IDS_GLOBAL = Ids(1000)


class Efecto(str):
    """XML de un efecto (``<p:par>``) + metadatos: spid, nodo, clase, dur (ms, nominal),
    retraso, dur_total (ms que ocupa en la cadena), grp (lleva grpId), parrafo, fondo."""


def elemento(xml):
    """Parsea un fragmento con prefijos p:/a:/mc:/p14:/p15:/p159: y devuelve el elemento."""
    w = etree.fromstring(f"<w {_DECL}>{xml}</w>".encode("utf-8"))
    return w[0]


def _ser_limpio(el):
    """Serializa sin declaraciones xmlns (fragmento con prefijos p:/a:)."""
    s = etree.tostring(el, encoding="unicode")
    return re.sub(r'\s+xmlns:\w+="[^"]*"', "", s)


def _mitad_arriba(x):
    """Redondeo 'mitad hacia arriba' con aritmetica exacta (como PowerPoint al reescalar)."""
    return int(math.floor(Fraction(x) + Fraction(1, 2)))


def nombre_morph(nombre):
    """Nombre de forma que fuerza el emparejamiento de Transformacion (Morph): '!!nombre'."""
    return nombre if nombre.startswith("!!") else "!!" + nombre


# ======================================================================================
# efectos
# ======================================================================================
def claves(familia=None):
    """Lista de claves de presets ('entrada.fade', ...), opcionalmente de una familia."""
    return [k for k in PRESETS if familia is None or k.startswith(familia + ".")]


def _resolver(familia, tipo, direccion=None):
    if direccion is not None:
        tipo = f"{tipo}_{direccion}"
    clave = f"{familia}.{tipo}"
    if clave in PRESETS:
        return clave
    cands = [k for k in PRESETS if k.startswith(clave + "_")]
    if cands:
        return cands[0]
    raise KeyError(f"preset desconocido: {clave}. Disponibles: {', '.join(claves(familia))}")


def _escalar(body, d0, d):
    """Reescala los tiempos de los comportamientos como PowerPoint cuando se cambia la
    duracion del efecto: cada dur y cada delay se multiplican por d/d0 (redondeo mitad
    hacia arriba); los <p:set> instantaneos (dur=1) al final quedan en d-1; los
    comportamientos con dur="indefinite" pasan a durar d."""
    if d0 is None:  # preset de duracion indefinida ("hasta el final de la diapositiva")
        if d in (None, "indefinite"):
            return
        for c in body.iter(_P + "cTn"):
            if c.get("dur") == "indefinite":
                c.set("dur", str(d))
        return
    if d is None or d == d0:
        return
    f = Fraction(d, d0)
    for c in body.iter(_P + "cTn"):
        dur = c.get("dur")
        cond = c.find(f"{_P}stCondLst/{_P}cond")
        delay = cond.get("delay") if cond is not None else None
        if dur == "1":
            if delay and delay.isdigit():
                if int(delay) + 1 == d0:  # set al final (salidas; y Appear/Disappear, d0 = 1)
                    cond.set("delay", str(d - 1))
                elif int(delay) > 0:
                    cond.set("delay", str(_mitad_arriba(int(delay) * f)))
            continue
        if dur and dur.isdigit():
            c.set("dur", str(_mitad_arriba(int(dur) * f)))
        if delay and delay.isdigit():
            cond.set("delay", str(_mitad_arriba(int(delay) * f)))


def _ajustes(body, cTn, clave, kw):
    """Parametros propios de cada efecto (angulo, escala, color, opacidad, radios, ruta...)."""
    if kw.get("angulo") is not None:
        for r in body.iter(_P + "animRot"):
            r.set("by", str(int(round(kw["angulo"] * 60000))))
    if kw.get("escala") is not None:
        e = kw["escala"]
        ex, ey = (e, e) if not isinstance(e, (tuple, list)) else e
        for s in body.iter(_P + "animScale"):
            for t in ("by", "from", "to"):
                for x in s.findall(_P + t):
                    s.remove(x)
            by = etree.SubElement(s, _P + "by")
            by.set("x", str(int(round(ex * 100000))))
            by.set("y", str(int(round(ey * 100000))))
            s.remove(by)
            s.insert(1, by)  # orden del esquema: cBhvr, by, from, to
    if kw.get("color") is not None:
        col = kw["color"].lstrip("#")
        tema = col in _COLORES_TEMA
        for c in list(body.iter(_A + "srgbClr", _A + "schemeClr")):
            if tema:  # color del tema: <a:schemeClr val="accent2"/> (lo que PowerPoint pone por defecto)
                nuevo = etree.Element(_A + "schemeClr", nsmap={"a": NS["a"]})
                nuevo.set("val", col)
            else:
                nuevo = etree.Element(_A + "srgbClr", nsmap={"a": NS["a"]})
                nuevo.set("val", col.upper())
            c.getparent().replace(c, nuevo)
    if kw.get("opacidad") is not None:
        v = kw["opacidad"]
        for s in body.iter(_P + "set"):
            an = s.find(f".//{_P}attrName")
            if an is not None and an.text == "style.opacity":
                s.find(f"{_P}to/{_P}strVal").set("val", f"{v:g}")
        for ae in body.iter(_P + "animEffect"):
            if ae.get("filter") == "image":
                ae.set("prLst", f"opacity: {v:g}")
    if kw.get("fuente") is not None:
        for s in body.iter(_P + "set"):
            an = s.find(f".//{_P}attrName")
            if an is not None and an.text == "style.fontFamily":
                s.find(f"{_P}to/{_P}strVal").set("val", kw["fuente"])
    if kw.get("radios") is not None:
        n = int(kw["radios"])
        cTn.set("presetSubtype", str(n))
        for ae in body.iter(_P + "animEffect"):
            if (ae.get("filter") or "").startswith("wheel("):
                ae.set("filter", f"wheel({n})")
    if kw.get("ruta") is not None:
        for m in body.iter(_P + "animMotion"):
            m.set("path", kw["ruta"])


def efecto(clave, spid, ids=None, nodo="clickEffect", retraso=0, dur=None, grpId=0,
           repetir=None, autorev=False, rep_dur=None, accel=None, decel=None,
           suave_inicio=False, suave_fin=False, rebobinar=False, unidad="preset",
           retraso_unidad=10, n_unidades=None, parrafo=None, fondo=True, **kw):
    """Genera el ``<p:par>`` de un efecto a partir de un preset capturado de PowerPoint.

    clave        'entrada.fade', 'salida.fly_out_left', 'enfasis.spin', 'trayectoria.path_circle'...
    spid         id de la forma (python-pptx: shape.shape_id)
    nodo         clickEffect | withEffect | afterEffect
    retraso      ms, retraso propio del efecto (Timing.TriggerDelayTime)
    dur          ms; None = por defecto del preset; en presets 'indefinidos' (Transparency,
                 Change Font...) None deja dur="indefinite" (hasta el final de la diapositiva)
    grpId        entero o None (None = sin grpId ni bldP: imagenes, Fill/Line Color).
                 timing() lo renumera por forma.
    repetir      None | n | "indefinite" | "untilNextClick"
    autorev      True -> autoRev="1" (ida y vuelta)
    rep_dur      ms -> repeatDur (repetir durante n ms)
    accel/decel  fraccion 0..1 (Inicio/Fin suaves = 0.5)
    rebobinar    True -> fill="remove" (Rebobinar al terminar)
    unidad       "preset" (lo del preset) | "wd" (por palabra) | "lt" (por letra) | None (objeto)
    retraso_unidad  % de la duracion entre palabras/letras (PowerPoint: 10)
    n_unidades   palabras/letras (sin espacios) para encadenar bien un 'despues de anterior'
    parrafo      indice de parrafo (0..) -> <p:txEl><p:pRg st=k end=k/> (compilacion por parrafo)
    fondo        la forma tiene relleno/linea -> animBg="1" en su bldP
    kw           angulo (grados), escala (float o (x, y)), color ('RRGGBB'), opacidad (0..1),
                 fuente, radios (Wheel), ruta (trayectorias)
    """
    if nodo not in NODOS:
        raise ValueError(f"nodo invalido: {nodo}")
    pre = PRESETS[clave]
    ids = ids or _IDS_GLOBAL
    d0 = pre["dur"]
    if dur == "indefinite":
        dur = None if d0 is None else dur
    d = d0 if dur is None else dur
    body = elemento(f"<p:childTnLst>{pre['cuerpo'].replace('{SPID}', str(int(spid)))}</p:childTnLst>")
    ctn = etree.Element(_P + "cTn", nsmap=NS)
    ctn.set("id", str(ids()))
    ctn.set("presetID", str(pre["presetID"]))
    ctn.set("presetClass", pre["presetClass"])
    ctn.set("presetSubtype", str(pre["presetSubtype"]))
    _ajustes(body, ctn, clave, kw)
    _escalar(body, d0, None if d == "indefinite" else d)
    if repetir is not None:
        if repetir in ("indefinite", "untilNextClick"):
            ctn.set("repeatCount", "indefinite")
        else:
            ctn.set("repeatCount", str(int(round(float(repetir) * 1000))))
    elif rep_dur is not None:
        ctn.set("repeatDur", str(int(rep_dur)))
    a = accel if accel is not None else (0.5 if suave_inicio else None)
    de = decel if decel is not None else (0.5 if suave_fin else None)
    a = pre.get("accel") if a is None else int(round(a * 100000))
    de = pre.get("decel") if de is None else int(round(de * 100000))
    if a:
        ctn.set("accel", str(a))
    if de:
        ctn.set("decel", str(de))
    if autorev:
        ctn.set("autoRev", "1")
    fill = "remove" if rebobinar else pre.get("fill")
    if fill:
        ctn.set("fill", fill)
    usa_grp = pre.get("grp", True) and grpId is not None
    if usa_grp:
        ctn.set("grpId", str(int(grpId)))
    ctn.set("nodeType", nodo)
    st = etree.SubElement(ctn, _P + "stCondLst")
    etree.SubElement(st, _P + "cond").set("delay", str(int(retraso)))
    if repetir == "untilNextClick":
        ec = etree.SubElement(ctn, _P + "endCondLst")
        c = etree.SubElement(ec, _P + "cond")
        c.set("evt", "onNext")
        c.set("delay", "0")
        etree.SubElement(etree.SubElement(c, _P + "tgtEl"), _P + "sldTgt")
    it = pre.get("iterate") if unidad == "preset" else (
        (unidad, "tmPct", int(round(retraso_unidad * 1000))) if unidad else None)
    if it:
        ite = etree.SubElement(ctn, _P + "iterate")
        if it[0] != "el":
            ite.set("type", it[0])
        etree.SubElement(ite, _P + it[1]).set("val", str(it[2]))
    if parrafo is not None:
        for sp in body.iter(_P + "spTgt"):
            for x in list(sp):
                sp.remove(x)
            tx = etree.SubElement(sp, _P + "txEl")
            pr = etree.SubElement(tx, _P + "pRg")
            pr.set("st", str(parrafo))
            pr.set("end", str(parrafo))
    ctn.append(body)
    for c in body.iter(_P + "cTn"):  # ids de los comportamientos, en orden de documento
        attrs = dict(c.attrib)
        c.attrib.clear()
        c.set("id", str(ids()))
        for k, v in attrs.items():
            c.set(k, v)
    par = etree.Element(_P + "par", nsmap=NS)
    par.append(ctn)
    xml = _ser_limpio(par)
    e = Efecto(xml)
    e.clave = clave
    e.spid = int(spid)
    e.nodo = nodo
    e.clase = pre["presetClass"]
    e.dur = None if d in (None, "indefinite") else int(d)
    e.retraso = int(retraso)
    e.grp = usa_grp
    e.parrafo = parrafo
    e.fondo = fondo
    total = 0 if e.dur is None else e.dur
    if it and n_unidades:  # comprobado: 6 palabras, 500 ms, 10 % -> 750 ms
        if it[1] == "tmPct":
            total = _mitad_arriba(total * (1 + Fraction(it[2], 100000) * (int(n_unidades) - 1)))
        else:
            total += it[2] * (int(n_unidades) - 1)
    if repetir not in (None, "indefinite", "untilNextClick"):
        total = _mitad_arriba(total * Fraction(str(repetir)))
    if autorev:
        total *= 2
    e.dur_total = total
    e.repetir = repetir
    e.autorev = autorev
    e.com = pre["com"]
    return e


def entrada(tipo, spid, direccion=None, **kw):
    """Efecto de entrada. tipo: 'fade', 'fly_bottom' (o 'fly' + direccion='left'), 'zoom',
    'wipe_left', 'split_vertical_out', 'grow_turn', 'rise_up', 'bounce', ... (ver claves('entrada'))."""
    return efecto(_resolver("entrada", tipo, direccion), spid, **kw)


def salida(tipo, spid, direccion=None, **kw):
    """Efecto de salida. tipo: 'fade', 'fly_out_bottom', 'zoom', 'wipe_left', 'shrink_turn'..."""
    return efecto(_resolver("salida", tipo, direccion), spid, **kw)


def enfasis(tipo, spid, direccion=None, **kw):
    """Efecto de enfasis. tipo: 'pulse', 'color_pulse', 'teeter', 'spin' (angulo=), 'grow_shrink'
    (escala=), 'transparency' (opacidad=), 'fill_color'/'font_color'/'line_color'/'object_color'
    (color='RRGGBB'), 'wave', 'bold_flash', 'brush_color', ..."""
    alias = {"spin": "spin_360", "transparency": "transparency_50"}
    return efecto(_resolver("enfasis", alias.get(tipo, tipo), direccion), spid, **kw)


def trayectoria(ruta=None, spid=None, preset=None, dur=None, preset_id=None, rAng=None,
                ptsTypes=None, centro=None, **kw):
    """Trayectoria de movimiento (bucleable con repetir=/autorev=).

    preset   'path_right', 'path_circle', 'path_arc_down', 'path_heart'... (claves('trayectoria'))
    ruta     cadena de trazado relativa al tamano de la diapositiva ('M 0 0 L 0.25 0 E');
             con preset=None se usa el preset 'path_right' con presetID=0 (trayectoria
             personalizada) salvo que se indique preset_id.
    rAng, ptsTypes, centro   atributos que escribe la interfaz al dibujar a mano (opcionales):
             centro=(x, y) en fraccion de diapositiva -> <p:rCtr x= y=/> (x100000);
             centro="auto" lo calcula como el centro de la caja de la ruta.
    """
    clave = f"trayectoria.{preset or 'path_right'}"
    if ruta is not None:
        kw["ruta"] = ruta
    e = efecto(clave, spid, dur=dur, **kw)
    if preset is None and ruta is not None or preset_id is not None or rAng is not None \
            or ptsTypes is not None or centro is not None:
        par = elemento(e)
        ctn = par.find(_P + "cTn")
        if preset is None and preset_id is None:
            preset_id = 0
        if preset_id is not None:
            ctn.set("presetID", str(preset_id))
        m = par.find(f".//{_P}animMotion")
        if rAng is not None:
            _reordenar_attr(m, "rAng", str(int(round(rAng * 60000))), antes_de="ptsTypes")
        if ptsTypes is not None:
            m.set("ptsTypes", ptsTypes)
        if centro is not None:
            if centro == "auto":
                centro = _centro_ruta(m.get("path"))
            rc = etree.SubElement(m, _P + "rCtr")
            rc.set("x", str(int(round(centro[0] * 100000))))
            rc.set("y", str(int(round(centro[1] * 100000))))
        nuevo = Efecto(_ser_limpio(par))
        nuevo.__dict__.update(e.__dict__)
        if ctn.get("presetID") == "0":
            # trayectoria personalizada: COM la reporta como msoAnimEffectCustom (0)
            nuevo.com = [0, 0, None]
            nuevo.clave = "trayectoria.personalizada"
        e = nuevo
    return e


def _reordenar_attr(el, nombre, valor, antes_de):
    attrs = list(el.attrib.items())
    el.attrib.clear()
    puesto = False
    for k, v in attrs:
        if k == antes_de and not puesto:
            el.set(nombre, valor)
            puesto = True
        el.set(k, v)
    if not puesto:
        el.set(nombre, valor)


def _centro_ruta(path):
    nums = [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?(?:[eE]-?\d+)?", path)]
    xs, ys = nums[0::2], nums[1::2]
    return ((min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2)


def texto_por_palabra(tipo, spid, familia="entrada", unidad="wd", retraso_pct=10, **kw):
    """Anima el texto de una forma palabra a palabra (unidad='wd') o letra a letra ('lt'):
    PowerPoint escribe <p:iterate type="wd|lt"><p:tmPct val="10000"/></p:iterate> en el cTn
    del efecto; el destino sigue siendo la forma entera (sin txEl) y el bldP no lleva build."""
    fn = {"entrada": entrada, "salida": salida, "enfasis": enfasis}[familia]
    return fn(tipo, spid, unidad=unidad, retraso_unidad=retraso_pct, **kw)


def esperado_com(e, interactiva=False):
    """Valores que PowerPoint debe reportar por COM para el efecto generado."""
    tipo, salida_, direc = e.com
    d = {"clave": e.clave, "EffectType": tipo, "Exit": salida_,
         "TriggerType": 4 if interactiva else {"clickEffect": 1, "withEffect": 2, "afterEffect": 3}[e.nodo],
         "TriggerDelayTime": round(e.retraso / 1000, 4),
         "AutoReverse": -1 if e.autorev else 0}
    d["Duration"] = "inf" if e.dur is None else round(e.dur / 1000, 4)
    if e.repetir in ("indefinite", "untilNextClick"):
        d["RepeatCount"] = INDEFINIDO_COM
    elif e.repetir is not None:
        d["RepeatCount"] = int(float(e.repetir))
    else:
        d["RepeatCount"] = 1
    if direc is not None:
        d["Direction"] = direc
    return d


# ======================================================================================
# <p:timing>
# ======================================================================================
def _E(tag, parent=None, **attrs):
    el = etree.Element(_P + tag) if parent is None else etree.SubElement(parent, _P + tag)
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _meta(e):
    """Metadatos de un efecto (Efecto o cadena XML cualquiera)."""
    el = elemento(e)
    ctn = el.find(_P + "cTn")
    m = {"el": el, "nodo": ctn.get("nodeType"), "clase": ctn.get("presetClass"),
         "grp": ctn.get("grpId") is not None}
    sp = el.find(f".//{_P}spTgt")
    m["spid"] = int(sp.get("spid")) if sp is not None else None
    m["parrafo"] = el.find(f".//{_P}txEl/{_P}pRg") is not None
    cond = ctn.find(f"{_P}stCondLst/{_P}cond")
    m["retraso"] = int(cond.get("delay")) if cond is not None and cond.get("delay", "").isdigit() else 0
    if isinstance(e, Efecto):
        m["dur_total"] = e.dur_total
        m["fondo"] = e.fondo
    else:  # deducirlo del XML
        fin = 0
        for c in ctn.find(_P + "childTnLst").iter(_P + "cTn"):
            du = c.get("dur")
            if du and du.isdigit():
                cd = c.find(f"{_P}stCondLst/{_P}cond")
                de = int(cd.get("delay")) if cd is not None and cd.get("delay", "").isdigit() else 0
                fin = max(fin, de + int(du))
        rc = ctn.get("repeatCount")
        if rc and rc.isdigit():
            fin = _mitad_arriba(Fraction(fin * int(rc), 1000))
        if ctn.get("autoRev") == "1":
            fin *= 2
        m["dur_total"] = fin
        m["fondo"] = True
    return m


def _grupos(parent, grupos, interactiva):
    """Anida grupos de clic -> sub-par (cadena con/despues) -> efectos."""
    for gi, grupo in enumerate(grupos):
        if not grupo:
            continue
        metas = [_meta(e) for e in grupo]
        par = _E("par", parent)
        ctn = _E("cTn", par, id=0, fill="hold")
        st = _E("stCondLst", ctn)
        if interactiva:
            _E("cond", st, delay="0" if gi == 0 else "indefinite")
        else:
            _E("cond", st, delay="indefinite")
            if gi == 0 and metas[0]["nodo"] != "clickEffect":
                c = _E("cond", st, evt="onBegin", delay="0")
                _E("tn", c, val="2")
        hijos = _E("childTnLst", ctn)
        t = 0
        fin_sub = 0
        sub = None
        for k, m in enumerate(metas):
            if sub is None or m["nodo"] == "afterEffect" and k > 0:
                if sub is not None:
                    t = t + fin_sub
                spar = _E("par", hijos)
                sctn = _E("cTn", spar, id=0, fill="hold")
                _E("cond", _E("stCondLst", sctn), delay=str(t))
                sub = _E("childTnLst", sctn)
                fin_sub = 0
            sub.append(m["el"])
            fin_sub = max(fin_sub, m["retraso"] + m["dur_total"])
    return parent


def timing(principal=(), interactivas=(), sin_fondo=(), autogrp=True):
    """Construye ``<p:timing>``.

    principal     lista de grupos de clic; cada grupo es una lista de efectos. El primer
                  efecto de cada grupo es el que dispara el clic (clickEffect); si el primer
                  efecto del PRIMER grupo es withEffect/afterEffect, el grupo arranca solo al
                  entrar en la diapositiva (onBegin del mainSeq).
    interactivas  lista de (spid_disparador, grupos): secuencias disparadas al hacer clic en
                  esa forma. Cada grupo = un clic sobre el disparador.
    sin_fondo     spids sin relleno ni linea (cuadros de texto): su bldP no lleva animBg.
    autogrp       renumera grpId por forma en orden de linea de tiempo (como PowerPoint).
    """
    raiz = etree.Element(_P + "timing", nsmap={"p": NS["p"], "a": NS["a"], "r": NS["r"]})
    tn = _E("tnLst", raiz)
    par = _E("par", tn)
    root = _E("cTn", par, id=0, dur="indefinite", restart="never", nodeType="tmRoot")
    hijos = _E("childTnLst", root)
    grupos_p = [g for g in principal if g]
    if grupos_p:
        seq = _E("seq", hijos, concurrent="1", nextAc="seek")
        ms = _E("cTn", seq, id=0, dur="indefinite", nodeType="mainSeq")
        _grupos(_E("childTnLst", ms), grupos_p, False)
        for tag, evt in (("prevCondLst", "onPrev"), ("nextCondLst", "onNext")):
            c = _E("cond", _E(tag, seq), evt=evt, delay="0")
            _E("sldTgt", _E("tgtEl", c))
    for spid_disp, grupos in interactivas:
        if grupos and isinstance(grupos[0], str):
            grupos = [grupos]
        seq = _E("seq", hijos, concurrent="1", nextAc="seek")
        ic = _E("cTn", seq, id=0, restart="whenNotActive", fill="hold",
                evtFilter="cancelBubble", nodeType="interactiveSeq")
        c = _E("cond", _E("stCondLst", ic), evt="onClick", delay="0")
        _E("spTgt", _E("tgtEl", c), spid=str(spid_disp))
        es = _E("endSync", ic, evt="end", delay="0")
        _E("rtn", es, val="all")
        _grupos(_E("childTnLst", ic), grupos, True)
        c = _E("cond", _E("nextCondLst", seq), evt="onClick", delay="0")
        _E("spTgt", _E("tgtEl", c), spid=str(spid_disp))
    # ids en orden de documento; <p:tn val> apunta al mainSeq
    main_id = None
    for n, c in enumerate(raiz.iter(_P + "cTn"), start=1):
        attrs = [(k, v) for k, v in c.attrib.items() if k != "id"]
        c.attrib.clear()
        c.set("id", str(n))
        for k, v in attrs:
            c.set(k, v)
        if c.get("nodeType") == "mainSeq":
            main_id = n
    for t in raiz.iter(_P + "tn"):
        t.set("val", str(main_id))
    # grpId + bldLst
    fondo = {}
    for e in list(principal_flat(principal)) + [x for _, g in interactivas for x in principal_flat(g)]:
        if isinstance(e, Efecto):
            fondo.setdefault(e.spid, e.fondo)
    bld = {}
    ultimo = {}
    cuenta = {}
    for ctn in raiz.iter(_P + "cTn"):
        if ctn.get("nodeType") not in NODOS or ctn.get("grpId") is None:
            continue
        sp = ctn.find(f".//{_P}spTgt")
        spid = int(sp.get("spid"))
        es_par = sp.find(f"{_P}txEl/{_P}pRg") is not None
        firma = (ctn.get("presetID"), ctn.get("presetClass"))
        if autogrp:
            prev = ultimo.get(spid)
            if es_par and prev and prev[1] and prev[2] == firma:
                g = prev[0]
            else:
                g = cuenta.get(spid, -1) + 1
                cuenta[spid] = g
            ctn.set("grpId", str(g))
            ultimo[spid] = (g, es_par, firma)
        g = int(ctn.get("grpId"))
        bld.setdefault((spid, g), es_par)
    if bld:
        bl = _E("bldLst", raiz)
        for (spid, g), es_par in sorted(bld.items()):
            b = _E("bldP", bl, spid=str(spid), grpId=str(g))
            if es_par:
                b.set("build", "p")
            if spid not in sin_fondo and fondo.get(spid, True) and not es_par:
                b.set("animBg", "1")
    return etree.tostring(raiz, encoding="unicode")


def principal_flat(grupos):
    for g in grupos:
        if isinstance(g, str):
            yield g
        else:
            for e in g:
                yield e


# ======================================================================================
# transiciones
# ======================================================================================
def _spd(dur):
    if dur <= 500:
        return None  # "fast" es el valor por defecto y PowerPoint no lo escribe
    return "med" if dur <= 750 else "slow"


def transicion(tipo, dur=1000, avance_ms=None, avance_clic=True):
    """``<p:transition>`` o ``<mc:AlternateContent>`` identico al de PowerPoint.

    tipo        clave de TRANSICIONES ('fade', 'fade_through_black', 'push_up', 'vortex_left',
                'morph_objeto', 'curtains', 'airplane_left', ...)
    dur         ms (SlideShowTransition.Duration). Con transiciones del esquema base y dur de
                500/750/1000 PowerPoint escribe <p:transition> simple (sin p14:dur).
    avance_ms   avance automatico (advTm); avance_clic=False -> advClick="0".
    """
    t = TRANSICIONES[tipo]
    dur = int(dur)
    spd = _spd(dur)
    extra = ""
    if not avance_clic:
        extra += ' advClick="0"'
    if avance_ms is not None:
        extra += f' advTm="{int(avance_ms)}"'
    s_spd = f' spd="{spd}"' if spd else ""
    el = t["el"]
    ns = t["ns"]
    if ns == "p" and dur in (500, 750, 1000):
        return f"<p:transition{s_spd}{extra}>{el}</p:transition>"
    fb = t["fb"] or el
    if ns in ("p", "p14"):
        choice = (f'<mc:Choice Requires="p14"><p:transition{s_spd} p14:dur="{dur}"{extra}>{el}'
                  f'</p:transition></mc:Choice>')
        decl = f'xmlns:mc="{NS["mc"]}" xmlns:p14="{NS["p14"]}"'
    else:
        choice = (f'<mc:Choice Requires="{ns}"><p:transition{s_spd} xmlns:p14="{NS["p14"]}" '
                  f'p14:dur="{dur}"{extra}>{el}</p:transition></mc:Choice>')
        decl = f'xmlns:mc="{NS["mc"]}" xmlns:{ns}="{NS[ns]}"'
    return (f"<mc:AlternateContent {decl}>{choice}<mc:Fallback xmlns=\"\">"
            f"<p:transition{s_spd}{extra}>{fb}</p:transition></mc:Fallback></mc:AlternateContent>")


# ======================================================================================
# insercion en <p:sld>
# ======================================================================================
def _es(el, local):
    if el.tag == _P + local:
        return True
    return el.tag == _MC + "AlternateContent" and el.find(f".//{_P}{local}") is not None


def inyectar(sld, transicion=None, timing=None):
    """Inserta (reemplazando) transicion y/o timing en un ``<p:sld>`` de lxml
    (python-pptx: ``slide._element``) respetando el orden del esquema:
    cSld, clrMapOvr, transition/AlternateContent, timing, extLst."""
    for local, xml in (("transition", transicion), ("timing", timing)):
        if xml is None:
            continue
        for ch in list(sld):
            if _es(ch, local):
                sld.remove(ch)
    nuevos = []
    if transicion is not None:
        nuevos.append(elemento(transicion))
    if timing is not None:
        nuevos.append(etree.fromstring(timing.encode("utf-8")) if timing.lstrip().startswith("<p:timing xmlns")
                      else elemento(timing))
    for el in nuevos:
        es_trans = _es(el, "transition")
        ancla = None
        for ch in sld:
            loc = etree.QName(ch).localname
            if loc in ("cSld", "clrMapOvr") or (not es_trans and _es(ch, "transition")):
                ancla = ch
        if ancla is None:
            sld.insert(0, el)
        else:
            ancla.addnext(el)
    return sld


# ======================================================================================
# === DATOS GENERADOS (s04_datos.py a partir de referencia.pptx; no editar a mano) ===
# PRESETS[clave] = {ui, presetID, presetClass, presetSubtype, accel, decel, fill, grp,
#   iterate (tipo, tmPct), dur (ms por defecto; None = indefinida), com [EffectType, Exit,
#   Direction] que reporta COM, cuerpo (hijos de <p:childTnLst>, spid='{SPID}', sin ids)}
PRESETS = {'entrada.fade': {'ui': 'Fade',
                  'presetID': 10,
                  'presetClass': 'entr',
                  'presetSubtype': 0,
                  'accel': None,
                  'decel': None,
                  'fill': 'hold',
                  'grp': True,
                  'iterate': None,
                  'dur': 500,
                  'com': [10, 0, None],
                  'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                            'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                            'filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.appear': {'ui': 'Appear',
                    'presetID': 1,
                    'presetClass': 'entr',
                    'presetSubtype': 0,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 1,
                    'com': [1, 0, None],
                    'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="visible"/></p:to></p:set>'},
 'entrada.fly_bottom': {'ui': 'Fly In desde abajo',
                        'presetID': 2,
                        'presetClass': 'entr',
                        'presetSubtype': 4,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 500,
                        'com': [2, 0, 3],
                        'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                  'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="1+#ppt_h/2"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.fly_left': {'ui': 'Fly In desde la izquierda',
                      'presetID': 2,
                      'presetClass': 'entr',
                      'presetSubtype': 8,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 500,
                      'com': [2, 0, 4],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                'fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="0-#ppt_w/2"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                'fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.fly_right': {'ui': 'Fly In desde la derecha',
                       'presetID': 2,
                       'presetClass': 'entr',
                       'presetSubtype': 2,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 500,
                       'com': [2, 0, 2],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                 'fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="1+#ppt_w/2"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                 'fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.fly_top': {'ui': 'Fly In desde arriba',
                     'presetID': 2,
                     'presetClass': 'entr',
                     'presetSubtype': 1,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 500,
                     'com': [2, 0, 1],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                               'fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                               'fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="0-#ppt_h/2"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.fly_bottom_left': {'ui': 'Fly In desde abajo-izquierda',
                             'presetID': 2,
                             'presetClass': 'entr',
                             'presetSubtype': 12,
                             'accel': None,
                             'decel': None,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 500,
                             'com': [2, 0, 9],
                             'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                       'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                       'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                       'valueType="num"><p:cBhvr additive="base"><p:cTn dur="500" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                       'tm="0"><p:val><p:strVal val="0-#ppt_w/2"/></p:val></p:tav><p:tav '
                                       'tm="100000"><p:val><p:strVal '
                                       'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                       'calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                       'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                       'tm="0"><p:val><p:strVal val="1+#ppt_h/2"/></p:val></p:tav><p:tav '
                                       'tm="100000"><p:val><p:strVal '
                                       'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.wipe_left': {'ui': 'Wipe desde la izquierda',
                       'presetID': 22,
                       'presetClass': 'entr',
                       'presetSubtype': 8,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 500,
                       'com': [22, 0, 4],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                 'filter="wipe(left)"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.wipe_bottom': {'ui': 'Wipe desde abajo',
                         'presetID': 22,
                         'presetClass': 'entr',
                         'presetSubtype': 4,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [22, 0, 3],
                         'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                   'filter="wipe(down)"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.wipe_right': {'ui': 'Wipe desde la derecha',
                        'presetID': 22,
                        'presetClass': 'entr',
                        'presetSubtype': 2,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 500,
                        'com': [22, 0, 2],
                        'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                  'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                  'filter="wipe(right)"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.wipe_top': {'ui': 'Wipe desde arriba',
                      'presetID': 22,
                      'presetClass': 'entr',
                      'presetSubtype': 1,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 500,
                      'com': [22, 0, 1],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                'filter="wipe(up)"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.float_in_up': {'ui': 'Float In (hacia arriba) = Ascend',
                         'presetID': 42,
                         'presetClass': 'entr',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 1000,
                         'com': [39, 0, None],
                         'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                   'filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                                   'valueType="num"><p:cBhvr><p:cTn dur="1000" '
                                   'fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                   'valueType="num"><p:cBhvr><p:cTn dur="1000" '
                                   'fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal val="#ppt_y+.1"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.float_in_down': {'ui': 'Float In (hacia abajo) = Descend',
                           'presetID': 47,
                           'presetClass': 'entr',
                           'presetSubtype': 0,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 1000,
                           'com': [42, 0, None],
                           'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                     'filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="1000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="1000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="#ppt_y-.1"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.zoom': {'ui': 'Zoom (moderno) = Faded Zoom',
                  'presetID': 53,
                  'presetClass': 'entr',
                  'presetSubtype': 16,
                  'accel': None,
                  'decel': None,
                  'fill': 'hold',
                  'grp': True,
                  'iterate': None,
                  'dur': 500,
                  'com': [48, 0, 19],
                  'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                            'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="in" '
                            'filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.split_vertical_out': {'ui': 'Split vertical hacia fuera',
                                'presetID': 16,
                                'presetClass': 'entr',
                                'presetSubtype': 37,
                                'accel': None,
                                'decel': None,
                                'fill': 'hold',
                                'grp': True,
                                'iterate': None,
                                'dur': 500,
                                'com': [16, 0, 26],
                                'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                          'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                          'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                          'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                          'filter="barn(outVertical)"><p:cBhvr><p:cTn '
                                          'dur="500"/><p:tgtEl><p:spTgt '
                                          'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.split_vertical_in': {'ui': 'Split vertical hacia dentro',
                               'presetID': 16,
                               'presetClass': 'entr',
                               'presetSubtype': 21,
                               'accel': None,
                               'decel': None,
                               'fill': 'hold',
                               'grp': True,
                               'iterate': None,
                               'dur': 500,
                               'com': [16, 0, 25],
                               'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                         'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                         'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                         'filter="barn(inVertical)"><p:cBhvr><p:cTn '
                                         'dur="500"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.split_horizontal_out': {'ui': 'Split horizontal hacia fuera',
                                  'presetID': 16,
                                  'presetClass': 'entr',
                                  'presetSubtype': 42,
                                  'accel': None,
                                  'decel': None,
                                  'fill': 'hold',
                                  'grp': True,
                                  'iterate': None,
                                  'dur': 500,
                                  'com': [16, 0, 24],
                                  'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                            'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                            'filter="barn(outHorizontal)"><p:cBhvr><p:cTn '
                                            'dur="500"/><p:tgtEl><p:spTgt '
                                            'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.grow_turn': {'ui': 'Grow & Turn',
                       'presetID': 31,
                       'presetClass': 'entr',
                       'presetSubtype': 0,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 1000,
                       'com': [31, 0, None],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:fltVal val="90"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:fltVal '
                                 'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="in" '
                                 'filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.rise_up': {'ui': 'Rise Up',
                     'presetID': 37,
                     'presetClass': 'entr',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 1000,
                     'com': [34, 0, None],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                               'filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="900" decel="100000" '
                               'fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_y+1"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_y-.03"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="100" accel="100000" '
                               'fill="hold"><p:stCondLst><p:cond '
                               'delay="900"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_y-.03"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.peek_bottom': {'ui': 'Peek In desde abajo',
                         'presetID': 12,
                         'presetClass': 'entr',
                         'presetSubtype': 4,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [12, 0, 3],
                         'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                   'valueType="num"><p:cBhvr additive="base"><p:cTn '
                                   'dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal '
                                   'val="#ppt_y+#ppt_h*1.125000"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                   'transition="in" filter="wipe(up)"><p:cBhvr><p:cTn '
                                   'dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.peek_left': {'ui': 'Peek In desde la izquierda',
                       'presetID': 12,
                       'presetClass': 'entr',
                       'presetSubtype': 8,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 500,
                       'com': [12, 0, 4],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr additive="base"><p:cTn '
                                 'dur="500"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal '
                                 'val="#ppt_x-#ppt_w*1.125000"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                 'transition="in" filter="wipe(right)"><p:cBhvr><p:cTn '
                                 'dur="500"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.stretch_across': {'ui': 'Stretch (across)',
                            'presetID': 17,
                            'presetClass': 'entr',
                            'presetSubtype': 10,
                            'accel': None,
                            'decel': None,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 500,
                            'com': [17, 0, 16],
                            'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                      'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                      'valueType="num"><p:cBhvr><p:cTn dur="500" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                      'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="500" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.expand': {'ui': 'Expand',
                    'presetID': 55,
                    'presetClass': 'entr',
                    'presetSubtype': 0,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 1000,
                    'com': [50, 0, None],
                    'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="#ppt_w*0.70"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:strVal '
                              'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:strVal '
                              'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                              'transition="in" filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.compress': {'ui': 'Compress (Stretchy)',
                      'presetID': 50,
                      'presetClass': 'entr',
                      'presetSubtype': 0,
                      'accel': None,
                      'decel': 100000,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 1000,
                      'com': [45, 0, None],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="#ppt_w+.3"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                'transition="in" filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.swivel': {'ui': 'Swivel',
                    'presetID': 19,
                    'presetClass': 'entr',
                    'presetSubtype': 10,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 5000,
                    'com': [19, 0, 16],
                    'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="5000" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0" fmla="#ppt_w*sin(2.5*pi*$)"><p:val><p:fltVal '
                              'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                              'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="5000" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:strVal '
                              'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.bounce': {'ui': 'Bounce',
                    'presetID': 26,
                    'presetClass': 'entr',
                    'presetSubtype': 0,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 2000,
                    'com': [26, 0, None],
                    'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                              'filter="wipe(down)"><p:cBhvr><p:cTn dur="580"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="1822" tmFilter="0,0; 0.14,0.36; '
                              '0.43,0.73; 0.71,0.91; 1.0,1.0"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="#ppt_x-0.25"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:strVal '
                              'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="664" tmFilter="0.0,0.0; 0.25,0.07; '
                              '0.50,0.2; 0.75,0.467; 1.0,1.0"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0" fmla="#ppt_y-sin(pi*$)/3"><p:val><p:fltVal '
                              'val="0.5"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                              'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="664" tmFilter="0, 0; 0.125,0.2665; '
                              '0.25,0.4; 0.375,0.465; 0.5,0.5;  0.625,0.535; 0.75,0.6; 0.875,0.7335; '
                              '1,1"><p:stCondLst><p:cond '
                              'delay="664"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0" fmla="#ppt_y-sin(pi*$)/9"><p:val><p:fltVal '
                              'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                              'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="332" tmFilter="0, 0; 0.125,0.2665; '
                              '0.25,0.4; 0.375,0.465; 0.5,0.5;  0.625,0.535; 0.75,0.6; 0.875,0.7335; '
                              '1,1"><p:stCondLst><p:cond '
                              'delay="1324"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0" fmla="#ppt_y-sin(pi*$)/27"><p:val><p:fltVal '
                              'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                              'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="164" tmFilter="0, 0; 0.125,0.2665; '
                              '0.25,0.4; 0.375,0.465; 0.5,0.5;  0.625,0.535; 0.75,0.6; 0.875,0.7335; '
                              '1,1"><p:stCondLst><p:cond '
                              'delay="1656"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0" fmla="#ppt_y-sin(pi*$)/81"><p:val><p:fltVal '
                              'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                              'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:animScale><p:cBhvr><p:cTn '
                              'dur="26"><p:stCondLst><p:cond '
                              'delay="650"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="60000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                              'decel="50000"><p:stCondLst><p:cond '
                              'delay="676"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="100000"/></p:animScale><p:animScale><p:cBhvr><p:cTn '
                              'dur="26"><p:stCondLst><p:cond '
                              'delay="1312"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="80000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                              'decel="50000"><p:stCondLst><p:cond '
                              'delay="1338"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="100000"/></p:animScale><p:animScale><p:cBhvr><p:cTn '
                              'dur="26"><p:stCondLst><p:cond '
                              'delay="1642"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="90000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                              'decel="50000"><p:stCondLst><p:cond '
                              'delay="1668"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="100000"/></p:animScale><p:animScale><p:cBhvr><p:cTn '
                              'dur="26"><p:stCondLst><p:cond '
                              'delay="1808"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="95000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                              'decel="50000"><p:stCondLst><p:cond '
                              'delay="1834"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                              'y="100000"/></p:animScale>'},
 'entrada.random_bars_h': {'ui': 'Random Bars horizontal',
                           'presetID': 14,
                           'presetClass': 'entr',
                           'presetSubtype': 10,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 500,
                           'com': [14, 0, 16],
                           'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                     'filter="randombar(horizontal)"><p:cBhvr><p:cTn '
                                     'dur="500"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.random_bars_v': {'ui': 'Random Bars vertical',
                           'presetID': 14,
                           'presetClass': 'entr',
                           'presetSubtype': 5,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 500,
                           'com': [14, 0, 17],
                           'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                     'filter="randombar(vertical)"><p:cBhvr><p:cTn '
                                     'dur="500"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.wheel_1': {'ui': 'Wheel 1 radio',
                     'presetID': 21,
                     'presetClass': 'entr',
                     'presetSubtype': 1,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 2000,
                     'com': [21, 0, None],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                               'filter="wheel(1)"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.wheel_4': {'ui': 'Wheel 4 radios',
                     'presetID': 21,
                     'presetClass': 'entr',
                     'presetSubtype': 4,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 2000,
                     'com': [21, 0, None],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                               'filter="wheel(4)"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.shape_circle_out': {'ui': 'Shape circulo hacia fuera',
                              'presetID': 6,
                              'presetClass': 'entr',
                              'presetSubtype': 32,
                              'accel': None,
                              'decel': None,
                              'fill': 'hold',
                              'grp': True,
                              'iterate': None,
                              'dur': 2000,
                              'com': [6, 0, 20],
                              'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                        'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                        'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                        'filter="circle(out)"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.shape_circle_in': {'ui': 'Shape circulo hacia dentro',
                             'presetID': 6,
                             'presetClass': 'entr',
                             'presetSubtype': 16,
                             'accel': None,
                             'decel': None,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [6, 0, 19],
                             'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                       'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                       'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                       'filter="circle(in)"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.shape_diamond_out': {'ui': 'Shape rombo hacia fuera',
                               'presetID': 8,
                               'presetClass': 'entr',
                               'presetSubtype': 32,
                               'accel': None,
                               'decel': None,
                               'fill': 'hold',
                               'grp': True,
                               'iterate': None,
                               'dur': 2000,
                               'com': [8, 0, 20],
                               'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                         'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                         'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                         'filter="diamond(out)"><p:cBhvr><p:cTn '
                                         'dur="2000"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.shape_plus_out': {'ui': 'Shape mas hacia fuera',
                            'presetID': 13,
                            'presetClass': 'entr',
                            'presetSubtype': 32,
                            'accel': None,
                            'decel': None,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 2000,
                            'com': [13, 0, 20],
                            'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                      'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                      'filter="plus(out)"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.shape_box_out': {'ui': 'Shape cuadro hacia fuera',
                           'presetID': 4,
                           'presetClass': 'entr',
                           'presetSubtype': 32,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 2000,
                           'com': [4, 0, 20],
                           'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                     'filter="box(out)"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.basic_zoom': {'ui': 'Basic Zoom',
                        'presetID': 23,
                        'presetClass': 'entr',
                        'presetSubtype': 16,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 500,
                        'com': [23, 0, 19],
                        'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                  'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.center_revolve': {'ui': 'Center Revolve',
                            'presetID': 43,
                            'presetClass': 'entr',
                            'presetSubtype': 0,
                            'accel': None,
                            'decel': None,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 1000,
                            'com': [40, 0, None],
                            'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                      'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                      'filter="fade"><p:cBhvr><p:cTn dur="100"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim '
                                      'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="400" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                      'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="400" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:strVal val="#ppt_y+0.31"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_y+0.31"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                      'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="600" '
                                      'decel="50000" fill="hold"><p:stCondLst><p:cond '
                                      'delay="400"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                      'tm="5000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0242"/></p:val></p:tav><p:tav '
                                      'tm="10000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0479"/></p:val></p:tav><p:tav '
                                      'tm="15000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0704"/></p:val></p:tav><p:tav '
                                      'tm="20000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0911"/></p:val></p:tav><p:tav '
                                      'tm="25000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1096"/></p:val></p:tav><p:tav '
                                      'tm="30000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1254"/></p:val></p:tav><p:tav '
                                      'tm="35000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1381"/></p:val></p:tav><p:tav '
                                      'tm="40000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1474"/></p:val></p:tav><p:tav '
                                      'tm="45000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1531"/></p:val></p:tav><p:tav '
                                      'tm="50000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1550"/></p:val></p:tav><p:tav '
                                      'tm="55000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1531"/></p:val></p:tav><p:tav '
                                      'tm="60000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1474"/></p:val></p:tav><p:tav '
                                      'tm="65000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1381"/></p:val></p:tav><p:tav '
                                      'tm="70000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1254"/></p:val></p:tav><p:tav '
                                      'tm="75000"><p:val><p:strVal '
                                      'val="#ppt_x+0.1096"/></p:val></p:tav><p:tav '
                                      'tm="80000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0911"/></p:val></p:tav><p:tav '
                                      'tm="85000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0704"/></p:val></p:tav><p:tav '
                                      'tm="90000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0479"/></p:val></p:tav><p:tav '
                                      'tm="95000"><p:val><p:strVal '
                                      'val="#ppt_x+0.0242"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                      'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="600" '
                                      'decel="50000" fill="hold"><p:stCondLst><p:cond '
                                      'delay="400"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:strVal val="#ppt_y+0.31"/></p:val></p:tav><p:tav '
                                      'tm="5000"><p:val><p:strVal val="#ppt_y+0.308"/></p:val></p:tav><p:tav '
                                      'tm="10000"><p:val><p:strVal '
                                      'val="#ppt_y+0.3024"/></p:val></p:tav><p:tav '
                                      'tm="15000"><p:val><p:strVal '
                                      'val="#ppt_y+0.2931"/></p:val></p:tav><p:tav '
                                      'tm="20000"><p:val><p:strVal '
                                      'val="#ppt_y+0.2804"/></p:val></p:tav><p:tav '
                                      'tm="25000"><p:val><p:strVal '
                                      'val="#ppt_y+0.2646"/></p:val></p:tav><p:tav '
                                      'tm="30000"><p:val><p:strVal '
                                      'val="#ppt_y+0.2461"/></p:val></p:tav><p:tav '
                                      'tm="35000"><p:val><p:strVal '
                                      'val="#ppt_y+0.2253"/></p:val></p:tav><p:tav '
                                      'tm="40000"><p:val><p:strVal '
                                      'val="#ppt_y+0.2029"/></p:val></p:tav><p:tav '
                                      'tm="45000"><p:val><p:strVal '
                                      'val="#ppt_y+0.1792"/></p:val></p:tav><p:tav '
                                      'tm="50000"><p:val><p:strVal '
                                      'val="#ppt_y+0.155"/></p:val></p:tav><p:tav '
                                      'tm="55000"><p:val><p:strVal '
                                      'val="#ppt_y+0.1307"/></p:val></p:tav><p:tav '
                                      'tm="60000"><p:val><p:strVal '
                                      'val="#ppt_y+0.1071"/></p:val></p:tav><p:tav '
                                      'tm="65000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0846"/></p:val></p:tav><p:tav '
                                      'tm="70000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0639"/></p:val></p:tav><p:tav '
                                      'tm="75000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0454"/></p:val></p:tav><p:tav '
                                      'tm="80000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0296"/></p:val></p:tav><p:tav '
                                      'tm="85000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0169"/></p:val></p:tav><p:tav '
                                      'tm="90000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0076"/></p:val></p:tav><p:tav '
                                      'tm="95000"><p:val><p:strVal '
                                      'val="#ppt_y+0.0019"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.pinwheel': {'ui': 'Pinwheel',
                      'presetID': 35,
                      'presetClass': 'entr',
                      'presetSubtype': 0,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 2000,
                      'com': [33, 0, None],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                'filter="fade"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:fltVal val="720"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:fltVal '
                                'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.dissolve_in': {'ui': 'Dissolve In',
                         'presetID': 9,
                         'presetClass': 'entr',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [9, 0, None],
                         'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                   'filter="dissolve"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.blinds_h': {'ui': 'Blinds horizontal',
                      'presetID': 3,
                      'presetClass': 'entr',
                      'presetSubtype': 10,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 500,
                      'com': [3, 0, 16],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                'filter="blinds(horizontal)"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.checkerboard_across': {'ui': 'Checkerboard across',
                                 'presetID': 5,
                                 'presetClass': 'entr',
                                 'presetSubtype': 10,
                                 'accel': None,
                                 'decel': None,
                                 'fill': 'hold',
                                 'grp': True,
                                 'iterate': None,
                                 'dur': 500,
                                 'com': [5, 0, 16],
                                 'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                           'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                           'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                           'filter="checkerboard(across)"><p:cBhvr><p:cTn '
                                           'dur="500"/><p:tgtEl><p:spTgt '
                                           'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.strips_down_right': {'ui': 'Strips abajo-derecha',
                               'presetID': 18,
                               'presetClass': 'entr',
                               'presetSubtype': 6,
                               'accel': None,
                               'decel': None,
                               'fill': 'hold',
                               'grp': True,
                               'iterate': None,
                               'dur': 500,
                               'com': [18, 0, 8],
                               'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                         'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                         'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                         'filter="strips(downRight)"><p:cBhvr><p:cTn '
                                         'dur="500"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.wedge': {'ui': 'Wedge',
                   'presetID': 20,
                   'presetClass': 'entr',
                   'presetSubtype': 0,
                   'accel': None,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 2000,
                   'com': [20, 0, None],
                   'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                             'filter="wedge"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.spinner': {'ui': 'Spinner',
                     'presetID': 49,
                     'presetClass': 'entr',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': 100000,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 500,
                     'com': [44, 0, None],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:fltVal val="360"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:fltVal '
                               'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="in" '
                               'filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.curve_up': {'ui': 'Curve Up (Arc Up)',
                      'presetID': 52,
                      'presetClass': 'entr',
                      'presetSubtype': 0,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 1000,
                      'com': [47, 0, None],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:animScale><p:cBhvr><p:cTn dur="1000" '
                                'decel="50000" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="250000" y="250000"/><p:to '
                                'x="100000" y="100000"/></p:animScale><p:animMotion origin="layout" path="M '
                                '-0.46736 0.92887  C -0.37517 0.88508  -0.02552 0.75279  0.0908 0.66613  C  '
                                '0.20747 0.57948  0.21649 0.50394  0.23177 0.40825  C 0.24705 0.31256  '
                                '0.22118 0.15964   0.18264 0.09152  C 0.1441 0.02341  0.03802 0.0  0.0 0.0  '
                                '" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="1000" '
                                'decel="50000" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion><p:animEffect '
                                'transition="in" filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.flip': {'ui': 'Flip',
                  'presetID': 56,
                  'presetClass': 'entr',
                  'presetSubtype': 0,
                  'accel': None,
                  'decel': None,
                  'fill': 'hold',
                  'grp': True,
                  'iterate': ('lt', 'tmPct', 10000),
                  'dur': 1000,
                  'com': [51, 0, None],
                  'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                            'val="visible"/></p:to></p:set><p:anim by="(-#ppt_w*2)" calcmode="lin" '
                            'valueType="num"><p:cBhvr rctx="PPT"><p:cTn dur="500" autoRev="1" '
                            'fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:anim '
                            'by="(#ppt_w*0.50)" calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="500" '
                            'decel="50000" autoRev="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:anim '
                            'from="(-#ppt_h/2)" to="(#ppt_y)" calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                            'dur="1000" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:animRot '
                            'by="21600000"><p:cBhvr><p:cTn dur="1000" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'entrada.boomerang': {'ui': 'Boomerang',
                       'presetID': 25,
                       'presetClass': 'entr',
                       'presetSubtype': 0,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 1000,
                       'com': [25, 0, None],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" decel="50000" '
                                 'fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:fltVal val="-90"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:fltVal '
                                 'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" decel="50000" '
                                 'fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_w"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_w*.05"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                 'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="500" accel="50000" '
                                 'fill="hold"><p:stCondLst><p:cond '
                                 'delay="500"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_w*.05"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" decel="50000" '
                                 'fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_x+.4"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" decel="50000" '
                                 'fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_y-.2"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_y+.1"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                 'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="500" accel="50000" '
                                 'fill="hold"><p:stCondLst><p:cond '
                                 'delay="500"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_y+.1"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                 'transition="in" filter="fade"><p:cBhvr><p:cTn dur="1000" '
                                 'decel="50000"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.credits': {'ui': 'Credits',
                     'presetID': 28,
                     'presetClass': 'entr',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 15000,
                     'com': [28, 0, None],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="15000" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="15000" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_y+1"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_y-1"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.spiral_in': {'ui': 'Spiral In',
                       'presetID': 15,
                       'presetClass': 'entr',
                       'presetSubtype': 0,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 1000,
                       'com': [15, 0, None],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0" '
                                 'fmla="#ppt_x+(cos(-2*pi*(1-$))*-#ppt_x-sin(-2*pi*(1-$))*(1-#ppt_y))*(1-$)"><p:val><p:fltVal '
                                 'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                                 'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0" '
                                 'fmla="#ppt_y+(sin(-2*pi*(1-$))*-#ppt_x+cos(-2*pi*(1-$))*(1-#ppt_y))*(1-$)"><p:val><p:fltVal '
                                 'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                                 'val="1"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.whip': {'ui': 'Whip',
                  'presetID': 41,
                  'presetClass': 'entr',
                  'presetSubtype': 0,
                  'accel': None,
                  'decel': None,
                  'fill': 'hold',
                  'grp': True,
                  'iterate': ('lt', 'tmPct', 10000),
                  'dur': 500,
                  'com': [38, 0, None],
                  'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                            'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                            'tm="50000"><p:val><p:strVal val="#ppt_x+.1"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_h/10"/></p:val></p:tav><p:tav '
                            'tm="50000"><p:val><p:strVal val="#ppt_h+.01"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_w/10"/></p:val></p:tav><p:tav '
                            'tm="50000"><p:val><p:strVal val="#ppt_w+.01"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="in" '
                            'filter="fade"><p:cBhvr><p:cTn dur="500" tmFilter="0,0; .5, 1; 1, '
                            '1"/><p:tgtEl><p:spTgt spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.float': {'ui': 'Float (antiguo)',
                   'presetID': 30,
                   'presetClass': 'entr',
                   'presetSubtype': 0,
                   'accel': None,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 1000,
                   'com': [30, 0, None],
                   'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                             'filter="fade"><p:cBhvr><p:cTn dur="800" decel="100000"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="800" decel="100000" '
                             'fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:fltVal val="-90"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:fltVal '
                             'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="800" decel="100000" '
                             'fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_x+0.4"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_x-0.05"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="800" decel="100000" '
                             'fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_y-0.4"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y+0.1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="200" accel="100000" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="800"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_x-0.05"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="200" accel="100000" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="800"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_y+0.1"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.basic_swivel': {'ui': 'Faded Swivel / Basic Swivel',
                          'presetID': 45,
                          'presetClass': 'entr',
                          'presetSubtype': 0,
                          'accel': None,
                          'decel': None,
                          'fill': 'hold',
                          'grp': True,
                          'iterate': None,
                          'dur': 2000,
                          'com': [41, 0, None],
                          'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                    'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                    'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                                    'filter="fade"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim '
                                    'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="2000" '
                                    'fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                    'tm="0" fmla="#ppt_w*sin(2.5*pi*$)"><p:val><p:fltVal '
                                    'val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:fltVal '
                                    'val="1"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                    'valueType="num"><p:cBhvr><p:cTn dur="2000" '
                                    'fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                    'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                                    'tm="100000"><p:val><p:strVal '
                                    'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.fold': {'ui': 'Fold',
                  'presetID': 58,
                  'presetClass': 'entr',
                  'presetSubtype': 0,
                  'accel': 100000,
                  'decel': None,
                  'fill': 'hold',
                  'grp': True,
                  'iterate': None,
                  'dur': 500,
                  'com': [53, 0, None],
                  'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                            'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_w*2.5"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_h*0.01"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                            'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                            'tm="0"><p:val><p:strVal val="#ppt_h+1"/></p:val></p:tav><p:tav '
                            'tm="100000"><p:val><p:strVal '
                            'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="in" '
                            'filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.glide': {'ui': 'Glide',
                   'presetID': 54,
                   'presetClass': 'entr',
                   'presetSubtype': 0,
                   'accel': 100000,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 500,
                   'com': [49, 0, None],
                   'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_w*0.05"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_x-.2"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                             'transition="in" filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.sling': {'ui': 'Sling',
                   'presetID': 48,
                   'presetClass': 'entr',
                   'presetSubtype': 0,
                   'accel': 50000,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 1000,
                   'com': [43, 0, None],
                   'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:fltVal val="90"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:fltVal val="90"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:fltVal val="90"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:fltVal '
                             'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:fltVal val="-1"/></p:val></p:tav><p:tav '
                             'tm="50000"><p:val><p:fltVal val="0.95"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                             'transition="in" filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.light_speed': {'ui': 'Light Speed',
                         'presetID': 34,
                         'presetClass': 'entr',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 1000,
                         'com': [32, 0, None],
                         'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="visible"/></p:to></p:set><p:anim from="(-#ppt_w/2)" to="(#ppt_x)" '
                                   'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="600" '
                                   'fill="hold"><p:stCondLst><p:cond '
                                   'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:anim '
                                   'from="0" to="-1.0" calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                   'dur="200" decel="50000" autoRev="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="600"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>xshear</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:animScale><p:cBhvr><p:cTn '
                                   'dur="200" decel="100000" autoRev="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="600"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="100000" y="100000"/><p:to '
                                   'x="80000" y="100000"/></p:animScale><p:anim by="(#ppt_h/3+#ppt_w*0.1)" '
                                   'calcmode="lin" valueType="num"><p:cBhvr additive="sum"><p:cTn dur="200" '
                                   'decel="100000" autoRev="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="600"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr></p:anim>'},
 'entrada.thin_line': {'ui': 'Thin Line',
                       'presetID': 39,
                       'presetClass': 'entr',
                       'presetSubtype': 0,
                       'accel': 100000,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 500,
                       'com': [36, 0, None],
                       'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_h/20"/></p:val></p:tav><p:tav '
                                 'tm="50000"><p:val><p:strVal val="#ppt_h/20"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_w+.3"/></p:val></p:tav><p:tav '
                                 'tm="50000"><p:val><p:strVal val="#ppt_w+.3"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_x-.3"/></p:val></p:tav><p:tav '
                                 'tm="50000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.unfold': {'ui': 'Unfold',
                    'presetID': 40,
                    'presetClass': 'entr',
                    'presetSubtype': 0,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': ('lt', 'tmPct', 10000),
                    'dur': 1000,
                    'com': [37, 0, None],
                    'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                              'filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="#ppt_x-.1"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:strVal '
                              'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:strVal '
                              'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.zip': {'ui': 'Zip',
                 'presetID': 51,
                 'presetClass': 'entr',
                 'presetSubtype': 0,
                 'accel': None,
                 'decel': None,
                 'fill': 'hold',
                 'grp': True,
                 'iterate': None,
                 'dur': 2000,
                 'com': [46, 0, None],
                 'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                           'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                           'val="visible"/></p:to></p:set><p:animEffect transition="in" '
                           'filter="fade"><p:cBhvr><p:cTn dur="770" decel="100000"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:animScale><p:cBhvr><p:cTn '
                           'dur="770" decel="100000"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="10000" y="10000"/><p:to x="200000" '
                           'y="450000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="1230" accel="100000" '
                           'fill="hold"><p:stCondLst><p:cond '
                           'delay="770"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="200000" y="450000"/><p:to '
                           'x="100000" y="100000"/></p:animScale><p:set><p:cBhvr><p:cTn dur="770" '
                           'fill="hold"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                           'val="(0.5)"/></p:to></p:set><p:anim from="(0.5)" to="(#ppt_x)" calcmode="lin" '
                           'valueType="num"><p:cBhvr><p:cTn dur="1230" accel="100000" '
                           'fill="hold"><p:stCondLst><p:cond '
                           'delay="770"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:set><p:cBhvr><p:cTn '
                           'dur="770" fill="hold"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                           'val="(#ppt_y+0.4)"/></p:to></p:set><p:anim from="(#ppt_y+0.4)" to="(#ppt_y)" '
                           'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="1230" accel="100000" '
                           'fill="hold"><p:stCondLst><p:cond '
                           'delay="770"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:anim>'},
 'entrada.crawl_in': {'ui': 'Crawl In',
                      'presetID': 7,
                      'presetClass': 'entr',
                      'presetSubtype': 4,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 5000,
                      'com': [7, 0, 3],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr additive="base"><p:cTn dur="5000" '
                                'fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr additive="base"><p:cTn dur="5000" '
                                'fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="1+#ppt_h/2"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.swish': {'ui': 'Swish',
                   'presetID': 38,
                   'presetClass': 'entr',
                   'presetSubtype': 0,
                   'accel': 50000,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': ('lt', 'tmPct', 50000),
                   'dur': 1000,
                   'com': [35, 0, None],
                   'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="visible"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="455" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="-45.0"/></p:to></p:set><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="455" fill="hold"><p:stCondLst><p:cond '
                             'delay="455"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:fltVal val="-45"/></p:val></p:tav><p:tav '
                             'tm="69900"><p:val><p:fltVal val="45"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:fltVal '
                             'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="455" fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="#ppt_y-1"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y-(0.354*#ppt_w-0.172*#ppt_h)"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                             'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="156" decel="50000" '
                             'autoRev="1" fill="hold"><p:stCondLst><p:cond '
                             'delay="455"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal '
                             'val="#ppt_y-(0.354*#ppt_w-0.172*#ppt_h)"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y-(0.354*#ppt_w-0.172*#ppt_h)-#ppt_h/2"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                             'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="136" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="864"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal '
                             'val="#ppt_y-(0.354*#ppt_w-0.172*#ppt_h)"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'},
 'entrada.ease_in': {'ui': 'Ease In',
                     'presetID': 29,
                     'presetClass': 'entr',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 1000,
                     'com': [29, 0, None],
                     'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="visible"/></p:to></p:set><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_x-.2"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="1000" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                               'transition="in" filter="wipe(right)" prLst="gradientSize: '
                               '0.1"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'entrada.flash_once': {'ui': 'Flash Once',
                        'presetID': 11,
                        'presetClass': 'entr',
                        'presetSubtype': 0,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 1000,
                        'com': [11, 0, None],
                        'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1000"><p:stCondLst><p:cond '
                                  'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="visible"/></p:to></p:set>'},
 'salida.fade': {'ui': 'Fade (salida)',
                 'presetID': 10,
                 'presetClass': 'exit',
                 'presetSubtype': 0,
                 'accel': None,
                 'decel': None,
                 'fill': 'hold',
                 'grp': True,
                 'iterate': None,
                 'dur': 500,
                 'com': [10, -1, None],
                 'cuerpo': '<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn '
                           'dur="500"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn dur="1" '
                           'fill="hold"><p:stCondLst><p:cond '
                           'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                           'val="hidden"/></p:to></p:set>'},
 'salida.disappear': {'ui': 'Disappear',
                      'presetID': 1,
                      'presetClass': 'exit',
                      'presetSubtype': 0,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 1,
                      'com': [1, -1, None],
                      'cuerpo': '<p:set><p:cBhvr><p:cTn dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="hidden"/></p:to></p:set>'},
 'salida.fly_out_bottom': {'ui': 'Fly Out hacia abajo',
                           'presetID': 2,
                           'presetClass': 'exit',
                           'presetSubtype': 4,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 500,
                           'com': [2, -1, 3],
                           'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                     'dur="500"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                     'dur="500"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="1+ppt_h/2"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                     'dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="hidden"/></p:to></p:set>'},
 'salida.fly_out_left': {'ui': 'Fly Out hacia la izquierda',
                         'presetID': 2,
                         'presetClass': 'exit',
                         'presetSubtype': 8,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [2, -1, 4],
                         'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                   'dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="0-ppt_w/2"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                   'calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                   'dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                   'dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="hidden"/></p:to></p:set>'},
 'salida.fly_out_right': {'ui': 'Fly Out hacia la derecha',
                          'presetID': 2,
                          'presetClass': 'exit',
                          'presetSubtype': 2,
                          'accel': None,
                          'decel': None,
                          'fill': 'hold',
                          'grp': True,
                          'iterate': None,
                          'dur': 500,
                          'com': [2, -1, 2],
                          'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                    'dur="500"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                    'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                    'tm="100000"><p:val><p:strVal '
                                    'val="1+ppt_w/2"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                    'calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                    'dur="500"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                    'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                    'tm="100000"><p:val><p:strVal '
                                    'val="ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                    'dur="1" fill="hold"><p:stCondLst><p:cond '
                                    'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                    'val="hidden"/></p:to></p:set>'},
 'salida.fly_out_top': {'ui': 'Fly Out hacia arriba',
                        'presetID': 2,
                        'presetClass': 'exit',
                        'presetSubtype': 1,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 500,
                        'com': [2, -1, 1],
                        'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn '
                                  'dur="500"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr additive="base"><p:cTn '
                                  'dur="500"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="0-ppt_h/2"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                  'dur="1" fill="hold"><p:stCondLst><p:cond '
                                  'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="hidden"/></p:to></p:set>'},
 'salida.float_out_up': {'ui': 'Float Out hacia arriba (Ascend)',
                         'presetID': 42,
                         'presetClass': 'exit',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 1000,
                         'com': [39, -1, None],
                         'cuerpo': '<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn '
                                   'dur="1000"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                                   'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                   'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                   'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                   'tm="100000"><p:val><p:strVal '
                                   'val="ppt_y+.1"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                   'dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="hidden"/></p:to></p:set>'},
 'salida.float_out_down': {'ui': 'Float Out hacia abajo (Descend)',
                           'presetID': 47,
                           'presetClass': 'exit',
                           'presetSubtype': 0,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 1000,
                           'com': [42, -1, None],
                           'cuerpo': '<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn '
                                     'dur="1000"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                     'dur="1000"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                     'dur="1000"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_y-.1"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                     'dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="hidden"/></p:to></p:set>'},
 'salida.zoom': {'ui': 'Zoom (salida, faded zoom)',
                 'presetID': 53,
                 'presetClass': 'exit',
                 'presetSubtype': 32,
                 'accel': None,
                 'decel': None,
                 'fill': 'hold',
                 'grp': True,
                 'iterate': None,
                 'dur': 500,
                 'com': [48, -1, 20],
                 'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                           'dur="500"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                           'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                           'tm="100000"><p:val><p:fltVal '
                           'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                           'valueType="num"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                           'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                           'tm="100000"><p:val><p:fltVal '
                           'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="out" '
                           'filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn dur="1" '
                           'fill="hold"><p:stCondLst><p:cond '
                           'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                           'val="hidden"/></p:to></p:set>'},
 'salida.basic_zoom': {'ui': 'Basic Zoom (salida)',
                       'presetID': 23,
                       'presetClass': 'exit',
                       'presetSubtype': 16,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 500,
                       'com': [23, -1, 19],
                       'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                 'dur="500"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="4*ppt_w"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                 'valueType="num"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                 'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                                 'tm="100000"><p:val><p:strVal '
                                 'val="4*ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                 'dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="hidden"/></p:to></p:set>'},
 'salida.wipe_left': {'ui': 'Wipe hacia la izquierda',
                      'presetID': 22,
                      'presetClass': 'exit',
                      'presetSubtype': 8,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 500,
                      'com': [22, -1, 4],
                      'cuerpo': '<p:animEffect transition="out" filter="wipe(left)"><p:cBhvr><p:cTn '
                                'dur="500"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                'dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="hidden"/></p:to></p:set>'},
 'salida.wipe_bottom': {'ui': 'Wipe hacia abajo',
                        'presetID': 22,
                        'presetClass': 'exit',
                        'presetSubtype': 4,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 500,
                        'com': [22, -1, 3],
                        'cuerpo': '<p:animEffect transition="out" filter="wipe(down)"><p:cBhvr><p:cTn '
                                  'dur="500"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                  'dur="1" fill="hold"><p:stCondLst><p:cond '
                                  'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="hidden"/></p:to></p:set>'},
 'salida.shrink_turn': {'ui': 'Shrink & Turn',
                        'presetID': 31,
                        'presetClass': 'exit',
                        'presetSubtype': 0,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 1000,
                        'com': [31, -1, None],
                        'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                  'dur="1000"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:fltVal '
                                  'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:fltVal '
                                  'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                  'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:fltVal '
                                  'val="90"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                  'transition="out" filter="fade"><p:cBhvr><p:cTn '
                                  'dur="1000"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                  'dur="1" fill="hold"><p:stCondLst><p:cond '
                                  'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="hidden"/></p:to></p:set>'},
 'salida.split_vertical_in': {'ui': 'Split vertical hacia dentro (salida)',
                              'presetID': 16,
                              'presetClass': 'exit',
                              'presetSubtype': 21,
                              'accel': None,
                              'decel': None,
                              'fill': 'hold',
                              'grp': True,
                              'iterate': None,
                              'dur': 500,
                              'com': [16, -1, 25],
                              'cuerpo': '<p:animEffect transition="out" '
                                        'filter="barn(inVertical)"><p:cBhvr><p:cTn '
                                        'dur="500"/><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                        'dur="1" fill="hold"><p:stCondLst><p:cond '
                                        'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                        'val="hidden"/></p:to></p:set>'},
 'salida.random_bars_h': {'ui': 'Random Bars horizontal (salida)',
                          'presetID': 14,
                          'presetClass': 'exit',
                          'presetSubtype': 10,
                          'accel': None,
                          'decel': None,
                          'fill': 'hold',
                          'grp': True,
                          'iterate': None,
                          'dur': 500,
                          'com': [14, -1, 16],
                          'cuerpo': '<p:animEffect transition="out" '
                                    'filter="randombar(horizontal)"><p:cBhvr><p:cTn '
                                    'dur="500"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                    'dur="1" fill="hold"><p:stCondLst><p:cond '
                                    'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                    'val="hidden"/></p:to></p:set>'},
 'salida.swivel': {'ui': 'Swivel (salida)',
                   'presetID': 19,
                   'presetClass': 'exit',
                   'presetSubtype': 10,
                   'accel': None,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 5000,
                   'com': [19, -1, 16],
                   'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                             'dur="5000"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="5000"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                             'tm="5000"><p:val><p:strVal val="0.92*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="10000"><p:val><p:strVal val="0.71*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="15000"><p:val><p:strVal val="0.38*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="20000"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                             'tm="25000"><p:val><p:strVal val="-0.38*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="30000"><p:val><p:strVal val="-0.71*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="35000"><p:val><p:strVal val="-0.92*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="40000"><p:val><p:strVal val="-ppt_w"/></p:val></p:tav><p:tav '
                             'tm="45000"><p:val><p:strVal val="-0.92*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="50000"><p:val><p:strVal val="-0.71*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="55000"><p:val><p:strVal val="-0.38*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="60000"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                             'tm="65000"><p:val><p:strVal val="0.38*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="70000"><p:val><p:strVal val="0.71*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="75000"><p:val><p:strVal val="0.92*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                             'tm="85000"><p:val><p:strVal val="0.92*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="90000"><p:val><p:strVal val="0.71*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="95000"><p:val><p:strVal val="0.38*ppt_w"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:fltVal '
                             'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn dur="1" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="4999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="hidden"/></p:to></p:set>'},
 'salida.bounce': {'ui': 'Bounce (salida)',
                   'presetID': 26,
                   'presetClass': 'exit',
                   'presetSubtype': 0,
                   'accel': None,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 2000,
                   'com': [26, -1, None],
                   'cuerpo': '<p:animEffect transition="out" filter="wipe(down)"><p:cBhvr><p:cTn dur="180" '
                             'accel="50000"><p:stCondLst><p:cond '
                             'delay="1820"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="1822" tmFilter="0,0; 0.14,0.31; '
                             '0.43,0.73; 0.71,0.91; 1.0,1.0"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="#ppt_x+0.25"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="178"><p:stCondLst><p:cond '
                             'delay="1822"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="664" '
                             'tmFilter="0.0,0.0;0.25,0.07;0.50,0.2;0.75,0.467;1.0,1.0"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                             'tm="5000"><p:val><p:strVal val="ppt_y+0.026"/></p:val></p:tav><p:tav '
                             'tm="10000"><p:val><p:strVal val="ppt_y+0.052"/></p:val></p:tav><p:tav '
                             'tm="15000"><p:val><p:strVal val="ppt_y+0.078"/></p:val></p:tav><p:tav '
                             'tm="20000"><p:val><p:strVal val="ppt_y+0.103"/></p:val></p:tav><p:tav '
                             'tm="30000"><p:val><p:strVal val="ppt_y+0.151"/></p:val></p:tav><p:tav '
                             'tm="40000"><p:val><p:strVal val="ppt_y+0.196"/></p:val></p:tav><p:tav '
                             'tm="50000"><p:val><p:strVal val="ppt_y+0.236"/></p:val></p:tav><p:tav '
                             'tm="60000"><p:val><p:strVal val="ppt_y+0.270"/></p:val></p:tav><p:tav '
                             'tm="70000"><p:val><p:strVal val="ppt_y+0.297"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:strVal val="ppt_y+0.317"/></p:val></p:tav><p:tav '
                             'tm="90000"><p:val><p:strVal val="ppt_y+0.329"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_y+0.333"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="664" tmFilter="0, 0; 0.125,0.2665; '
                             '0.25,0.4; 0.375,0.465; 0.5,0.5;  0.625,0.535; 0.75,0.6; 0.875,0.7335; '
                             '1,1"><p:stCondLst><p:cond delay="664"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                             'tm="10000"><p:val><p:strVal val="ppt_y-0.034"/></p:val></p:tav><p:tav '
                             'tm="20000"><p:val><p:strVal val="ppt_y-0.065"/></p:val></p:tav><p:tav '
                             'tm="30000"><p:val><p:strVal val="ppt_y-0.090"/></p:val></p:tav><p:tav '
                             'tm="40000"><p:val><p:strVal val="ppt_y-0.106"/></p:val></p:tav><p:tav '
                             'tm="50000"><p:val><p:strVal val="ppt_y-0.111"/></p:val></p:tav><p:tav '
                             'tm="60000"><p:val><p:strVal val="ppt_y-0.106"/></p:val></p:tav><p:tav '
                             'tm="70000"><p:val><p:strVal val="ppt_y-0.090"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:strVal val="ppt_y-0.065"/></p:val></p:tav><p:tav '
                             'tm="90000"><p:val><p:strVal val="ppt_y-0.034"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="332" tmFilter="0, 0; 0.125,0.2665; '
                             '0.25,0.4; 0.375,0.465; 0.5,0.5;  0.625,0.535; 0.75,0.6; 0.875,0.7335; '
                             '1,1"><p:stCondLst><p:cond '
                             'delay="1324"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                             'tm="10000"><p:val><p:strVal val="ppt_y-0.011"/></p:val></p:tav><p:tav '
                             'tm="20000"><p:val><p:strVal val="ppt_y-0.022"/></p:val></p:tav><p:tav '
                             'tm="30000"><p:val><p:strVal val="ppt_y-0.030"/></p:val></p:tav><p:tav '
                             'tm="40000"><p:val><p:strVal val="ppt_y-0.035"/></p:val></p:tav><p:tav '
                             'tm="50000"><p:val><p:strVal val="ppt_y-0.037"/></p:val></p:tav><p:tav '
                             'tm="60000"><p:val><p:strVal val="ppt_y-0.035"/></p:val></p:tav><p:tav '
                             'tm="70000"><p:val><p:strVal val="ppt_y-0.030"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:strVal val="ppt_y-0.022"/></p:val></p:tav><p:tav '
                             'tm="90000"><p:val><p:strVal val="ppt_y-0.011"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="164" tmFilter="0, 0; 0.125,0.2665; '
                             '0.25,0.4; 0.375,0.465; 0.5,0.5;  0.625,0.535; 0.75,0.6; 0.875,0.7335; '
                             '1,1"><p:stCondLst><p:cond '
                             'delay="1656"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                             'tm="10000"><p:val><p:strVal val="ppt_y-0.004"/></p:val></p:tav><p:tav '
                             'tm="20000"><p:val><p:strVal val="ppt_y-0.007"/></p:val></p:tav><p:tav '
                             'tm="30000"><p:val><p:strVal val="ppt_y-0.010"/></p:val></p:tav><p:tav '
                             'tm="40000"><p:val><p:strVal val="ppt_y-0.012"/></p:val></p:tav><p:tav '
                             'tm="50000"><p:val><p:strVal val="ppt_y-0.0123"/></p:val></p:tav><p:tav '
                             'tm="60000"><p:val><p:strVal val="ppt_y-0.012"/></p:val></p:tav><p:tav '
                             'tm="70000"><p:val><p:strVal val="ppt_y-0.010"/></p:val></p:tav><p:tav '
                             'tm="80000"><p:val><p:strVal val="ppt_y-0.007"/></p:val></p:tav><p:tav '
                             'tm="90000"><p:val><p:strVal val="ppt_y-0.004"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                             'valueType="num"><p:cBhvr><p:cTn dur="180" accel="50000"><p:stCondLst><p:cond '
                             'delay="1820"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                             'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                             'tm="100000"><p:val><p:strVal '
                             'val="ppt_y+ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:animScale><p:cBhvr><p:cTn '
                             'dur="26"><p:stCondLst><p:cond '
                             'delay="620"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="60000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                             'decel="50000"><p:stCondLst><p:cond '
                             'delay="646"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="100000"/></p:animScale><p:animScale><p:cBhvr><p:cTn '
                             'dur="26"><p:stCondLst><p:cond '
                             'delay="1312"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="80000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                             'decel="50000"><p:stCondLst><p:cond '
                             'delay="1338"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="100000"/></p:animScale><p:animScale><p:cBhvr><p:cTn '
                             'dur="26"><p:stCondLst><p:cond '
                             'delay="1642"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="90000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                             'decel="50000"><p:stCondLst><p:cond '
                             'delay="1668"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="100000"/></p:animScale><p:animScale><p:cBhvr><p:cTn '
                             'dur="26"><p:stCondLst><p:cond '
                             'delay="1808"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="95000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="166" '
                             'decel="50000"><p:stCondLst><p:cond '
                             'delay="1834"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="100000" '
                             'y="100000"/></p:animScale><p:set><p:cBhvr><p:cTn dur="1" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="1999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="hidden"/></p:to></p:set>'},
 'salida.shape_circle_in': {'ui': 'Shape circulo hacia dentro (salida)',
                            'presetID': 6,
                            'presetClass': 'exit',
                            'presetSubtype': 16,
                            'accel': None,
                            'decel': None,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 2000,
                            'com': [6, -1, 19],
                            'cuerpo': '<p:animEffect transition="out" filter="circle(in)"><p:cBhvr><p:cTn '
                                      'dur="2000"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                      'dur="1" fill="hold"><p:stCondLst><p:cond '
                                      'delay="1999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="hidden"/></p:to></p:set>'},
 'salida.wheel_1': {'ui': 'Wheel (salida)',
                    'presetID': 21,
                    'presetClass': 'exit',
                    'presetSubtype': 1,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 2000,
                    'com': [21, -1, None],
                    'cuerpo': '<p:animEffect transition="out" filter="wheel(1)"><p:cBhvr><p:cTn '
                              'dur="2000"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                              'dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="1999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="hidden"/></p:to></p:set>'},
 'salida.peek_out_bottom': {'ui': 'Peek Out hacia abajo',
                            'presetID': 12,
                            'presetClass': 'exit',
                            'presetSubtype': 4,
                            'accel': None,
                            'decel': None,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 500,
                            'com': [12, -1, 3],
                            'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr '
                                      'additive="base"><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                      'tm="0"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav><p:tav '
                                      'tm="100000"><p:val><p:strVal '
                                      'val="#ppt_y+#ppt_h*1.125000"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                      'transition="out" filter="wipe(down)"><p:cBhvr><p:cTn '
                                      'dur="500"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                      'dur="1" fill="hold"><p:stCondLst><p:cond '
                                      'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="hidden"/></p:to></p:set>'},
 'salida.sink_down': {'ui': 'Sink Down (Rise Up salida)',
                      'presetID': 37,
                      'presetClass': 'exit',
                      'presetSubtype': 0,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 1000,
                      'com': [34, -1, None],
                      'cuerpo': '<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn '
                                'dur="1000"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="100" decel="100000"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="ppt_y-.03"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                                'valueType="num"><p:cBhvr><p:cTn dur="900" '
                                'accel="100000"><p:stCondLst><p:cond '
                                'delay="100"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                'tm="100000"><p:val><p:strVal '
                                'val="ppt_y+1"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn '
                                'dur="1" fill="hold"><p:stCondLst><p:cond '
                                'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                'val="hidden"/></p:to></p:set>'},
 'salida.center_revolve': {'ui': 'Center Revolve (salida)',
                           'presetID': 43,
                           'presetClass': 'exit',
                           'presetSubtype': 0,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 1000,
                           'com': [40, -1, None],
                           'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="600" '
                                     'decel="50000"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                     'tm="5000"><p:val><p:strVal val="ppt_x+0.0242"/></p:val></p:tav><p:tav '
                                     'tm="10000"><p:val><p:strVal val="ppt_x+0.0479"/></p:val></p:tav><p:tav '
                                     'tm="15000"><p:val><p:strVal val="ppt_x+0.0704"/></p:val></p:tav><p:tav '
                                     'tm="20000"><p:val><p:strVal val="ppt_x+0.0911"/></p:val></p:tav><p:tav '
                                     'tm="25000"><p:val><p:strVal val="ppt_x+0.1096"/></p:val></p:tav><p:tav '
                                     'tm="30000"><p:val><p:strVal val="ppt_x+0.1254"/></p:val></p:tav><p:tav '
                                     'tm="35000"><p:val><p:strVal val="ppt_x+0.1381"/></p:val></p:tav><p:tav '
                                     'tm="40000"><p:val><p:strVal val="ppt_x+0.1474"/></p:val></p:tav><p:tav '
                                     'tm="45000"><p:val><p:strVal val="ppt_x+0.1531"/></p:val></p:tav><p:tav '
                                     'tm="50000"><p:val><p:strVal val="ppt_x+0.155"/></p:val></p:tav><p:tav '
                                     'tm="55000"><p:val><p:strVal val="ppt_x+0.1531"/></p:val></p:tav><p:tav '
                                     'tm="60000"><p:val><p:strVal val="ppt_x+0.1474"/></p:val></p:tav><p:tav '
                                     'tm="65000"><p:val><p:strVal val="ppt_x+0.1381"/></p:val></p:tav><p:tav '
                                     'tm="70000"><p:val><p:strVal val="ppt_x+0.1254"/></p:val></p:tav><p:tav '
                                     'tm="75000"><p:val><p:strVal val="ppt_x+0.1096"/></p:val></p:tav><p:tav '
                                     'tm="80000"><p:val><p:strVal val="ppt_x+0.0911"/></p:val></p:tav><p:tav '
                                     'tm="85000"><p:val><p:strVal val="ppt_x+0.0704"/></p:val></p:tav><p:tav '
                                     'tm="90000"><p:val><p:strVal val="ppt_x+0.0479"/></p:val></p:tav><p:tav '
                                     'tm="95000"><p:val><p:strVal val="ppt_x+0.0242"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                     'dur="400"><p:stCondLst><p:cond '
                                     'delay="600"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_x"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_x"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="600" '
                                     'decel="50000"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                     'tm="5000"><p:val><p:strVal val="ppt_y+0.0019"/></p:val></p:tav><p:tav '
                                     'tm="10000"><p:val><p:strVal val="ppt_y+0.0076"/></p:val></p:tav><p:tav '
                                     'tm="15000"><p:val><p:strVal val="ppt_y+0.0169"/></p:val></p:tav><p:tav '
                                     'tm="20000"><p:val><p:strVal val="ppt_y+0.0296"/></p:val></p:tav><p:tav '
                                     'tm="25000"><p:val><p:strVal val="ppt_y+0.0454"/></p:val></p:tav><p:tav '
                                     'tm="30000"><p:val><p:strVal val="ppt_y+0.0639"/></p:val></p:tav><p:tav '
                                     'tm="35000"><p:val><p:strVal val="ppt_y+0.0846"/></p:val></p:tav><p:tav '
                                     'tm="40000"><p:val><p:strVal val="ppt_y+0.1071"/></p:val></p:tav><p:tav '
                                     'tm="45000"><p:val><p:strVal val="ppt_y+0.1307"/></p:val></p:tav><p:tav '
                                     'tm="50000"><p:val><p:strVal val="ppt_y+0.155"/></p:val></p:tav><p:tav '
                                     'tm="55000"><p:val><p:strVal val="ppt_y+0.1792"/></p:val></p:tav><p:tav '
                                     'tm="60000"><p:val><p:strVal val="ppt_y+0.2029"/></p:val></p:tav><p:tav '
                                     'tm="65000"><p:val><p:strVal val="ppt_y+0.2253"/></p:val></p:tav><p:tav '
                                     'tm="70000"><p:val><p:strVal val="ppt_y+0.2461"/></p:val></p:tav><p:tav '
                                     'tm="75000"><p:val><p:strVal val="ppt_y+0.2646"/></p:val></p:tav><p:tav '
                                     'tm="80000"><p:val><p:strVal val="ppt_y+0.2804"/></p:val></p:tav><p:tav '
                                     'tm="85000"><p:val><p:strVal val="ppt_y+0.2931"/></p:val></p:tav><p:tav '
                                     'tm="90000"><p:val><p:strVal val="ppt_y+0.3024"/></p:val></p:tav><p:tav '
                                     'tm="95000"><p:val><p:strVal val="ppt_y+0.308"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_y+0.31"/></p:val></p:tav></p:tavLst></p:anim><p:anim '
                                     'calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                                     'dur="400"><p:stCondLst><p:cond '
                                     'delay="600"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                     'tm="0"><p:val><p:strVal val="ppt_y"/></p:val></p:tav><p:tav '
                                     'tm="100000"><p:val><p:strVal '
                                     'val="ppt_y"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                                     'transition="out" filter="fade"><p:cBhvr><p:cTn '
                                     'dur="100"><p:stCondLst><p:cond '
                                     'delay="900"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                     'dur="1" fill="hold"><p:stCondLst><p:cond '
                                     'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="hidden"/></p:to></p:set>'},
 'salida.pinwheel': {'ui': 'Pinwheel (salida)',
                     'presetID': 35,
                     'presetClass': 'exit',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 2000,
                     'com': [33, -1, None],
                     'cuerpo': '<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn '
                               'dur="2000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:fltVal '
                               'val="720"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:fltVal '
                               'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:fltVal '
                               'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:set><p:cBhvr><p:cTn dur="1" '
                               'fill="hold"><p:stCondLst><p:cond '
                               'delay="1999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="hidden"/></p:to></p:set>'},
 'salida.dissolve_out': {'ui': 'Dissolve Out',
                         'presetID': 9,
                         'presetClass': 'exit',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [9, -1, None],
                         'cuerpo': '<p:animEffect transition="out" filter="dissolve"><p:cBhvr><p:cTn '
                                   'dur="500"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                   'dur="1" fill="hold"><p:stCondLst><p:cond '
                                   'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="hidden"/></p:to></p:set>'},
 'salida.collapse': {'ui': 'Collapse (Expand salida)',
                     'presetID': 55,
                     'presetClass': 'exit',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 1000,
                     'com': [50, -1, None],
                     'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                               'dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="ppt_w*0.70"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                               'transition="out" filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                               'dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="hidden"/></p:to></p:set>'},
 'salida.stretchy': {'ui': 'Stretchy / Compress (salida)',
                     'presetID': 50,
                     'presetClass': 'exit',
                     'presetSubtype': 0,
                     'accel': 100000,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 1000,
                     'com': [45, -1, None],
                     'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                               'dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="ppt_w+.3"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                               'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                               'tm="100000"><p:val><p:strVal '
                               'val="ppt_h"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect '
                               'transition="out" filter="fade"><p:cBhvr><p:cTn dur="1000"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                               'dur="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="hidden"/></p:to></p:set>'},
 'salida.curve_down': {'ui': 'Curve Down (Arc Up salida)',
                       'presetID': 52,
                       'presetClass': 'exit',
                       'presetSubtype': 0,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': None,
                       'dur': 1000,
                       'com': [47, -1, None],
                       'cuerpo': '<p:animScale><p:cBhvr><p:cTn dur="1000" accel="50000"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="100000" y="100000"/><p:to '
                                 'x="250000" y="250000"/></p:animScale><p:animMotion origin="layout" path="M '
                                 '0.0000 0.0000 C 0.03802 0.0 0.1441 0.02341 0.1826 0.0915 C 0.22118 0.15964 '
                                 '0.24705 0.31256 0.2318 0.4083 C 0.21649 0.50394 0.20747 0.57948 0.0908 '
                                 '0.6661 C -0.02552 0.75279 -0.37517 0.88508 -0.4674 0.9289" '
                                 'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="1000" '
                                 'accel="50000"><p:stCondLst><p:cond '
                                 'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion><p:animEffect '
                                 'transition="out" filter="fade"><p:cBhvr><p:cTn '
                                 'dur="1000"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                                 'dur="1" fill="hold"><p:stCondLst><p:cond '
                                 'delay="999"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="hidden"/></p:to></p:set>'},
 'salida.spinner': {'ui': 'Spinner (salida)',
                    'presetID': 49,
                    'presetClass': 'exit',
                    'presetSubtype': 0,
                    'accel': 100000,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 500,
                    'com': [44, -1, None],
                    'cuerpo': '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn '
                              'dur="500"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="ppt_w"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:fltVal '
                              'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:strVal val="ppt_h"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:fltVal '
                              'val="0"/></p:val></p:tav></p:tavLst></p:anim><p:anim calcmode="lin" '
                              'valueType="num"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                              'tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav '
                              'tm="100000"><p:val><p:fltVal '
                              'val="360"/></p:val></p:tav></p:tavLst></p:anim><p:animEffect transition="out" '
                              'filter="fade"><p:cBhvr><p:cTn dur="500"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn '
                              'dur="1" fill="hold"><p:stCondLst><p:cond '
                              'delay="499"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="hidden"/></p:to></p:set>'},
 'enfasis.pulse': {'ui': 'Pulse (Flash Bulb)',
                   'presetID': 26,
                   'presetClass': 'emph',
                   'presetSubtype': 0,
                   'accel': None,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 500,
                   'com': [75, 0, None],
                   'cuerpo': '<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn dur="500" '
                             'tmFilter="0, 0; .2, .5; .8, .5; 1, 0"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect><p:animScale><p:cBhvr><p:cTn '
                             'dur="250" autoRev="1" fill="hold"/><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:by x="105000" '
                             'y="105000"/></p:animScale>'},
 'enfasis.color_pulse': {'ui': 'Color Pulse? (Flicker)',
                         'presetID': 27,
                         'presetClass': 'emph',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'remove',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [76, 0, None],
                         'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                   'dur="250" autoRev="1" fill="remove"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:schemeClr '
                                   'val="bg1"/></p:to></p:animClr><p:animClr clrSpc="rgb" '
                                   'dir="cw"><p:cBhvr><p:cTn dur="250" autoRev="1" '
                                   'fill="remove"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:schemeClr '
                                   'val="bg1"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="250" '
                                   'autoRev="1" fill="remove"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="solid"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="250" autoRev="1" '
                                   'fill="remove"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.on</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="true"/></p:to></p:set>'},
 'enfasis.teeter': {'ui': 'Teeter',
                    'presetID': 32,
                    'presetClass': 'emph',
                    'presetSubtype': 0,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 1000,
                    'com': [80, 0, None],
                    'cuerpo': '<p:animRot by="120000"><p:cBhvr><p:cTn dur="100" '
                              'fill="hold"><p:stCondLst><p:cond '
                              'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                              'by="-240000"><p:cBhvr><p:cTn dur="200" fill="hold"><p:stCondLst><p:cond '
                              'delay="200"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                              'by="240000"><p:cBhvr><p:cTn dur="200" fill="hold"><p:stCondLst><p:cond '
                              'delay="400"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                              'by="-240000"><p:cBhvr><p:cTn dur="200" fill="hold"><p:stCondLst><p:cond '
                              'delay="600"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                              'by="120000"><p:cBhvr><p:cTn dur="200" fill="hold"><p:stCondLst><p:cond '
                              'delay="800"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'enfasis.spin_360': {'ui': 'Spin 360',
                      'presetID': 8,
                      'presetClass': 'emph',
                      'presetSubtype': 0,
                      'accel': None,
                      'decel': None,
                      'fill': 'hold',
                      'grp': True,
                      'iterate': None,
                      'dur': 2000,
                      'com': [61, 0, None],
                      'cuerpo': '<p:animRot by="21600000"><p:cBhvr><p:cTn dur="2000" '
                                'fill="hold"/><p:tgtEl><p:spTgt '
                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'enfasis.spin_90': {'ui': 'Spin 90 (parcial)',
                     'presetID': 8,
                     'presetClass': 'emph',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 2000,
                     'com': [61, 0, None],
                     'cuerpo': '<p:animRot by="5400000"><p:cBhvr><p:cTn dur="2000" '
                               'fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'enfasis.spin_ccw_180': {'ui': 'Spin -180 (antihorario)',
                          'presetID': 8,
                          'presetClass': 'emph',
                          'presetSubtype': 0,
                          'accel': None,
                          'decel': None,
                          'fill': 'hold',
                          'grp': True,
                          'iterate': None,
                          'dur': 2000,
                          'com': [61, 0, None],
                          'cuerpo': '<p:animRot by="-10800000"><p:cBhvr><p:cTn dur="2000" '
                                    'fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'enfasis.grow_shrink_110': {'ui': 'Grow/Shrink 110%',
                             'presetID': 6,
                             'presetClass': 'emph',
                             'presetSubtype': 0,
                             'accel': None,
                             'decel': None,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [59, 0, None],
                             'cuerpo': '<p:animScale><p:cBhvr><p:cTn dur="2000" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:by x="110000" '
                                       'y="110000"/></p:animScale>'},
 'enfasis.grow_shrink_150': {'ui': 'Grow/Shrink 150%',
                             'presetID': 6,
                             'presetClass': 'emph',
                             'presetSubtype': 0,
                             'accel': None,
                             'decel': None,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [59, 0, None],
                             'cuerpo': '<p:animScale><p:cBhvr><p:cTn dur="2000" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:by x="150000" '
                                       'y="150000"/></p:animScale>'},
 'enfasis.transparency_50': {'ui': 'Transparency 50%',
                             'presetID': 9,
                             'presetClass': 'emph',
                             'presetSubtype': 0,
                             'accel': None,
                             'decel': None,
                             'fill': None,
                             'grp': True,
                             'iterate': None,
                             'dur': None,
                             'com': [62, 0, None],
                             'cuerpo': '<p:set><p:cBhvr><p:cTn dur="indefinite"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.opacity</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                       'val="0.5"/></p:to></p:set><p:animEffect filter="image" '
                                       'prLst="opacity: 0.5"><p:cBhvr rctx="IE"><p:cTn '
                                       'dur="indefinite"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl></p:cBhvr></p:animEffect>'},
 'enfasis.fill_color': {'ui': 'Fill Color (rojo)',
                        'presetID': 1,
                        'presetClass': 'emph',
                        'presetSubtype': 2,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': False,
                        'iterate': None,
                        'dur': 2000,
                        'com': [54, 0, None],
                        'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr><p:cTn dur="2000" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                  'val="C00000"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="2000" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="solid"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="2000" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.on</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="true"/></p:to></p:set>'},
 'enfasis.object_color': {'ui': 'Object Color? (Color Blend)',
                          'presetID': 19,
                          'presetClass': 'emph',
                          'presetSubtype': 0,
                          'accel': None,
                          'decel': None,
                          'fill': 'hold',
                          'grp': True,
                          'iterate': None,
                          'dur': 500,
                          'com': [68, 0, None],
                          'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                    'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                    'val="C00000"/></p:to></p:animClr><p:animClr clrSpc="rgb" '
                                    'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                    'val="C00000"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="500" '
                                    'fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                    'val="solid"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                    'fill="hold"/><p:tgtEl><p:spTgt '
                                    'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.on</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                    'val="true"/></p:to></p:set>'},
 'enfasis.line_color': {'ui': 'Line Color',
                        'presetID': 7,
                        'presetClass': 'emph',
                        'presetSubtype': 2,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': False,
                        'iterate': None,
                        'dur': 2000,
                        'com': [60, 0, None],
                        'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr><p:cTn dur="2000" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                  'val="C00000"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="2000" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.on</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="true"/></p:to></p:set>'},
 'enfasis.font_color': {'ui': 'Font Color',
                        'presetID': 3,
                        'presetClass': 'emph',
                        'presetSubtype': 2,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 2000,
                        'com': [56, 0, None],
                        'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                  'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                  'val="C00000"/></p:to></p:animClr>'},
 'enfasis.brush_color': {'ui': 'Brush On Color',
                         'presetID': 16,
                         'presetClass': 'emph',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': ('lt', 'tmPct', 4000),
                         'dur': 500,
                         'com': [66, 0, None],
                         'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn dur="500" '
                                   'fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:clrVal><a:srgbClr '
                                   'val="C00000"/></p:clrVal></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                   'fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:clrVal><a:srgbClr '
                                   'val="C00000"/></p:clrVal></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                   'fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="solid"/></p:to></p:set>'},
 'enfasis.underline': {'ui': 'Brush On Underline',
                       'presetID': 18,
                       'presetClass': 'emph',
                       'presetSubtype': 0,
                       'accel': None,
                       'decel': None,
                       'fill': 'hold',
                       'grp': True,
                       'iterate': ('lt', 'tmPct', 4000),
                       'dur': 500,
                       'com': [67, 0, None],
                       'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn dur="500" '
                                 'fill="hold"/><p:tgtEl><p:spTgt '
                                 'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.textDecorationUnderline</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                 'val="true"/></p:to></p:set>'},
 'enfasis.bold_flash': {'ui': 'Bold Flash',
                        'presetID': 10,
                        'presetClass': 'emph',
                        'presetSubtype': 0,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 2000,
                        'com': [63, 0, None],
                        'cuerpo': '<p:anim calcmode="discrete" valueType="str"><p:cBhvr '
                                  'override="childStyle"><p:cTn dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontWeight</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav '
                                  'tm="0"><p:val><p:strVal val="normal"/></p:val></p:tav><p:tav '
                                  'tm="50000"><p:val><p:strVal val="bold"/></p:val></p:tav><p:tav '
                                  'tm="60000"><p:val><p:strVal val="normal"/></p:val></p:tav><p:tav '
                                  'tm="100000"><p:val><p:strVal '
                                  'val="normal"/></p:val></p:tav></p:tavLst></p:anim>'},
 'enfasis.bold_reveal': {'ui': 'Bold Reveal',
                         'presetID': 15,
                         'presetClass': 'emph',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': None,
                         'grp': True,
                         'iterate': ('lt', 'tmAbs', 25),
                         'dur': None,
                         'com': [65, 0, None],
                         'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn '
                                   'dur="indefinite"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontWeight</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="bold"/></p:to></p:set>'},
 'enfasis.wave': {'ui': 'Wave',
                  'presetID': 34,
                  'presetClass': 'emph',
                  'presetSubtype': 0,
                  'accel': None,
                  'decel': None,
                  'fill': 'hold',
                  'grp': True,
                  'iterate': ('lt', 'tmPct', 10000),
                  'dur': 500,
                  'com': [82, 0, None],
                  'cuerpo': '<p:animMotion origin="layout" path="M 0.0 0.0 L 0.0 -0.07213" '
                            'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="250" accel="50000" '
                            'decel="50000" autoRev="1" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion><p:animRot '
                            'by="1500000"><p:cBhvr><p:cTn dur="125" fill="hold"><p:stCondLst><p:cond '
                            'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                            'by="-1500000"><p:cBhvr><p:cTn dur="125" fill="hold"><p:stCondLst><p:cond '
                            'delay="125"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                            'by="-1500000"><p:cBhvr><p:cTn dur="125" fill="hold"><p:stCondLst><p:cond '
                            'delay="250"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot><p:animRot '
                            'by="1500000"><p:cBhvr><p:cTn dur="125" fill="hold"><p:stCondLst><p:cond '
                            'delay="375"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                            'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'enfasis.desaturate': {'ui': 'Desaturate',
                        'presetID': 25,
                        'presetClass': 'emph',
                        'presetSubtype': 0,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': None,
                        'dur': 500,
                        'com': [74, 0, None],
                        'cuerpo': '<p:animClr clrSpc="hsl" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                  'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                  'h="0" s="-70588" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                  'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                  'h="0" s="-70588" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                  'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                  'h="0" s="-70588" l="0"/></p:by></p:animClr><p:set><p:cBhvr><p:cTn '
                                  'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="solid"/></p:to></p:set>'},
 'enfasis.darken': {'ui': 'Darken',
                    'presetID': 24,
                    'presetClass': 'emph',
                    'presetSubtype': 0,
                    'accel': None,
                    'decel': None,
                    'fill': 'hold',
                    'grp': True,
                    'iterate': None,
                    'dur': 500,
                    'com': [73, 0, None],
                    'cuerpo': '<p:animClr clrSpc="hsl" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                              'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                              'h="0" s="-12549" l="-25098"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                              'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                              'h="0" s="-12549" l="-25098"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                              'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                              'h="0" s="-12549" l="-25098"/></p:by></p:animClr><p:set><p:cBhvr><p:cTn '
                              'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                              'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                              'val="solid"/></p:to></p:set>'},
 'enfasis.lighten': {'ui': 'Lighten',
                     'presetID': 30,
                     'presetClass': 'emph',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': None,
                     'dur': 500,
                     'com': [78, 0, None],
                     'cuerpo': '<p:animClr clrSpc="hsl" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                               'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                               'h="0" s="12549" l="25098"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                               'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                               'h="0" s="12549" l="25098"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                               'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                               'h="0" s="12549" l="25098"/></p:by></p:animClr><p:set><p:cBhvr><p:cTn '
                               'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                               'val="solid"/></p:to></p:set>'},
 'enfasis.complementary': {'ui': 'Complementary Color',
                           'presetID': 21,
                           'presetClass': 'emph',
                           'presetSubtype': 0,
                           'accel': None,
                           'decel': None,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 500,
                           'com': [70, 0, None],
                           'cuerpo': '<p:animClr clrSpc="hsl" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                     'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                     'h="7200000" s="0" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                     'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                     'h="7200000" s="0" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                     'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                     'h="7200000" s="0" l="0"/></p:by></p:animClr><p:set><p:cBhvr><p:cTn '
                                     'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="solid"/></p:to></p:set>'},
 'enfasis.complementary2': {'ui': 'Complementary Color 2',
                            'presetID': 22,
                            'presetClass': 'emph',
                            'presetSubtype': 0,
                            'accel': None,
                            'decel': None,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 500,
                            'com': [71, 0, None],
                            'cuerpo': '<p:animClr clrSpc="hsl" dir="cw"><p:cBhvr '
                                      'override="childStyle"><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                      'h="-7200000" s="0" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                      'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                      'h="-7200000" s="0" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                      'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                      'h="-7200000" s="0" l="0"/></p:by></p:animClr><p:set><p:cBhvr><p:cTn '
                                      'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="solid"/></p:to></p:set>'},
 'enfasis.contrasting': {'ui': 'Contrasting Color',
                         'presetID': 23,
                         'presetClass': 'emph',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 500,
                         'com': [72, 0, None],
                         'cuerpo': '<p:animClr clrSpc="hsl" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                   'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                   'h="10842353" s="0" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                   'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                   'h="10842353" s="0" l="0"/></p:by></p:animClr><p:animClr clrSpc="hsl" '
                                   'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>stroke.color</p:attrName></p:attrNameLst></p:cBhvr><p:by><p:hsl '
                                   'h="10842353" s="0" l="0"/></p:by></p:animClr><p:set><p:cBhvr><p:cTn '
                                   'dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="solid"/></p:to></p:set>'},
 'enfasis.color_wave': {'ui': 'Color Wave',
                        'presetID': 20,
                        'presetClass': 'emph',
                        'presetSubtype': 0,
                        'accel': None,
                        'decel': None,
                        'fill': 'hold',
                        'grp': True,
                        'iterate': ('lt', 'tmPct', 10000),
                        'dur': 1000,
                        'com': [69, 0, None],
                        'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn dur="500" autoRev="1" '
                                  'fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:clrVal><a:srgbClr '
                                  'val="C00000"/></p:clrVal></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                  'autoRev="1" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:clrVal><a:srgbClr '
                                  'val="C00000"/></p:clrVal></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                  'autoRev="1" fill="hold"/><p:tgtEl><p:spTgt '
                                  'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                  'val="solid"/></p:to></p:set>'},
 'enfasis.grow_with_color': {'ui': 'Grow With Color',
                             'presetID': 28,
                             'presetClass': 'emph',
                             'presetSubtype': 0,
                             'accel': None,
                             'decel': None,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': ('lt', 'tmPct', 10000),
                             'dur': 500,
                             'com': [77, 0, None],
                             'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr '
                                       'override="childStyle"><p:cTn dur="500" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                       'val="C00000"/></p:to></p:animClr><p:animClr clrSpc="rgb" '
                                       'dir="cw"><p:cBhvr><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:srgbClr '
                                       'val="C00000"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="500" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                       'val="solid"/></p:to></p:set><p:anim to="1.5" calcmode="lin" '
                                       'valueType="num"><p:cBhvr override="childStyle"><p:cTn dur="500" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontSize</p:attrName></p:attrNameLst></p:cBhvr></p:anim>'},
 'enfasis.vertical_grow': {'ui': 'Vertical Grow',
                           'presetID': 33,
                           'presetClass': 'emph',
                           'presetSubtype': 0,
                           'accel': None,
                           'decel': None,
                           'fill': 'remove',
                           'grp': True,
                           'iterate': None,
                           'dur': 3000,
                           'com': [81, 0, None],
                           'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                                     'dur="1500" accel="50000" autoRev="1" fill="hold" tmFilter="0, 0; '
                                     '.33333, 1; 1, 1"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:schemeClr '
                                     'val="accent2"/></p:to></p:animClr><p:animClr clrSpc="rgb" '
                                     'dir="cw"><p:cBhvr><p:cTn dur="1500" accel="50000" autoRev="1" '
                                     'fill="hold" tmFilter="0, 0; .33333, 1; 1, 1"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillcolor</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:schemeClr '
                                     'val="accent2"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="3000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="solid"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="3000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.on</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                     'val="true"/></p:to></p:set><p:animScale><p:cBhvr><p:cTn dur="1500" '
                                     'accel="50000" autoRev="1" fill="hold" tmFilter="0, 0; .33333, 1; 1, '
                                     '1"><p:stCondLst><p:cond '
                                     'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="100000" '
                                     'y="100000"/><p:to x="100000" y="140000"/></p:animScale>'},
 'enfasis.style_emphasis': {'ui': 'Style Emphasis',
                            'presetID': 31,
                            'presetClass': 'emph',
                            'presetSubtype': 0,
                            'accel': None,
                            'decel': None,
                            'fill': None,
                            'grp': True,
                            'iterate': None,
                            'dur': 500,
                            'com': [79, 0, None],
                            'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn dur="500" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:clrVal><a:schemeClr '
                                      'val="accent2"/></p:clrVal></p:to></p:set><p:set><p:cBhvr '
                                      'override="childStyle"><p:cTn dur="500" fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontStyle</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="italic"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontWeight</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="bold"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="500" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.textDecorationUnderline</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                      'val="true"/></p:to></p:set>'},
 'enfasis.change_font': {'ui': 'Change Font',
                         'presetID': 2,
                         'presetClass': 'emph',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': None,
                         'grp': True,
                         'iterate': None,
                         'dur': None,
                         'com': [55, 0, None],
                         'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn '
                                   'dur="indefinite"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontFamily</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                   'val="Courier New"/></p:to></p:set>'},
 'enfasis.change_font_size': {'ui': 'Change Font Size',
                              'presetID': 4,
                              'presetClass': 'emph',
                              'presetSubtype': 2,
                              'accel': None,
                              'decel': None,
                              'fill': 'hold',
                              'grp': True,
                              'iterate': None,
                              'dur': 2000,
                              'com': [57, 0, None],
                              'cuerpo': '<p:anim to="1.5" calcmode="lin" valueType="num"><p:cBhvr '
                                        'override="childStyle"><p:cTn dur="2000" '
                                        'fill="hold"/><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontSize</p:attrName></p:attrNameLst></p:cBhvr></p:anim>'},
 'enfasis.change_font_style': {'ui': 'Change Font Style',
                               'presetID': 5,
                               'presetClass': 'emph',
                               'presetSubtype': 1,
                               'accel': None,
                               'decel': None,
                               'fill': None,
                               'grp': True,
                               'iterate': None,
                               'dur': None,
                               'com': [58, 0, None],
                               'cuerpo': '<p:set><p:cBhvr override="childStyle"><p:cTn '
                                         'dur="indefinite"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontStyle</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                         'val="normal"/></p:to></p:set><p:set><p:cBhvr '
                                         'override="childStyle"><p:cTn dur="indefinite"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.fontWeight</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                         'val="bold"/></p:to></p:set><p:set><p:cBhvr '
                                         'override="childStyle"><p:cTn dur="indefinite"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.textDecorationUnderline</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                                         'val="false"/></p:to></p:set>'},
 'enfasis.blast': {'ui': 'Blast',
                   'presetID': 14,
                   'presetClass': 'emph',
                   'presetSubtype': 0,
                   'accel': None,
                   'decel': None,
                   'fill': 'hold',
                   'grp': True,
                   'iterate': None,
                   'dur': 2000,
                   'com': [64, 0, None],
                   'cuerpo': '<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr override="childStyle"><p:cTn '
                             'dur="1900" fill="hold"><p:stCondLst><p:cond '
                             'delay="100"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:schemeClr '
                             'val="accent2"/></p:to></p:animClr><p:animClr clrSpc="rgb" '
                             'dir="cw"><p:cBhvr><p:cTn dur="1900" fill="hold"><p:stCondLst><p:cond '
                             'delay="100"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fillColor</p:attrName></p:attrNameLst></p:cBhvr><p:to><a:schemeClr '
                             'val="accent2"/></p:to></p:animClr><p:set><p:cBhvr><p:cTn dur="1900" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="100"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.type</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="solid"/></p:to></p:set><p:set><p:cBhvr><p:cTn dur="1900" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="100"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>fill.on</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal '
                             'val="true"/></p:to></p:set><p:animScale><p:cBhvr><p:cTn dur="200" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="100000" y="100000"/><p:to '
                             'x="100000" y="5000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="200" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="200"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:from x="100000" y="5000"/><p:to '
                             'x="120000" y="150000"/></p:animScale><p:animScale><p:cBhvr><p:cTn dur="600" '
                             'fill="hold"><p:stCondLst><p:cond '
                             'delay="1400"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                             'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="120000" '
                             'y="150000"/></p:animScale>'},
 'enfasis.shimmer': {'ui': 'Shimmer',
                     'presetID': 36,
                     'presetClass': 'emph',
                     'presetSubtype': 0,
                     'accel': None,
                     'decel': None,
                     'fill': 'hold',
                     'grp': True,
                     'iterate': ('lt', 'tmPct', 10000),
                     'dur': 500,
                     'com': [52, 0, None],
                     'cuerpo': '<p:animScale><p:cBhvr><p:cTn dur="250" autoRev="1" '
                               'fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:to x="80000" '
                               'y="100000"/></p:animScale><p:anim by="(#ppt_w*0.10)" calcmode="lin" '
                               'valueType="num"><p:cBhvr><p:cTn dur="250" autoRev="1" '
                               'fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:anim '
                               'by="(-#ppt_w*0.10)" calcmode="lin" valueType="num"><p:cBhvr><p:cTn dur="250" '
                               'autoRev="1" fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:anim><p:animRot '
                               'by="-480000"><p:cBhvr><p:cTn dur="250" autoRev="1" '
                               'fill="hold"><p:stCondLst><p:cond '
                               'delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt '
                               'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'},
 'trayectoria.path_right': {'ui': 'Linea a la derecha',
                            'presetID': 63,
                            'presetClass': 'path',
                            'presetSubtype': 0,
                            'accel': 50000,
                            'decel': 50000,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 2000,
                            'com': [149, 0, None],
                            'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.25 0 E" '
                                      'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                      'fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_left': {'ui': 'Linea a la izquierda',
                           'presetID': 35,
                           'presetClass': 'path',
                           'presetSubtype': 0,
                           'accel': 50000,
                           'decel': 50000,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 2000,
                           'com': [120, 0, None],
                           'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L -0.25 0 E" '
                                     'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_down': {'ui': 'Linea hacia abajo',
                           'presetID': 42,
                           'presetClass': 'path',
                           'presetSubtype': 0,
                           'accel': 50000,
                           'decel': 50000,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 2000,
                           'com': [127, 0, None],
                           'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0 0.25 E" '
                                     'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_up': {'ui': 'Linea hacia arriba',
                         'presetID': 64,
                         'presetClass': 'path',
                         'presetSubtype': 0,
                         'accel': 50000,
                         'decel': 50000,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 2000,
                         'com': [148, 0, None],
                         'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0 -0.25 E" '
                                   'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                   'fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_diag_down_right': {'ui': 'Diagonal abajo-derecha',
                                      'presetID': 49,
                                      'presetClass': 'path',
                                      'presetSubtype': 0,
                                      'accel': 50000,
                                      'decel': 50000,
                                      'fill': 'hold',
                                      'grp': True,
                                      'iterate': None,
                                      'dur': 2000,
                                      'com': [134, 0, None],
                                      'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.25 0.25 E" '
                                                'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                                'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                                'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_circle': {'ui': 'Circulo',
                             'presetID': 1,
                             'presetClass': 'path',
                             'presetSubtype': 0,
                             'accel': 50000,
                             'decel': 50000,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [86, 0, None],
                             'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0.069 0 0.125 0.056 '
                                       '0.125 0.125 C 0.125 0.194 0.069 0.25 0 0.25 C -0.069 0.25 -0.125 '
                                       '0.194 -0.125 0.125 C -0.125 0.056 -0.069 0 0 0 Z" '
                                       'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_arc_down': {'ui': 'Arco hacia abajo',
                               'presetID': 37,
                               'presetClass': 'path',
                               'presetSubtype': 0,
                               'accel': 50000,
                               'decel': 50000,
                               'fill': 'hold',
                               'grp': True,
                               'iterate': None,
                               'dur': 2000,
                               'com': [122, 0, None],
                               'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.067 0.04 C 0.081 '
                                         '0.049 0.102 0.054 0.124 0.054 C 0.149 0.054 0.169 0.049 0.183 0.04 '
                                         'L 0.25 0 E" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                         'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_arc_up': {'ui': 'Arco hacia arriba',
                             'presetID': 44,
                             'presetClass': 'path',
                             'presetSubtype': 0,
                             'accel': 50000,
                             'decel': 50000,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [129, 0, None],
                             'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.067 -0.04 C 0.081 '
                                       '-0.049 0.102 -0.054 0.124 -0.054 C 0.149 -0.054 0.169 -0.049 0.183 '
                                       '-0.04 L 0.25 0 E" pathEditMode="relative" '
                                       'ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_arc_right': {'ui': 'Arco a la derecha',
                                'presetID': 58,
                                'presetClass': 'path',
                                'presetSubtype': 0,
                                'accel': 50000,
                                'decel': 50000,
                                'fill': 'hold',
                                'grp': True,
                                'iterate': None,
                                'dur': 2000,
                                'com': [143, 0, None],
                                'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.04 0.067 C 0.049 '
                                          '0.081 0.054 0.102 0.054 0.124 C 0.054 0.149 0.049 0.169 0.04 '
                                          '0.183 L 0 0.25 E" pathEditMode="relative" '
                                          'ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                          'fill="hold"/><p:tgtEl><p:spTgt '
                                          'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_s_curve': {'ui': 'Curva en S 1',
                              'presetID': 59,
                              'presetClass': 'path',
                              'presetSubtype': 0,
                              'accel': 50000,
                              'decel': 50000,
                              'fill': 'hold',
                              'grp': True,
                              'iterate': None,
                              'dur': 2000,
                              'com': [144, 0, None],
                              'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0 -0.035 0.028 -0.062 '
                                        '0.062 -0.062 C 0.097 -0.062 0.125 -0.035 0.125 0 C 0.125 0.035 '
                                        '0.153 0.062 0.188 0.062 C 0.222 0.062 0.25 0.035 0.25 0 E" '
                                        'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                        'fill="hold"/><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_zigzag': {'ui': 'Zigzag',
                             'presetID': 38,
                             'presetClass': 'path',
                             'presetSubtype': 0,
                             'accel': 50000,
                             'decel': 50000,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [123, 0, None],
                             'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.016 0.099 L 0.031 0 L '
                                       '0.047 0.099 L 0.063 0 L 0.078 0.099 L 0.094 0 L 0.109 0.099 L 0.125 '
                                       '0 L 0.141 0.099 L 0.156 0 L 0.172 0.099 L 0.187 0 L 0.203 0.099 L '
                                       '0.219 0 L 0.234 0.099 L 0.25 0 E" pathEditMode="relative" '
                                       'ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                       'fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_sine_wave': {'ui': 'Onda senoidal',
                                'presetID': 40,
                                'presetClass': 'path',
                                'presetSubtype': 0,
                                'accel': 50000,
                                'decel': 50000,
                                'fill': 'hold',
                                'grp': True,
                                'iterate': None,
                                'dur': 2000,
                                'com': [125, 0, None],
                                'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0.003 -0.019 0.007 '
                                          '-0.037 0.015 -0.037 C 0.024 -0.037 0.027 -0.019 0.03 0 C 0.034 '
                                          '0.021 0.037 0.042 0.047 0.042 C 0.056 0.042 0.059 0.021 0.063 0 C '
                                          '0.065 -0.019 0.069 -0.037 0.078 -0.037 C 0.086 -0.037 0.09 -0.019 '
                                          '0.093 0 C 0.096 0.021 0.1 0.042 0.109 0.042 C 0.118 0.042 0.125 0 '
                                          '0.125 0 C 0.128 -0.019 0.131 -0.037 0.14 -0.037 C 0.149 -0.037 '
                                          '0.152 -0.019 0.155 0 C 0.159 0.021 0.162 0.042 0.172 0.042 C '
                                          '0.181 0.042 0.184 0.021 0.187 0 C 0.191 -0.019 0.194 -0.037 0.203 '
                                          '-0.037 C 0.211 -0.037 0.215 -0.019 0.218 0 C 0.221 0.021 0.225 '
                                          '0.042 0.234 0.042 C 0.243 0.042 0.246 0.021 0.25 0 E" '
                                          'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                          'fill="hold"/><p:tgtEl><p:spTgt '
                                          'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_wave': {'ui': 'Onda',
                           'presetID': 47,
                           'presetClass': 'path',
                           'presetSubtype': 0,
                           'accel': 50000,
                           'decel': 50000,
                           'fill': 'hold',
                           'grp': True,
                           'iterate': None,
                           'dur': 2000,
                           'com': [132, 0, None],
                           'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0.002 0.063 0.009 0.108 '
                                     '0.016 0.108 C 0.023 0.108 0.029 0.063 0.031 0 C 0.034 0.063 0.04 0.108 '
                                     '0.047 0.108 C 0.054 0.108 0.06 0.063 0.062 0 C 0.065 0.063 0.071 0.108 '
                                     '0.078 0.108 C 0.085 0.108 0.092 0.063 0.094 0 C 0.096 0.063 0.102 '
                                     '0.108 0.11 0.108 C 0.116 0.108 0.123 0.063 0.125 0 C 0.127 0.063 0.134 '
                                     '0.108 0.141 0.108 C 0.148 0.108 0.154 0.063 0.156 0 C 0.159 0.063 '
                                     '0.165 0.108 0.172 0.108 C 0.179 0.108 0.185 0.063 0.188 0 C 0.19 0.063 '
                                     '0.196 0.108 0.203 0.108 C 0.21 0.108 0.217 0.063 0.219 0 C 0.221 0.063 '
                                     '0.227 0.108 0.235 0.108 C 0.242 0.108 0.248 0.063 0.25 0 E" '
                                     'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                     'fill="hold"/><p:tgtEl><p:spTgt '
                                     'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_5_point_star': {'ui': 'Estrella de 5 puntas',
                                   'presetID': 5,
                                   'presetClass': 'path',
                                   'presetSubtype': 0,
                                   'accel': 50000,
                                   'decel': 50000,
                                   'fill': 'hold',
                                   'grp': True,
                                   'iterate': None,
                                   'dur': 2000,
                                   'com': [90, 0, None],
                                   'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.029 0.091 L '
                                             '0.125 0.091 L 0.048 0.147 L 0.077 0.238 L 0 0.182 L -0.077 '
                                             '0.238 L -0.048 0.147 L -0.125 0.091 L -0.029 0.091 L 0 0 Z" '
                                             'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                             'fill="hold"/><p:tgtEl><p:spTgt '
                                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_heart': {'ui': 'Corazon',
                            'presetID': 9,
                            'presetClass': 'path',
                            'presetSubtype': 0,
                            'accel': 50000,
                            'decel': 50000,
                            'fill': 'hold',
                            'grp': True,
                            'iterate': None,
                            'dur': 2000,
                            'com': [94, 0, None],
                            'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0.012 -0.018 0.033 -0.044 '
                                      '0.058 -0.044 C 0.095 -0.044 0.125 -0.017 0.125 0.017 C 0.125 0.028 '
                                      '0.122 0.038 0.116 0.047 C 0.117 0.047 0 0.182 0 0.183 C 0 0.182 '
                                      '-0.117 0.047 -0.116 0.047 C -0.122 0.038 -0.125 0.028 -0.125 0.017 C '
                                      '-0.125 -0.017 -0.095 -0.044 -0.057 -0.044 C -0.033 -0.044 -0.012 '
                                      '-0.018 0 0 Z" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                      'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                      'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_loop_de_loop': {'ui': 'Rizo (loop de loop)',
                                   'presetID': 24,
                                   'presetClass': 'path',
                                   'presetSubtype': 0,
                                   'accel': 50000,
                                   'decel': 50000,
                                   'fill': 'hold',
                                   'grp': True,
                                   'iterate': None,
                                   'dur': 2000,
                                   'com': [109, 0, None],
                                   'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0.023 0.001 0.042 '
                                             '0.009 0.052 0.021 L 0.075 0.049 C 0.08 0.055 0.088 0.058 0.098 '
                                             '0.058 C 0.112 0.058 0.124 0.05 0.125 0.038 C 0.124 0.028 0.112 '
                                             '0.019 0.098 0.019 C 0.088 0.019 0.08 0.023 0.075 0.028 L 0.052 '
                                             '0.056 C 0.042 0.068 0.023 0.076 0 0.077 C -0.023 0.076 -0.042 '
                                             '0.068 -0.052 0.056 L -0.075 0.028 C -0.08 0.023 -0.088 0.019 '
                                             '-0.098 0.019 C -0.112 0.019 -0.124 0.028 -0.125 0.038 C -0.124 '
                                             '0.05 -0.112 0.058 -0.098 0.058 C -0.088 0.058 -0.08 0.055 '
                                             '-0.075 0.049 L -0.052 0.021 C -0.042 0.009 -0.023 0.001 0 0 Z" '
                                             'pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                             'fill="hold"/><p:tgtEl><p:spTgt '
                                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_spiral_right': {'ui': 'Espiral a la derecha',
                                   'presetID': 46,
                                   'presetClass': 'path',
                                   'presetSubtype': 0,
                                   'accel': 50000,
                                   'decel': 50000,
                                   'fill': 'hold',
                                   'grp': True,
                                   'iterate': None,
                                   'dur': 2000,
                                   'com': [131, 0, None],
                                   'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C -0.004 -0.067 '
                                             '0.046 -0.125 0.113 -0.129 C 0.177 -0.134 0.237 -0.089 0.241 '
                                             '-0.024 C 0.246 0.036 0.204 0.092 0.144 0.096 C 0.089 0.099 '
                                             '0.037 0.062 0.033 0.006 C 0.029 -0.045 0.064 -0.093 0.115 '
                                             '-0.097 C 0.162 -0.1 0.206 -0.069 0.209 -0.022 C 0.212 0.02 '
                                             '0.184 0.061 0.142 0.063 C 0.104 0.066 0.068 0.042 0.065 0.004 '
                                             'C 0.063 -0.03 0.084 -0.063 0.117 -0.065 C 0.146 -0.067 0.175 '
                                             '-0.049 0.177 -0.02 C 0.179 0.005 0.164 0.029 0.14 0.031 C 0.12 '
                                             '0.033 0.099 0.022 0.098 0.002 C 0.096 -0.014 0.104 -0.031 '
                                             '0.119 -0.033 C 0.131 -0.033 0.143 -0.029 0.145 -0.018 C 0.146 '
                                             '-0.011 0.144 -0.004 0.138 -0.001 C 0.135 0 0.133 0 0.13 -0.001 '
                                             'E" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                             'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_turn_right': {'ui': 'Giro a la derecha',
                                 'presetID': 36,
                                 'presetClass': 'path',
                                 'presetSubtype': 0,
                                 'accel': 50000,
                                 'decel': 50000,
                                 'fill': 'hold',
                                 'grp': True,
                                 'iterate': None,
                                 'dur': 2000,
                                 'com': [121, 0, None],
                                 'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0 0.125 C 0 0.181 '
                                           '0.069 0.25 0.125 0.25 L 0.25 0.25 E" pathEditMode="relative" '
                                           'ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                           'fill="hold"/><p:tgtEl><p:spTgt '
                                           'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_bounce_right': {'ui': 'Rebote a la derecha',
                                   'presetID': 54,
                                   'presetClass': 'path',
                                   'presetSubtype': 0,
                                   'accel': 50000,
                                   'decel': 50000,
                                   'fill': 'hold',
                                   'grp': True,
                                   'iterate': None,
                                   'dur': 2000,
                                   'com': [139, 0, None],
                                   'cuerpo': '<p:animMotion origin="layout" path="M 0 0 c 0.004 -0.008 0.018 '
                                             '-0.016 0.023 -0.016 c 0.031 0 0.063 0.125 0.063 0.25 c 0 '
                                             '-0.063 0.016 -0.125 0.031 -0.125 c 0.016 0 0.031 0.063 0.031 '
                                             '0.125 c 0 -0.031 0.008 -0.063 0.016 -0.063 c 0.008 0 0.016 '
                                             '0.031 0.016 0.063 c 0 -0.016 0.004 -0.031 0.008 -0.031 c 0.004 '
                                             '0 0.008 0.016 0.008 0.031 c 0 -0.008 0.002 -0.016 0.004 -0.016 '
                                             'c 0.001 0 0.004 0.008 0.004 0.016 c 0 -0.004 0.001 -0.008 '
                                             '0.002 -0.008 c 0 0.001 0.002 0.004 0.002 0.008 c 0 -0.002 0 '
                                             '-0.004 0.001 -0.004 c 0 0.001 0.001 0.002 0.001 0.004 c 0 '
                                             '-0.001 0 -0.002 0 -0.003 c 0.001 0 0.001 0.001 0.001 0.002 c '
                                             '0.001 0 0.001 -0.001 0.001 -0.002 c 0.001 0 0.001 0.001 0.001 '
                                             '0.002 E" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                             'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                             'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_figure8_h': {'ui': 'Ocho horizontal',
                                'presetID': 26,
                                'presetClass': 'path',
                                'presetSubtype': 0,
                                'accel': 50000,
                                'decel': 50000,
                                'fill': 'hold',
                                'grp': True,
                                'iterate': None,
                                'dur': 2000,
                                'com': [111, 0, None],
                                'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0 0.033 0.027 0.06 '
                                          '0.06 0.06 C 0.099 0.06 0.113 0.03 0.119 0.012 L 0.125 -0.012 C '
                                          '0.131 -0.03 0.146 -0.06 0.19 -0.06 C 0.218 -0.06 0.25 -0.033 0.25 '
                                          '0 C 0.25 0.033 0.218 0.06 0.19 0.06 C 0.146 0.06 0.131 0.03 0.125 '
                                          '0.012 L 0.119 -0.012 C 0.113 -0.03 0.099 -0.06 0.06 -0.06 C 0.027 '
                                          '-0.06 0 -0.033 0 0 Z" pathEditMode="relative" '
                                          'ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                          'fill="hold"/><p:tgtEl><p:spTgt '
                                          'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_square': {'ui': 'Cuadrado',
                             'presetID': 7,
                             'presetClass': 'path',
                             'presetSubtype': 0,
                             'accel': 50000,
                             'decel': 50000,
                             'fill': 'hold',
                             'grp': True,
                             'iterate': None,
                             'dur': 2000,
                             'com': [92, 0, None],
                             'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.25 0 L 0.25 0.25 L 0 '
                                       '0.25 L 0 0 Z" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                       'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                       'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.path_triangle': {'ui': 'Triangulo equilatero',
                               'presetID': 13,
                               'presetClass': 'path',
                               'presetSubtype': 0,
                               'accel': 50000,
                               'decel': 50000,
                               'fill': 'hold',
                               'grp': True,
                               'iterate': None,
                               'dur': 2000,
                               'com': [98, 0, None],
                               'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.125 0.216 L -0.125 '
                                         '0.216 L 0 0 Z" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                         'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                         'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.custom_lines': {'ui': 'Personalizada (lineas)',
                              'presetID': 63,
                              'presetClass': 'path',
                              'presetSubtype': 0,
                              'accel': 50000,
                              'decel': 50000,
                              'fill': 'hold',
                              'grp': True,
                              'iterate': None,
                              'dur': 2000,
                              'com': [149, 0, None],
                              'cuerpo': '<p:animMotion origin="layout" path="M 0 0 L 0.15 -0.2 L 0.3 0 L '
                                        '0.45 -0.2 E" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn '
                                        'dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'trayectoria.custom_curve': {'ui': 'Personalizada (curva Bezier)',
                              'presetID': 63,
                              'presetClass': 'path',
                              'presetSubtype': 0,
                              'accel': 50000,
                              'decel': 50000,
                              'fill': 'hold',
                              'grp': True,
                              'iterate': None,
                              'dur': 2000,
                              'com': [149, 0, None],
                              'cuerpo': '<p:animMotion origin="layout" path="M 0 0 C 0.1 -0.25 0.3 -0.25 0.4 '
                                        '0 E" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn dur="2000" '
                                        'fill="hold"/><p:tgtEl><p:spTgt '
                                        'spid="{SPID}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>'},
 'enfasis.grow_shrink': {'ui': 'Grow/Shrink (por defecto 150%)',
                         'presetID': 6,
                         'presetClass': 'emph',
                         'presetSubtype': 0,
                         'accel': None,
                         'decel': None,
                         'fill': 'hold',
                         'grp': True,
                         'iterate': None,
                         'dur': 2000,
                         'com': [59, 0, None],
                         'cuerpo': '<p:animScale><p:cBhvr><p:cTn dur="2000" fill="hold"/><p:tgtEl><p:spTgt '
                                   'spid="{SPID}"/></p:tgtEl></p:cBhvr><p:by x="150000" '
                                   'y="150000"/></p:animScale>'}}

# TRANSICIONES[clave] = {ns (p|p14|p15|p159), el (elemento dentro de <p:transition>),
#   fb (elemento del Fallback si difiere), com (EntryEffect que reporta COM), pedido}
TRANSICIONES = {'fade_smoothly': {'ns': 'p', 'el': '<p:fade/>', 'fb': None, 'com': 3849, 'pedido': 3849},
 'fade_through_black': {'ns': 'p', 'el': '<p:fade thruBlk="1"/>', 'fb': None, 'com': 1793, 'pedido': 1793},
 'cut': {'ns': 'p', 'el': '<p:cut/>', 'fb': None, 'com': 257, 'pedido': 257},
 'cut_through_black': {'ns': 'p', 'el': '<p:cut thruBlk="1"/>', 'fb': None, 'com': 258, 'pedido': 258},
 'push_up': {'ns': 'p', 'el': '<p:push dir="u"/>', 'fb': None, 'com': 3855, 'pedido': 3855},
 'push_left': {'ns': 'p', 'el': '<p:push/>', 'fb': None, 'com': 3853, 'pedido': 3853},
 'wipe_right': {'ns': 'p', 'el': '<p:wipe dir="r"/>', 'fb': None, 'com': 2819, 'pedido': 2819},
 'wipe_down': {'ns': 'p', 'el': '<p:wipe dir="d"/>', 'fb': None, 'com': 2820, 'pedido': 2820},
 'split_vertical_out': {'ns': 'p', 'el': '<p:split orient="vert"/>', 'fb': None, 'com': 3587, 'pedido': 3587},
 'split_horizontal_in': {'ns': 'p', 'el': '<p:split dir="in"/>', 'fb': None, 'com': 3586, 'pedido': 3586},
 'reveal_smooth_left': {'ns': 'p14', 'el': '<p14:reveal/>', 'fb': '<p:fade/>', 'com': 3894, 'pedido': 3894},
 'reveal_black_left': {'ns': 'p14',
                       'el': '<p14:reveal thruBlk="1"/>',
                       'fb': '<p:fade/>',
                       'com': 3896,
                       'pedido': 3896},
 'random_bars_h': {'ns': 'p', 'el': '<p:randomBar/>', 'fb': None, 'com': 2305, 'pedido': 2305},
 'shape_circle': {'ns': 'p', 'el': '<p:circle/>', 'fb': None, 'com': 3845, 'pedido': 3845},
 'shape_diamond': {'ns': 'p', 'el': '<p:diamond/>', 'fb': None, 'com': 3846, 'pedido': 3846},
 'shape_plus': {'ns': 'p', 'el': '<p:plus/>', 'fb': None, 'com': 3851, 'pedido': 3851},
 'uncover_left': {'ns': 'p', 'el': '<p:pull/>', 'fb': None, 'com': 2049, 'pedido': 2049},
 'cover_left': {'ns': 'p', 'el': '<p:cover/>', 'fb': None, 'com': 1281, 'pedido': 1281},
 'flash': {'ns': 'p14', 'el': '<p14:flash/>', 'fb': '<p:fade/>', 'com': 3909, 'pedido': 3909},
 'fall_over_left': {'ns': 'p15',
                    'el': '<p15:prstTrans prst="fallOver"/>',
                    'fb': '<p:fade/>',
                    'com': 3934,
                    'pedido': 3934},
 'drape_left': {'ns': 'p15',
                'el': '<p15:prstTrans prst="drape"/>',
                'fb': '<p:fade/>',
                'com': 3936,
                'pedido': 3936},
 'curtains': {'ns': 'p15',
              'el': '<p15:prstTrans prst="curtains"/>',
              'fb': '<p:fade/>',
              'com': 3938,
              'pedido': 3938},
 'wind_left': {'ns': 'p15',
               'el': '<p15:prstTrans prst="wind" invX="1"/>',
               'fb': '<p:fade/>',
               'com': 3939,
               'pedido': 3939},
 'prestige': {'ns': 'p15',
              'el': '<p15:prstTrans prst="prestige"/>',
              'fb': '<p:fade/>',
              'com': 3941,
              'pedido': 3941},
 'fracture': {'ns': 'p15',
              'el': '<p15:prstTrans prst="fracture"/>',
              'fb': '<p:fade/>',
              'com': 3942,
              'pedido': 3942},
 'crush': {'ns': 'p15',
           'el': '<p15:prstTrans prst="crush"/>',
           'fb': '<p:fade/>',
           'com': 3943,
           'pedido': 3943},
 'peel_off_left': {'ns': 'p15',
                   'el': '<p15:prstTrans prst="peelOff"/>',
                   'fb': '<p:fade/>',
                   'com': 3944,
                   'pedido': 3944},
 'page_curl_double_left': {'ns': 'p15',
                           'el': '<p15:prstTrans prst="pageCurlDouble"/>',
                           'fb': '<p:fade/>',
                           'com': 3948,
                           'pedido': 3948},
 'page_curl_single_left': {'ns': 'p15',
                           'el': '<p15:prstTrans prst="pageCurlSingle"/>',
                           'fb': '<p:fade/>',
                           'com': 3946,
                           'pedido': 3946},
 'airplane_left': {'ns': 'p15',
                   'el': '<p15:prstTrans prst="airplane" invX="1"/>',
                   'fb': '<p:fade/>',
                   'com': 3950,
                   'pedido': 3950},
 'origami_left': {'ns': 'p15',
                  'el': '<p15:prstTrans prst="origami" invX="1"/>',
                  'fb': '<p:fade/>',
                  'com': 3952,
                  'pedido': 3952},
 'dissolve': {'ns': 'p', 'el': '<p:dissolve/>', 'fb': None, 'com': 1537, 'pedido': 1537},
 'checkerboard_across': {'ns': 'p', 'el': '<p:checker/>', 'fb': None, 'com': 1025, 'pedido': 1025},
 'blinds_h': {'ns': 'p', 'el': '<p:blinds/>', 'fb': None, 'com': 769, 'pedido': 769},
 'clock_wedge': {'ns': 'p', 'el': '<p:wedge/>', 'fb': None, 'com': 3856, 'pedido': 3856},
 'ripple_center': {'ns': 'p14', 'el': '<p14:ripple/>', 'fb': '<p:fade/>', 'com': 3867, 'pedido': 3867},
 'honeycomb': {'ns': 'p14', 'el': '<p14:honeycomb/>', 'fb': '<p:fade/>', 'com': 3898, 'pedido': 3898},
 'glitter_diamond_left': {'ns': 'p14',
                          'el': '<p14:glitter dir="r"/>',
                          'fb': '<p:fade/>',
                          'com': 3872,
                          'pedido': 3872},
 'glitter_hexagon_left': {'ns': 'p14',
                          'el': '<p14:glitter dir="r" pattern="hexagon"/>',
                          'fb': '<p:fade/>',
                          'com': 3876,
                          'pedido': 3876},
 'vortex_left': {'ns': 'p14', 'el': '<p14:vortex/>', 'fb': '<p:fade/>', 'com': 3863, 'pedido': 3863},
 'shred_strips_in': {'ns': 'p14', 'el': '<p14:shred/>', 'fb': '<p:fade/>', 'com': 3910, 'pedido': 3910},
 'shred_rectangle_in': {'ns': 'p14',
                        'el': '<p14:shred pattern="rectangle"/>',
                        'fb': '<p:fade/>',
                        'com': 3912,
                        'pedido': 3912},
 'switch_left': {'ns': 'p14', 'el': '<p14:switch dir="l"/>', 'fb': '<p:fade/>', 'com': 3901, 'pedido': 3901},
 'flip_left': {'ns': 'p14', 'el': '<p14:flip dir="l"/>', 'fb': '<p:fade/>', 'com': 3905, 'pedido': 3905},
 'gallery_left': {'ns': 'p14',
                  'el': '<p14:gallery dir="l"/>',
                  'fb': '<p:fade/>',
                  'com': 3880,
                  'pedido': 3880},
 'cube_left': {'ns': 'p14', 'el': '<p14:prism/>', 'fb': '<p:fade/>', 'com': 3914, 'pedido': 3914},
 'doors_vertical': {'ns': 'p14',
                    'el': '<p14:doors dir="vert"/>',
                    'fb': '<p:fade/>',
                    'com': 3884,
                    'pedido': 3884},
 'box_left': {'ns': 'p14',
              'el': '<p14:prism isInverted="1"/>',
              'fb': '<p:fade/>',
              'com': 3922,
              'pedido': 3922},
 'comb_horizontal': {'ns': 'p', 'el': '<p:comb/>', 'fb': None, 'com': 3847, 'pedido': 3847},
 'warp_in': {'ns': 'p14', 'el': '<p14:warp dir="in"/>', 'fb': '<p:fade/>', 'com': 3888, 'pedido': 3888},
 'newsflash': {'ns': 'p', 'el': '<p:newsflash/>', 'fb': None, 'com': 3850, 'pedido': 3850},
 'rotate_left': {'ns': 'p14',
                 'el': '<p14:prism isContent="1"/>',
                 'fb': '<p:fade/>',
                 'com': 3918,
                 'pedido': 3918},
 'window_vertical': {'ns': 'p14',
                     'el': '<p14:window dir="vert"/>',
                     'fb': '<p:fade/>',
                     'com': 3886,
                     'pedido': 3886},
 'orbit_left': {'ns': 'p14',
                'el': '<p14:prism isContent="1" isInverted="1"/>',
                'fb': '<p:fade/>',
                'com': 3926,
                'pedido': 3926},
 'flythrough_in': {'ns': 'p14', 'el': '<p14:flythrough/>', 'fb': '<p:fade/>', 'com': 3890, 'pedido': 3890},
 'flythrough_in_bounce': {'ns': 'p14',
                          'el': '<p14:flythrough hasBounce="1"/>',
                          'fb': '<p:fade/>',
                          'com': 3892,
                          'pedido': 3892},
 'ferris_left': {'ns': 'p14', 'el': '<p14:ferris dir="l"/>', 'fb': '<p:fade/>', 'com': 3899, 'pedido': 3899},
 'conveyor_left': {'ns': 'p14',
                   'el': '<p14:conveyor dir="l"/>',
                   'fb': '<p:fade/>',
                   'com': 3882,
                   'pedido': 3882},
 'pan_up': {'ns': 'p14', 'el': '<p14:pan dir="u"/>', 'fb': '<p:fade/>', 'com': 3931, 'pedido': 3931},
 'wheel_1': {'ns': 'p', 'el': '<p:wheel spokes="1"/>', 'fb': None, 'com': 3857, 'pedido': 3857},
 'wheel_reverse_1': {'ns': 'p14',
                     'el': '<p14:wheelReverse spokes="1"/>',
                     'fb': '<p:fade/>',
                     'com': 3862,
                     'pedido': 3862},
 'strips_up_left': {'ns': 'p', 'el': '<p:strips/>', 'fb': None, 'com': 2565, 'pedido': 2561},
 'box_out': {'ns': 'p', 'el': '<p:zoom/>', 'fb': None, 'com': 3073, 'pedido': 3073},
 'random': {'ns': 'p', 'el': '<p:random/>', 'fb': None, 'com': 513, 'pedido': 513},
 'morph_objeto': {'ns': 'p159',
                  'el': '<p159:morph option="byObject"/>',
                  'fb': '<p:fade/>',
                  'com': 3954,
                  'pedido': 3954},
 'morph_palabra': {'ns': 'p159',
                   'el': '<p159:morph option="byWord"/>',
                   'fb': '<p:fade/>',
                   'com': 3955,
                   'pedido': 3955},
 'morph_caracter': {'ns': 'p159',
                    'el': '<p159:morph option="byChar"/>',
                    'fb': '<p:fade/>',
                    'com': 3956,
                    'pedido': 3956},
 'fall_over_right': {'ns': 'p15',
                     'el': '<p15:prstTrans prst="fallOver" invX="1"/>',
                     'fb': '<p:fade/>',
                     'com': 3935,
                     'pedido': 3935},
 'drape_right': {'ns': 'p15',
                 'el': '<p15:prstTrans prst="drape" invX="1"/>',
                 'fb': '<p:fade/>',
                 'com': 3937,
                 'pedido': 3937},
 'wind_right': {'ns': 'p15',
                'el': '<p15:prstTrans prst="wind"/>',
                'fb': '<p:fade/>',
                'com': 3940,
                'pedido': 3940},
 'peel_off_right': {'ns': 'p15',
                    'el': '<p15:prstTrans prst="peelOff" invX="1"/>',
                    'fb': '<p:fade/>',
                    'com': 3945,
                    'pedido': 3945},
 'page_curl_double_right': {'ns': 'p15',
                            'el': '<p15:prstTrans prst="pageCurlDouble" invX="1"/>',
                            'fb': '<p:fade/>',
                            'com': 3949,
                            'pedido': 3949},
 'airplane_right': {'ns': 'p15',
                    'el': '<p15:prstTrans prst="airplane"/>',
                    'fb': '<p:fade/>',
                    'com': 3951,
                    'pedido': 3951},
 'origami_right': {'ns': 'p15',
                   'el': '<p15:prstTrans prst="origami"/>',
                   'fb': '<p:fade/>',
                   'com': 3953,
                   'pedido': 3953},
 'vortex_right': {'ns': 'p14', 'el': '<p14:vortex dir="r"/>', 'fb': '<p:fade/>', 'com': 3865, 'pedido': 3865},
 'glitter_diamond_right': {'ns': 'p14',
                           'el': '<p14:glitter/>',
                           'fb': '<p:fade/>',
                           'com': 3874,
                           'pedido': 3874},
 'switch_right': {'ns': 'p14', 'el': '<p14:switch dir="r"/>', 'fb': '<p:fade/>', 'com': 3903, 'pedido': 3903},
 'flip_right': {'ns': 'p14', 'el': '<p14:flip dir="r"/>', 'fb': '<p:fade/>', 'com': 3907, 'pedido': 3907},
 'gallery_right': {'ns': 'p14',
                   'el': '<p14:gallery dir="r"/>',
                   'fb': '<p:fade/>',
                   'com': 3881,
                   'pedido': 3881},
 'conveyor_right': {'ns': 'p14',
                    'el': '<p14:conveyor dir="r"/>',
                    'fb': '<p:fade/>',
                    'com': 3883,
                    'pedido': 3883},
 'ferris_right': {'ns': 'p14', 'el': '<p14:ferris dir="r"/>', 'fb': '<p:fade/>', 'com': 3900, 'pedido': 3900},
 'pan_down': {'ns': 'p14', 'el': '<p14:pan dir="d"/>', 'fb': '<p:fade/>', 'com': 3933, 'pedido': 3933},
 'cube_right': {'ns': 'p14', 'el': '<p14:prism dir="r"/>', 'fb': '<p:fade/>', 'com': 3916, 'pedido': 3916},
 'push_right': {'ns': 'p', 'el': '<p:push dir="r"/>', 'fb': None, 'com': 3854, 'pedido': 3854},
 'cover_right': {'ns': 'p', 'el': '<p:cover dir="r"/>', 'fb': None, 'com': 1283, 'pedido': 1283},
 'uncover_right': {'ns': 'p', 'el': '<p:pull dir="r"/>', 'fb': None, 'com': 2051, 'pedido': 2051},
 'reveal_smooth_right': {'ns': 'p14',
                         'el': '<p14:reveal dir="r"/>',
                         'fb': '<p:fade/>',
                         'com': 3895,
                         'pedido': 3895},
 'doors_horizontal': {'ns': 'p14', 'el': '<p14:doors/>', 'fb': '<p:fade/>', 'com': 3885, 'pedido': 3885},
 'window_horizontal': {'ns': 'p14', 'el': '<p14:window/>', 'fb': '<p:fade/>', 'com': 3887, 'pedido': 3887},
 'warp_out': {'ns': 'p14', 'el': '<p14:warp/>', 'fb': '<p:fade/>', 'com': 3889, 'pedido': 3889},
 'shred_strips_out': {'ns': 'p14',
                      'el': '<p14:shred dir="out"/>',
                      'fb': '<p:fade/>',
                      'com': 3911,
                      'pedido': 3911},
 'ripple_left_up': {'ns': 'p14',
                    'el': '<p14:ripple dir="lu"/>',
                    'fb': '<p:fade/>',
                    'com': 3869,
                    'pedido': 3869},
 'flythrough_out': {'ns': 'p14',
                    'el': '<p14:flythrough dir="out"/>',
                    'fb': '<p:fade/>',
                    'com': 3891,
                    'pedido': 3891},
 'random_bars_v': {'ns': 'p', 'el': '<p:randomBar dir="vert"/>', 'fb': None, 'com': 2306, 'pedido': 2306},
 'split_horizontal_out': {'ns': 'p', 'el': '<p:split/>', 'fb': None, 'com': 3585, 'pedido': 3585},
 'split_vertical_in': {'ns': 'p',
                       'el': '<p:split orient="vert" dir="in"/>',
                       'fb': None,
                       'com': 3588,
                       'pedido': 3588},
 'wipe_left': {'ns': 'p', 'el': '<p:wipe/>', 'fb': None, 'com': 2817, 'pedido': 2817},
 'wipe_up': {'ns': 'p', 'el': '<p:wipe dir="u"/>', 'fb': None, 'com': 2818, 'pedido': 2818},
 'strips_down_right': {'ns': 'p', 'el': '<p:strips dir="rd"/>', 'fb': None, 'com': 2568, 'pedido': 2564},
 'wheel_4': {'ns': 'p', 'el': '<p:wheel/>', 'fb': None, 'com': 3860, 'pedido': 3860},
 'comb_vertical': {'ns': 'p', 'el': '<p:comb dir="vert"/>', 'fb': None, 'com': 3848, 'pedido': 3848},
 'box_in': {'ns': 'p', 'el': '<p:zoom dir="in"/>', 'fb': None, 'com': 3074, 'pedido': 3074},
 'fade': {'ns': 'p', 'el': '<p:fade/>', 'fb': None, 'com': 3849, 'pedido': 3849},
 'morph': {'ns': 'p159',
           'el': '<p159:morph option="byObject"/>',
           'fb': '<p:fade/>',
           'com': 3954,
           'pedido': 3954}}
# === FIN DATOS GENERADOS ===
