"""Construye la presentación de defensa de la tesis EduFEM (PowerPoint editable).

Uso:
    python build_deck.py guion.json salida.pptx

El guion es un JSON con la lista de diapositivas; cada una declara un `layout`
del vocabulario de abajo y los campos que ese layout necesita. Todo lo que se
dibuja son formas, cuadros de texto, tablas e imágenes nativas de PowerPoint,
así que la presentación queda íntegramente editable.

Layouts: portada · seccion · bullets · dos_columnas · imagen · tabla · kpi ·
flujo · comparacion · ecuaciones · demo · cierre
"""
from __future__ import annotations

import json
import math
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Sistema visual
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1B, 0x2A, 0x4A)       # títulos, cabeceras de tabla, chips
NAVY_DEEP = RGBColor(0x0E, 0x18, 0x2E)  # fondo de sección y cierre
NAVY_SOFT = RGBColor(0x2C, 0x3E, 0x63)  # chips «ya vistos» sobre fondo oscuro
SLATE = RGBColor(0x3E, 0x5A, 0x86)      # tercera tarjeta de comparación
ACCENT = RGBColor(0xF2, 0x8C, 0x28)     # naranja: acentos, remates, hilo actual
ACCENT_TINT = RGBColor(0xFD, 0xF1, 0xE2)  # fondo de la columna/fila destacada
INK = RGBColor(0x1F, 0x29, 0x37)        # texto principal
MUTED = RGBColor(0x5B, 0x65, 0x75)      # texto secundario
CARD = RGBColor(0xF6, 0xF8, 0xFC)       # superficie de tarjetas y filas pares
LINE = RGBColor(0xDC, 0xE3, 0xEE)       # bordes y hairlines
TINT = RGBColor(0xEA, 0xF0, 0xFA)       # fondo de la barra de remate
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ON_DARK = RGBColor(0xC7, 0xD2, 0xE8)    # texto secundario sobre navy
Q4_BLUE = RGBColor(0x1F, 0x77, 0xB4)    # mismo azul que las curvas de la tesis
Q9_ORANGE = RGBColor(0xFF, 0x7F, 0x0E)  # mismo naranja que las curvas de la tesis

FONT = "Calibri"
MATH = "Cambria Math"

SLIDE_W = 13.333
SLIDE_H = 7.5
MX = 0.62           # margen lateral
KICKER_Y = 0.34     # rótulo del bloque, encima del título
TITLE_Y = 0.64
RULE_Y = 1.50       # acento naranja bajo el título
CONTENT_Y = 1.72
REMATE_Y = 6.30
REMATE_H = 0.58
FOOTER_Y = 7.04

ROOT = os.path.dirname(os.path.abspath(__file__))
FIGURAS = os.path.normpath(os.path.join(ROOT, "..", "figuras"))

AVISOS: list[str] = []


def aviso(msg: str) -> None:
    AVISOS.append(msg)
    print("  ! " + msg)


def hexa(color: RGBColor) -> str:
    return "%02X%02X%02X" % (color[0], color[1], color[2])


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------
def _set_lang(run) -> None:
    run._r.get_or_add_rPr().set("lang", "es-ES")


def _style_run(run, size, bold=False, color=INK, font=FONT, italic=False, spacing=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    _set_lang(run)
    if spacing:  # separación entre letras, en puntos (para rótulos en versalita)
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def add_text(slide, x, y, w, h, text, size, bold=False, color=INK, font=FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
             line_spacing=1.05, margins=(0.05, 0.03, 0.05, 0.03), spacing=None):
    """Cuadro de texto simple. `text` admite '\n' para varios párrafos."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        _style_run(r, size, bold, color, font, italic, spacing)
    return tb


def _add_bullet_xml(paragraph, color: RGBColor, level: int = 0, char: str = "•"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.38 + 0.38 * level))))
    pPr.set("indent", str(-int(Inches(0.38))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buSzPct", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(parse_xml(f'<a:buClr {nsdecls("a")}><a:srgbClr val="{hexa(color)}"/></a:buClr>'))
    pPr.append(parse_xml(f'<a:buSzPct {nsdecls("a")} val="100000"/>'))
    pPr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="Arial"/>'))
    pPr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="{char}"/>'))


def est_lines(text: str, size_pt: float, width_in: float, char_w: float = 0.50) -> int:
    """Estimación conservadora de líneas que ocupa un texto en Calibri."""
    chars_per_line = max(1, int(width_in * 72 / (size_pt * char_w)))
    n = 0
    for para in str(text).split("\n"):
        n += max(1, math.ceil(len(para) / chars_per_line))
    return n


def fit_size(items, width_in, height_in, start, floor, line_h=1.18, gap_pt=10, indent_in=0.0):
    """Mayor tamaño de fuente (de `start` hacia `floor`, de a 1 pt) con el que cabe la lista."""
    for s in range(int(start), int(floor) - 1, -1):
        total = sum(est_lines(t, s, width_in - indent_in) for t in items) * s * line_h / 72
        total += max(0, len(items) - 1) * gap_pt / 72
        if total <= height_in:
            return s
    return int(floor)


def add_bullets(slide, x, y, w, h, items, size=24, color=INK, bullet_color=ACCENT,
                gap_pt=10, min_size=20, anchor=MSO_ANCHOR.TOP, bold_prefix=True):
    """Lista con viñetas reales (editables). Si `bold_prefix`, un 'Texto: resto' resalta 'Texto:'."""
    items = [str(i) for i in items if str(i).strip()]
    if not items:
        return None
    s = fit_size(items, w, h, size, min_size, gap_pt=gap_pt, indent_in=0.38)
    if s < size:
        aviso(f"viñetas reducidas a {s} pt para que quepan: {items[0][:40]}…")
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(gap_pt)
        _add_bullet_xml(p, bullet_color)
        head, sep, tail = item.partition(":")
        if bold_prefix and sep and 0 < len(head) <= 32 and tail.strip():
            r1 = p.add_run()
            r1.text = head + ":"
            _style_run(r1, s, True, NAVY)
            r2 = p.add_run()
            r2.text = tail
            _style_run(r2, s, False, color)
        else:
            r = p.add_run()
            r.text = item
            _style_run(r, s, False, color)
    return tb


def add_rect(slide, x, y, w, h, fill=CARD, line=None, rounded=False, radius=0.06, line_w=0.75):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.text_frame.text = ""
    return shp


def add_card(slide, x, y, w, h, fill=WHITE, border=LINE, radius=0.05, bar=None, bar_h=0.10,
             bar_side="top"):
    """Tarjeta: fondo claro, borde fino y —si se pide— una barra de color de acento."""
    card = add_rect(slide, x, y, w, h, fill=fill, line=border, rounded=True, radius=radius)
    if bar is not None:
        if bar_side == "top":
            add_rect(slide, x + 0.02, y + 0.02, w - 0.04, bar_h, fill=bar)
        else:
            add_rect(slide, x + 0.02, y + 0.02, bar_h, h - 0.04, fill=bar)
    return card


def add_oval(slide, x, y, d, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.0)
    shp.text_frame.text = ""
    return shp


def shape_text(shp, text, size, bold=False, color=WHITE, font=FONT, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, margins=(0.08, 0.04, 0.08, 0.04), line_spacing=1.0):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        _style_run(r, size, bold, color, font)


def add_picture_fit(slide, path, x, y, w, h, border=True, align="center"):
    """Inserta la imagen manteniendo la proporción, centrada en la caja (x, y, w, h)."""
    if not os.path.isfile(path):
        aviso(f"imagen no encontrada: {path}")
        ph = add_rect(slide, x, y, w, h, fill=CARD, line=LINE)
        shape_text(ph, "[imagen no encontrada]\n" + os.path.basename(path), 16, color=MUTED)
        return ph
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    if w / h > ar:
        ph_h = h
        ph_w = h * ar
    else:
        ph_w = w
        ph_h = w / ar
    px = x + (w - ph_w) / 2 if align == "center" else x
    py = y + (h - ph_h) / 2
    pic = slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(ph_w), Inches(ph_h))
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def add_picture_card(slide, path, x, y, w, h, pad=0.13, radius=0.04):
    """Imagen centrada en la caja, con la tarjeta ajustada a su tamaño real (sin franjas muertas)."""
    if not os.path.isfile(path):
        return add_picture_fit(slide, path, x, y, w, h)
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    bw, bh = w - 2 * pad, h - 2 * pad
    pw, ph = (bh * ar, bh) if bw / bh > ar else (bw, bw / ar)
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    add_rect(slide, px - pad, py - pad, pw + 2 * pad, ph + 2 * pad, fill=CARD, line=LINE,
             rounded=True, radius=radius)
    return slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))


def mesh_motif(slide, x, y, w, h, nx, ny, line_color, node_color=None, lw=0.012, nd=0.055):
    """Malla de cuadriláteros como motivo de marca (líneas y nodos, todo editable)."""
    for j in range(ny + 1):
        yy = y + h * j / ny
        add_rect(slide, x, yy, w, lw, fill=line_color)
    for i in range(nx + 1):
        xx = x + w * i / nx
        add_rect(slide, xx, y, lw, h, fill=line_color)
    if node_color is not None:
        for j in range(ny + 1):
            for i in range(nx + 1):
                add_oval(slide, x + w * i / nx - nd / 2, y + h * j / ny - nd / 2, nd, node_color)


def set_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = str(text)


def add_slide_number(slide, x, y, w, h, color=MUTED, size=12, align=PP_ALIGN.RIGHT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    fld = parse_xml(
        f'<a:fld {nsdecls("a")} id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
        f'<a:rPr lang="es-ES" sz="{int(size * 100)}" dirty="0"><a:solidFill><a:srgbClr val="{hexa(color)}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}"/></a:rPr><a:t>1</a:t></a:fld>'
    )
    p._p.append(fld)
    return tb


# ---------------------------------------------------------------------------
# Marco común de las diapositivas de contenido
# ---------------------------------------------------------------------------
class Deck:
    def __init__(self, guion: dict):
        self.guion = guion
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.slides_data = guion["diapositivas"]
        # Bloques del hilo conductor: los títulos de las diapositivas "seccion", en orden.
        self.bloques = [d for d in self.slides_data if d.get("layout") == "seccion"]
        self.meta = guion.get("meta", {})

    # ---- piezas comunes -------------------------------------------------
    def new_slide(self, bg=None):
        s = self.prs.slides.add_slide(self.blank)
        if bg is not None:
            fill = s.background.fill
            fill.solid()
            fill.fore_color.rgb = bg
        return s

    def header(self, s, d, title_size=32):
        """Rótulo del bloque + título + acento, con el avance del hilo conductor a la derecha."""
        bloque = d.get("bloque", "")
        if bloque:
            idx = self.bloque_index(bloque)
            label = f"{idx:02d}   {bloque.upper()}" if idx else bloque.upper()
            add_rect(s, MX, KICKER_Y + 0.055, 0.055, 0.19, fill=ACCENT)
            add_text(s, MX + 0.15, KICKER_Y, 7.4, 0.30, label, 12, bold=True, color=ACCENT,
                     anchor=MSO_ANCHOR.MIDDLE, spacing=0.8)
            self.progreso(s, d)
        titulo = d.get("titulo", "")
        size = title_size if len(titulo) <= 46 else (29 if len(titulo) <= 58 else 26)
        add_text(s, MX - 0.02, TITLE_Y, SLIDE_W - 2 * MX, 0.82, titulo, size, bold=True, color=NAVY,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, MX, RULE_Y, 1.15, 0.055, fill=ACCENT)

    def progreso(self, s, d):
        """Puntos del hilo conductor arriba a la derecha: el bloque actual en naranja."""
        n = len(self.bloques)
        if not n:
            return
        idx = self.bloque_index(d.get("bloque", ""))
        dot, gap = 0.125, 0.115
        total = n * dot + (n - 1) * gap
        x = SLIDE_W - MX - total
        y = KICKER_Y + 0.085
        for i in range(n):
            actual = (i + 1) == idx
            if actual:
                add_oval(s, x - 0.035, y - 0.035, dot + 0.07, ACCENT)
            else:
                add_oval(s, x, y, dot, NAVY if (i + 1) < idx else LINE)
            x += dot + gap

    def bloque_index(self, bloque):
        for i, b in enumerate(self.bloques, start=1):
            if b.get("bloque") == bloque:
                return i
        return 0

    def footer(self, s, d):
        add_rect(s, MX, FOOTER_Y - 0.08, SLIDE_W - 2 * MX, 0.008, fill=LINE)
        add_text(s, MX, FOOTER_Y, 5.5, 0.32, self.meta.get("pie", "EduFEM · Defensa de tesis · Ingeniería Civil · UATF"),
                 12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        fuente = d.get("fuente", "")
        if fuente:
            if len(fuente) > 72:
                fuente = fuente[:70].rstrip(" ,;") + "…"
            add_text(s, 4.4, FOOTER_Y, 5.9, 0.32, "Tesis: " + fuente, 12 if len(fuente) <= 58 else 11, color=MUTED,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_slide_number(s, SLIDE_W - MX - 1.2, FOOTER_Y, 1.2, 0.32)

    @staticmethod
    def remate_geom(text):
        """(y, h, tamaño) de la barra de remate: crece a dos líneas si el texto lo pide."""
        if not text:
            return REMATE_Y, 0.0, 22
        w = SLIDE_W - 2 * MX - 0.35
        size = 22 if len(text) <= 95 else 20
        lines = est_lines(text, size, w, char_w=0.53)
        if lines <= 1:
            return REMATE_Y, REMATE_H, size
        h = 0.92
        return REMATE_Y + REMATE_H - h, h, size

    def remate(self, s, text):
        if not text:
            return
        y, h, size = self.remate_geom(text)
        x, w = MX, SLIDE_W - 2 * MX
        add_rect(s, x, y, w, h, fill=TINT)
        add_rect(s, x, y, 0.085, h, fill=ACCENT)
        add_text(s, x + 0.25, y, w - 0.35, h, text, size, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
                 line_spacing=1.0)

    def content_box(self, d):
        """(y, h) del área de contenido según haya remate o no, y de qué alto."""
        y, h, _ = self.remate_geom(d.get("remate"))
        if h:
            return CONTENT_Y, y - CONTENT_Y - 0.12
        return CONTENT_Y, REMATE_Y + REMATE_H - CONTENT_Y

    def subtitulo(self, s, d, y, h):
        """Línea de aclaración bajo el título; devuelve el (y, h) restante."""
        sub = d.get("subtitulo")
        if not sub:
            return y, h
        size = 18 if len(sub) <= 110 else 16
        lines = est_lines(sub, size, SLIDE_W - 2 * MX, char_w=0.50)
        sh = min(0.86, lines * size * 1.15 / 72 + 0.10)
        add_text(s, MX, y, SLIDE_W - 2 * MX, sh, sub, size, color=MUTED, italic=True, line_spacing=1.08)
        return y + sh + 0.12, h - sh - 0.12

    def fig(self, name):
        name = str(name or "").strip()
        if not name:
            return ""
        if not name.lower().endswith(".png"):
            name += ".png"
        return os.path.join(FIGURAS, name)

    # ---- layouts -----------------------------------------------------------
    def portada(self, d):
        s = self.new_slide(WHITE)
        m = self.meta
        panel_x = 9.05
        # Panel derecho: escudo, logo y marca sobre navy, con malla de fondo
        add_rect(s, panel_x, 0, SLIDE_W - panel_x, SLIDE_H, fill=NAVY_DEEP)
        add_rect(s, panel_x, 0, 0.06, SLIDE_H, fill=ACCENT)
        add_picture_fit(s, self.fig("logo_universidad"), panel_x + 0.75, 0.75, 2.8, 2.35, border=False)
        add_rect(s, panel_x + 1.55, 3.42, 1.2, 0.045, fill=ACCENT)
        add_picture_fit(s, self.fig("fig_logo_edufem"), panel_x + 1.05, 3.85, 2.2, 1.95, border=False)
        add_text(s, panel_x, 5.85, SLIDE_W - panel_x, 0.5, "EduFEM", 26, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER)
        add_text(s, panel_x + 0.35, 6.32, SLIDE_W - panel_x - 0.7, 0.8,
                 "Software educativo de elementos finitos", 13, color=ON_DARK, align=PP_ALIGN.CENTER)
        # Banda superior y columna izquierda
        add_rect(s, 0, 0, panel_x, 0.30, fill=NAVY)
        add_rect(s, 0, 0.30, panel_x, 0.05, fill=ACCENT)
        add_text(s, 0.8, 0.62, 7.9, 0.34, "DEFENSA DE TESIS", 13, bold=True, color=ACCENT,
                 anchor=MSO_ANCHOR.MIDDLE, spacing=1.4)
        add_text(s, 0.8, 1.00, 7.9, 0.42, m.get("universidad", "").upper(), 19, bold=True, color=NAVY)
        add_text(s, 0.8, 1.42, 7.9, 0.34, m.get("facultad", "").upper(), 14, bold=True, color=MUTED)
        add_text(s, 0.8, 1.72, 7.9, 0.34, m.get("carrera", "").upper(), 14, bold=True, color=MUTED)
        add_rect(s, 0.82, 2.22, 1.15, 0.055, fill=ACCENT)
        titulo = m.get("titulo", d.get("titulo", ""))
        tsize = 27 if len(titulo) <= 150 else 24
        add_text(s, 0.8, 2.45, 7.9, 2.15, titulo.upper(), tsize, bold=True, color=NAVY,
                 anchor=MSO_ANCHOR.TOP, line_spacing=1.08)
        add_text(s, 0.8, 4.72, 7.9, 0.42,
                 "Tesis para optar al " + m.get("grado", "Título de Licenciatura en Ingeniería Civil"),
                 17, color=MUTED)
        add_rect(s, 0.82, 5.25, 7.75, 0.008, fill=LINE)
        add_text(s, 0.8, 5.42, 7.9, 0.45, m.get("autor", ""), 23, bold=True, color=NAVY)
        if m.get("director"):
            add_text(s, 0.8, 5.92, 7.9, 0.40, "Tutor:  " + m["director"], 16, color=INK)
        add_text(s, 0.8, 6.52, 7.9, 0.42,
                 f"{m.get('ciudad', 'Potosí – Bolivia')}   ·   {m.get('fecha', '')}", 16, color=MUTED)
        set_notes(s, d.get("notas"))
        return s

    def seccion(self, d):
        s = self.new_slide(NAVY_DEEP)
        idx = self.bloque_index(d.get("bloque", "")) or (self.bloques.index(d) + 1 if d in self.bloques else 0)
        mesh_motif(s, 8.95, 1.05, 3.95, 2.95, 4, 3, NAVY_SOFT, node_color=NAVY_SOFT, lw=0.012, nd=0.085)
        add_text(s, 0.8, 1.55, 2.4, 2.0, f"{idx:02d}" if idx else "", 96, bold=True, color=ACCENT,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, 3.35, 1.85, 0.06, 1.45, fill=ACCENT)
        add_text(s, 3.7, 1.55, 8.9, 1.3, d.get("titulo", ""), 44, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, 3.7, 2.85, 8.9, 1.4, d.get("subtitulo", ""), 21, color=ON_DARK, anchor=MSO_ANCHOR.TOP,
                 line_spacing=1.15)
        if d.get("bullets"):
            add_bullets(s, 3.7, 3.7, 8.9, 2.2, d["bullets"], size=22, color=ON_DARK, bullet_color=ACCENT,
                        min_size=18, bold_prefix=False)
        # Hilo conductor: chips de todos los bloques, con el actual resaltado
        n = len(self.bloques)
        if n:
            gap = 0.12
            cw = (SLIDE_W - 2 * MX - gap * (n - 1)) / n
            y = 6.35
            for i, b in enumerate(self.bloques):
                x = MX + i * (cw + gap)
                actual = (b is d)
                pasado = self.bloques.index(b) < (self.bloques.index(d) if d in self.bloques else 0)
                fill = ACCENT if actual else (NAVY_SOFT if pasado else None)
                chip = add_rect(s, x, y, cw, 0.54, fill=fill, line=None if fill else NAVY_SOFT,
                                rounded=True, radius=0.28)
                shape_text(chip, f"{i + 1}. {b.get('bloque', b.get('titulo', ''))}", 12,
                           bold=actual, color=WHITE if (actual or pasado) else ON_DARK,
                           margins=(0.05, 0.02, 0.05, 0.02))
        set_notes(s, d.get("notas"))
        return s

    def bullets(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        items = d.get("bullets", [])
        add_bullets(s, MX + 0.1, y + 0.1, SLIDE_W - 2 * MX - 0.2, h - 0.1, items, size=26 if len(items) <= 4 else 24,
                    gap_pt=14 if len(items) <= 4 else 10,
                    anchor=MSO_ANCHOR.MIDDLE if not d.get("subtitulo") else MSO_ANCHOR.TOP)
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def dos_columnas(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        img = self.fig(d.get("imagen"))
        left_w = 5.75 if img else 5.9
        add_bullets(s, MX + 0.05, y + 0.05, left_w, h - 0.1, d.get("bullets", []), size=24, gap_pt=12,
                    min_size=19, anchor=MSO_ANCHOR.MIDDLE)
        rx = MX + left_w + 0.45
        rw = SLIDE_W - MX - rx
        if img:
            cap = d.get("caption", "")
            cap_h = 0.55 if cap else 0
            pic = add_picture_card(s, img, rx, y, rw, h - cap_h - 0.02)
            if cap:
                bottom = (pic.top + pic.height) / 914400
                add_text(s, rx, bottom + 0.05, rw, cap_h, cap, 15, color=MUTED,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, italic=True)
        elif d.get("columna_derecha"):
            add_card(s, rx, y, rw, h, fill=CARD, bar=ACCENT, bar_h=0.09)
            add_bullets(s, rx + 0.2, y + 0.3, rw - 0.4, h - 0.5, d["columna_derecha"], size=22, gap_pt=10,
                        anchor=MSO_ANCHOR.MIDDLE, bullet_color=NAVY)
        elif d.get("kpis"):
            self._kpi_cards(s, d["kpis"], rx, y, rw, h, cols=1 if len(d["kpis"]) <= 3 else 2)
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def imagen(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        cap = d.get("caption", "")
        cap_h = 0.5 if cap else 0
        pic = add_picture_card(s, self.fig(d.get("imagen")), MX, y, SLIDE_W - 2 * MX, h - cap_h - 0.02)
        if cap:
            bottom = (pic.top + pic.height) / 914400
            add_text(s, MX, bottom + 0.06, SLIDE_W - 2 * MX, cap_h, cap, 16,
                     color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, italic=True)
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def tabla(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        t = d.get("tabla") or {}
        cab = [str(c) for c in t.get("cabeceras", [])]
        filas = [[str(c) for c in f] for f in t.get("filas", [])]
        destacar_col = t.get("destacar_col")     # índice de columna a resaltar (0 = la primera)
        destacar_fila = t.get("destacar_fila")   # índice de fila de datos a resaltar (0 = la primera)
        ncols = max(len(cab), max((len(f) for f in filas), default=0))
        nrows = len(filas) + 1
        if ncols == 0 or nrows == 1:
            aviso(f"tabla vacía en la diapositiva {d.get('n')}")
            return s
        size = 22 if nrows <= 5 else (20 if nrows <= 7 else 18)
        # ancho por columna proporcional al contenido, con un mínimo que no parte la palabra más larga
        table_w = SLIDE_W - 2 * MX
        cols_txt = []
        for c in range(ncols):
            cols_txt.append([cab[c] if c < len(cab) else ""] + [f[c] if c < len(f) else "" for f in filas])
        lens = [max(4, max(len(x) for x in col)) for col in cols_txt]
        longest_word = [max((len(w) for x in col for w in x.split()), default=4) for col in cols_txt]

        def widths_for(sz):
            tot = sum(lens)
            ws = []
            for L, lw in zip(lens, longest_word):
                minimum = max(1.1, lw * sz * 0.55 / 72 + 0.25)
                ws.append(max(minimum, table_w * L / tot))
            k = table_w / sum(ws)
            return [w * k for w in ws]

        def row_heights_for(sz, ws):
            hs = []
            for r in range(nrows):
                lines = 1
                for c in range(ncols):
                    txt = cols_txt[c][r]
                    lines = max(lines, est_lines(txt, sz, ws[c] - 0.16, char_w=0.52))
                hs.append(lines * sz * 1.2 / 72 + 0.14)
            return hs

        while True:
            widths = widths_for(size)
            row_hs = row_heights_for(size, widths)
            if sum(row_hs) <= h or size <= 13:
                break
            size -= 1
        if size < (22 if nrows <= 5 else (20 if nrows <= 7 else 18)):
            aviso(f"tabla de la diapositiva {d.get('n')} reducida a {size} pt")
        # si sobra alto, repartilo (filas más cómodas, tope 0,62 in)
        extra = max(0.0, h - sum(row_hs))
        row_hs = [min(0.62, rh + extra / nrows) for rh in row_hs]
        table_h = sum(row_hs)
        if table_h > h:
            aviso(f"la tabla de la diapositiva {d.get('n')} no entra en el area de contenido")
        ty = max(y, y + (h - table_h) / 2)
        gt = s.shapes.add_table(nrows, ncols, Inches(MX), Inches(ty), Inches(table_w), Inches(table_h))
        tbl = gt.table
        # sin estilo predefinido: pintamos las celdas nosotros
        tblPr = tbl._tbl.tblPr
        for attr in ("bandRow", "firstRow"):
            tblPr.set(attr, "0")
        for c, wdt in enumerate(widths):
            tbl.columns[c].width = Inches(wdt)
        for r in range(nrows):
            tbl.rows[r].height = Inches(row_hs[r])
            for c in range(ncols):
                cell = tbl.cell(r, c)
                if r == 0:
                    txt = cab[c] if c < len(cab) else ""
                else:
                    f = filas[r - 1]
                    txt = f[c] if c < len(f) else ""
                marcada = (destacar_col is not None and c == destacar_col) or \
                          (destacar_fila is not None and r == destacar_fila + 1)
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = NAVY
                elif marcada:
                    cell.fill.fore_color.rgb = ACCENT_TINT
                else:
                    cell.fill.fore_color.rgb = WHITE if r % 2 else CARD
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                run = p.add_run()
                run.text = txt
                _style_run(run, size, bold=(r == 0 or c == 0 or marcada),
                           color=WHITE if r == 0 else INK)
        # línea de acento bajo la cabecera (el borde inferior lo pone la propia tabla)
        add_rect(s, MX, ty + row_hs[0] - 0.035, table_w, 0.045, fill=ACCENT)
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def _kpi_cards(self, s, kpis, x, y, w, h, cols=None):
        """Tarjetas de cifra: un solo tamaño y una sola línea de base para todas (sin huecos muertos)."""
        n = len(kpis)
        cols = cols or n
        rows = math.ceil(n / cols)
        gap = 0.28
        cw = (w - gap * (cols - 1)) / cols
        ch = min(3.3, (h - gap * (rows - 1)) / rows)
        y0 = y + (h - (ch * rows + gap * (rows - 1))) / 2
        # tamaños comunes: manda el valor más largo y la etiqueta más larga
        vs = 44
        for k in kpis:
            valor = str(k.get("valor", ""))
            v = 44 if len(valor) <= 7 else (36 if len(valor) <= 11 else (30 if len(valor) <= 14 else 24))
            while v > 18 and est_lines(valor, v, cw - 0.5, char_w=0.56) > 1:
                v -= 2
            vs = min(vs, v)
        es = 18 if all(len(str(k.get("etiqueta", ""))) <= 70 for k in kpis) else 16
        vh = vs * 1.25 / 72
        while True:  # el bloque acento + valor + etiqueta tiene que caber en la tarjeta
            lab_lines = max(est_lines(str(k.get("etiqueta", "")), es, cw - 0.5, char_w=0.50) for k in kpis)
            lh = lab_lines * es * 1.16 / 72
            bloque = 0.26 + vh + 0.10 + lh
            if bloque <= ch - 0.30 or es <= 13:
                break
            es -= 1
        dy = max(0.16, (ch - bloque) / 2)
        for i, k in enumerate(kpis):
            r, c = divmod(i, cols)
            cx = x + c * (cw + gap)
            cy = y0 + r * (ch + gap)
            add_card(s, cx, cy, cw, ch, fill=CARD, border=LINE, radius=0.06)
            add_rect(s, cx + 0.28, cy + dy, 0.62, 0.06, fill=ACCENT)
            add_text(s, cx + 0.25, cy + dy + 0.24, cw - 0.5, vh + 0.12, str(k.get("valor", "")), vs,
                     bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP, line_spacing=1.0)
            add_text(s, cx + 0.25, cy + dy + 0.24 + vh + 0.12, cw - 0.5, lh + 0.2, str(k.get("etiqueta", "")),
                     es, color=MUTED, anchor=MSO_ANCHOR.TOP, line_spacing=1.12)

    def kpi(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        kpis = d.get("kpis", [])
        cols = len(kpis) if len(kpis) <= 4 else 3
        self._kpi_cards(s, kpis, MX, y, SLIDE_W - 2 * MX, h, cols=cols)
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def flujo(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        pasos = d.get("pasos", [])
        n = len(pasos)
        if not n:
            aviso(f"flujo sin pasos en la diapositiva {d.get('n')}")
            return s
        per_row = n if n <= 5 else math.ceil(n / 2)
        rows = math.ceil(n / per_row)
        arrow_w = 0.38
        gap = 0.11
        cw = (SLIDE_W - 2 * MX - (per_row - 1) * (arrow_w + 2 * gap)) / per_row
        # alto del chip: lo que piden rótulo y sublínea, sin dejar hueco entre ambos
        ls_all, ss_all = [], []
        for p in pasos:
            label = str(p.get("etiqueta", ""))
            sub = str(p.get("sub", "") or "")
            ls_all.append(20 if est_lines(label, 20, cw - 0.28, char_w=0.55) <= 1 else 17)
            ss_all.append(13 if est_lines(sub, 13, cw - 0.28, char_w=0.50) <= 2 else 11)
        ls = min(ls_all)
        ss = min(ss_all)
        lab_lines = max(est_lines(str(p.get("etiqueta", "")), ls, cw - 0.28, char_w=0.55) for p in pasos)
        sub_lines = max(est_lines(str(p.get("sub", "") or ""), ss, cw - 0.28, char_w=0.50)
                        for p in pasos) if any(p.get("sub") for p in pasos) else 0
        lab_h = lab_lines * ls * 1.15 / 72
        sub_h = sub_lines * ss * 1.15 / 72 if sub_lines else 0
        texto_h = lab_h + (0.10 + sub_h if sub_h else 0)
        # la tarjeta crece hasta ocupar el alto disponible, con el texto centrado dentro
        row_gap = 0.5 if rows > 1 else 0.0
        ch = min(2.3 if rows > 1 else 2.7, max(texto_h + 0.56, (h - row_gap * (rows - 1)) / rows - 0.25))
        total_h = rows * ch + (rows - 1) * row_gap
        y0 = y + (h - total_h) / 2
        for i, p in enumerate(pasos):
            r, c = divmod(i, per_row)
            cx = MX + c * (cw + arrow_w + 2 * gap)
            cy = y0 + r * (ch + row_gap)
            add_card(s, cx, cy, cw, ch, fill=CARD, border=LINE, radius=0.09, bar=NAVY, bar_h=0.11)
            label = str(p.get("etiqueta", ""))
            sub = str(p.get("sub", "") or "")
            ty0 = cy + max(0.26, 0.11 + (ch - 0.11 - texto_h) / 2)
            add_text(s, cx, ty0, cw, lab_h + 0.12, label, ls, bold=True, color=NAVY,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, line_spacing=1.05,
                     margins=(0.10, 0.02, 0.10, 0.02))
            if sub:
                add_text(s, cx, ty0 + lab_h + 0.08, cw, sub_h + 0.14, sub, ss, color=MUTED,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, line_spacing=1.05,
                         margins=(0.10, 0.02, 0.10, 0.02))
            if c < per_row - 1 and i < n - 1:
                ax = cx + cw + gap
                arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(cy + ch / 2 - 0.13),
                                         Inches(arrow_w), Inches(0.26))
                arr.shadow.inherit = False
                arr.fill.solid()
                arr.fill.fore_color.rgb = ACCENT
                arr.line.fill.background()
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def comparacion(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        tarjetas = d.get("tarjetas", [])[:3]
        paleta = {"navy": NAVY, "accent": ACCENT, "slate": SLATE, "q4": Q4_BLUE, "q9": Q9_ORANGE}
        por_defecto = [NAVY, ACCENT, SLATE]
        n = max(1, len(tarjetas))
        gap = 0.4 if n == 2 else 0.32
        cw = (SLIDE_W - 2 * MX - gap * (n - 1)) / n
        head_h = 0.68
        for i, t in enumerate(tarjetas):
            cx = MX + i * (cw + gap)
            titulo = str(t.get("titulo", ""))
            up = titulo.upper()
            color = paleta.get(str(t.get("color", "")).lower())
            if color is None:
                color = Q4_BLUE if up.startswith("Q4") else (Q9_ORANGE if up.startswith("Q9") else por_defecto[i])
            add_rect(s, cx, y, cw, h, fill=CARD, line=LINE, rounded=True, radius=0.05)
            add_rect(s, cx, y, cw, head_h, fill=color, rounded=True, radius=0.12)
            add_rect(s, cx, y + head_h - 0.30, cw, 0.30, fill=color)  # cuadra la base de la cabecera
            ts = 23 if est_lines(titulo, 23, cw - 0.3, char_w=0.55) <= 1 else (20 if n == 2 else 18)
            add_text(s, cx, y, cw, head_h, titulo, ts, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
            add_bullets(s, cx + 0.22, y + head_h + 0.22, cw - 0.44, h - head_h - 0.36, t.get("lineas", []),
                        size=22 if n == 2 else 19, gap_pt=9, min_size=15, bullet_color=color,
                        anchor=MSO_ANCHOR.TOP, bold_prefix=(n == 2))
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def ecuaciones(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        y, h = self.subtitulo(s, d, y, h)
        eqs = d.get("ecuaciones", [])
        n = max(1, len(eqs))
        rh = min(1.35, h / n)
        y0 = y + (h - rh * n) / 2
        name_w = 3.7
        eq_w = SLIDE_W - 2 * MX - name_w - 0.25
        # un solo tamaño para todas las ecuaciones de la lámina: la más larga manda
        es = 30
        while es > 18 and any(est_lines(str(e.get("texto", "")), es, eq_w, char_w=0.56) > 1 for e in eqs):
            es -= 2
        for i, e in enumerate(eqs):
            ry = y0 + i * rh
            add_rect(s, MX, ry + 0.10, SLIDE_W - 2 * MX, rh - 0.20, fill=CARD, line=LINE, rounded=True, radius=0.07)
            add_rect(s, MX, ry + 0.10, 0.085, rh - 0.20, fill=ACCENT)
            add_text(s, MX + 0.3, ry + 0.10, name_w - 0.3, rh - 0.20, str(e.get("nombre", "")), 18, bold=True,
                     color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
            add_text(s, MX + name_w, ry + 0.10, eq_w, rh - 0.20, str(e.get("texto", "")), es, color=NAVY,
                     font=MATH, anchor=MSO_ANCHOR.MIDDLE)
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def demo(self, d):
        s = self.new_slide(WHITE)
        self.header(s, d)
        y, h = self.content_box(d)
        img = self.fig(d.get("imagen") or "fig_app_completa")
        iw = 7.9
        add_picture_card(s, img, MX, y, iw, h)
        rx = MX + iw + 0.4
        rw = SLIDE_W - MX - rx
        pasos = d.get("bullets") or [p.get("etiqueta", "") for p in d.get("pasos", [])]
        add_text(s, rx, y, rw, 0.42, d.get("subtitulo", "En vivo"), 20, bold=True, color=ACCENT,
                 anchor=MSO_ANCHOR.MIDDLE)
        size = fit_size([str(p) for p in pasos], rw - 0.65, h - 0.7, 20, 15, line_h=1.15, gap_pt=22)
        py = y + 0.58
        for i, p in enumerate(pasos):
            step_h = est_lines(str(p), size, rw - 0.65) * size * 1.15 / 72 + 0.30
            circ = add_oval(s, rx, py + 0.04, 0.46, NAVY)
            shape_text(circ, str(i + 1), 17, bold=True, color=WHITE, margins=(0, 0, 0, 0))
            add_text(s, rx + 0.62, py, rw - 0.62, step_h, str(p), size, color=INK, anchor=MSO_ANCHOR.TOP,
                     line_spacing=1.15)
            py += step_h
        self.remate(s, d.get("remate"))
        self.footer(s, d)
        set_notes(s, d.get("notas"))
        return s

    def cierre(self, d):
        s = self.new_slide(NAVY_DEEP)
        m = self.meta
        add_text(s, 0.8, 1.35, 8.0, 1.2, d.get("titulo", "Gracias"), 54, bold=True, color=WHITE,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, 0.85, 2.60, 1.3, 0.07, fill=ACCENT)
        add_text(s, 0.8, 2.88, 8.0, 1.2, d.get("subtitulo", ""), 21, color=ON_DARK, line_spacing=1.15)
        chips = d.get("chips") or []
        if chips:
            gap = 0.22
            cw = (8.0 - gap * (len(chips) - 1)) / len(chips)
            for i, c in enumerate(chips):
                chip = add_rect(s, 0.8 + i * (cw + gap), 4.40, cw, 0.62, fill=None, line=NAVY_SOFT,
                                rounded=True, radius=0.28)
                shape_text(chip, str(c), 14, bold=True, color=ON_DARK, margins=(0.08, 0.02, 0.08, 0.02))
        elif d.get("bullets"):
            add_bullets(s, 0.8, 4.20, 8.0, 1.5, d["bullets"], size=20, color=ON_DARK, bullet_color=ACCENT,
                        min_size=16, bold_prefix=False)
        add_text(s, 0.8, 5.72, 8.0, 0.5, m.get("autor", ""), 24, bold=True, color=WHITE)
        add_text(s, 0.8, 6.22, 8.0, 0.45, m.get("repo", "github.com/Sucullani/SoftwareED"), 19, color=ACCENT)
        add_picture_fit(s, self.fig("logo_universidad"), 9.55, 1.15, 2.6, 2.6, border=False)
        add_picture_fit(s, self.fig("fig_logo_edufem"), 9.75, 4.15, 2.2, 2.2, border=False)
        set_notes(s, d.get("notas"))
        return s

    # ---- construcción ----------------------------------------------------
    def build(self):
        builders = {
            "portada": self.portada, "seccion": self.seccion, "bullets": self.bullets,
            "dos_columnas": self.dos_columnas, "imagen": self.imagen, "tabla": self.tabla, "kpi": self.kpi,
            "flujo": self.flujo, "comparacion": self.comparacion, "ecuaciones": self.ecuaciones,
            "demo": self.demo, "cierre": self.cierre,
        }
        for d in self.slides_data:
            layout = d.get("layout", "bullets")
            fn = builders.get(layout)
            if fn is None:
                aviso(f"layout desconocido '{layout}' en la diapositiva {d.get('n')}; se usa bullets")
                fn = self.bullets
            print(f"  [{d.get('n'):>2}] {layout:<13} {d.get('titulo', '')[:60]}")
            fn(d)
        return self.prs


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    with open(sys.argv[1], encoding="utf-8") as fh:
        guion = json.load(fh)
    out = sys.argv[2]
    deck = Deck(guion)
    prs = deck.build()
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print(f"\nGuardado: {out}  ({len(prs.slides)} diapositivas)")
    if AVISOS:
        print(f"{len(AVISOS)} avisos de ajuste.")


if __name__ == "__main__":
    main()
